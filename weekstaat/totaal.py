"""Eén Excel voor alle medewerkers samen: vooraan het simpele deel (Samenvatting, Actielijst en alle facturen zonder
order), daarachter de verdieping (het uitgebreide totaal, alle opmerkingen en per medewerker de bladen Stand en Zonder
order).

De cijfers per medewerker komen uit de `Overzicht`-objecten en staan als waarden in het blad Totaal uitgebreid; de
optelling en de verwijzingen in de Samenvatting zijn formules. De bladen per medewerker zijn dezelfde bladen als in het
losse overzicht (`werkboek.stand` en `werkboek.zonder_order`), als waarden en zonder leeswijzer. Tekst uit een bron
(opmerkingen, acties, namen) gaat altijd langs `werkboek._rij`, zodat hij nooit een formule wordt."""

from collections import Counter
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

from . import overzicht as model
from . import werkboek
from .aansluiten import OW
from .instellingen import Instellingen
from .opmaak import nl_bedrag
from .overzicht import Overzicht
from .werkboek import EUR, KOP_VULLING, LETTER, TAB_DIEPTE, TAB_HOOFD, Formule, _laatste, _rij, _som

SAMENVATTING, ACTIELIJST, ALLE, TOT, OPMERKINGEN = (
    "Samenvatting", "Actielijst", "Zonder order alle", "Totaal uitgebreid", "Opmerkingen"
)  # fmt: skip
VOORKANT = (SAMENVATTING, ACTIELIJST, ALLE)
ALGEMEEN = "Algemeen"
GEEL = PatternFill("solid", fgColor="FFF2CC")
ONGELDIG = "/\\?*[]:"  # tekens die in een bladnaam van Excel niet mogen
MAX_KORT = 31 - len("Zonder order ")  # de langste bladnaam is "Zonder order <kort>"


# ---------- korte namen voor de bladen ----------


def _schoon(tekst: str) -> str:
    return "".join("_" if t in ONGELDIG else t for t in tekst).strip("'")


def korte_namen(namen: list[str]) -> dict[str, str]:
    """De korte naam van elke medewerker voor de bladnaam: de voornaam; bij een dubbele voornaam de voornaam plus de
    eerste letter van het laatste woord. Tekens die in een bladnaam niet mogen worden `_`, de naam is zo kort dat
    "Zonder order " plus de naam hooguit 31 tekens is, en na het afkappen blijft elke naam uniek (Excel let niet op
    hoofdletters)."""
    delen = {n: n.split() or [n] for n in namen}
    telling = Counter(_schoon(d[0]).lower() for d in delen.values())
    uit, gebruikt = {}, set()
    for naam, d in delen.items():
        voor = _schoon(d[0])
        if telling[voor.lower()] > 1 and len(d) > 1:
            kort = f"{voor[: MAX_KORT - 2]} {_schoon(d[-1])[:1]}"
        else:
            kort = voor[:MAX_KORT]
        basis, i = kort.strip() or "_", 1
        kort = basis
        while kort.lower() in gebruikt:  # na afkappen of bij dezelfde beginletter: de basisnaam met één volgnummer
            i += 1
            kort = f"{basis[: MAX_KORT - len(str(i)) - 1].rstrip()} {i}"
        gebruikt.add(kort.lower())
        uit[naam] = kort
    return uit


def _wie_bij_wie(inst: Instellingen) -> str:
    """'Uren van <bureau> bij <opdrachtgever>. ' als de instellingen eigen namen hebben; bij de standaardnamen ('het
    bureau', 'de opdrachtgever') leest dat niet en staat er niets."""
    if (inst.bureau, inst.opdrachtgever) == (Instellingen.bureau, Instellingen.opdrachtgever):  # de standaardwaarden
        return ""
    return f"Uren van {inst.bureau} bij {inst.opdrachtgever}. "


def _letterlijk(tekst: str) -> str:
    """Een tekst als tekenreeks in een formule: aanhalingstekens verdubbelen."""
    return '"' + tekst.replace('"', '""') + '"'


# ---------- Zonder order alle ----------


class _Alle:
    """Waar de kolommen en rijen van blad "Zonder order alle" staan, voor de formules van blok 4."""

    def __init__(self, kolommen: list[str], eerste: int, laatste: int):
        self.kolom = {k: L(i + 1) for i, k in enumerate(kolommen)}
        self.eerste, self.laatste = eerste, laatste

    def bereik(self, sleutel: str) -> str:
        k = self.kolom[sleutel]
        return f"'{ALLE}'!${k}${self.eerste}:${k}${self.laatste}"

    def som(self, sleutel: str, naam: str) -> str:
        """De formule (zonder =) die kolom `sleutel` optelt voor de rijen van deze medewerker."""
        return f"SUMPRODUCT(({self.bereik('medewerker')}={_letterlijk(naam)})*{self.bereik(sleutel)})"

    def aantal(self, naam: str) -> str:
        return f"SUMPRODUCT(({self.bereik('medewerker')}={_letterlijk(naam)})*1)"


def _zonder_order_alle(ws, overzichten: list[Overzicht], met_bank: bool, met_toezegging: bool) -> _Alle:
    """Elke factuur zonder order van alle medewerkers, overgenomen uit blok A van hun blad Zonder order."""
    inst = overzichten[0].inst
    kolommen = [  # sleutel, kop, breedte
        ("medewerker", "Medewerker", 22), ("factuur", "Factuur", 15), ("factuurdatum", "Factuurdatum", 12),
        ("weken", "Week", 13), ("dagen", "Dagen zonder order", 26), ("urenstaten", "Urenstaat", 22),
        ("uren", "Uren", 8), ("bedrag_uren", "Bedrag uren", 12), ("onkosten", "Onkosten", 11),
        ("zonder_order", "Zonder order", 13), ("factuur_totaal", "Factuur totaal", 12),
        ("order_bij", f"Order van {inst.opdrachtgever}", 46),
    ]  # fmt: skip
    if met_bank:
        kolommen.append(("ontvangen", f"Ontvangen van {inst.opdrachtgever}", 12))
    kolommen.append(("ander_stuk", f"Ander schriftelijk stuk van {inst.opdrachtgever}", 50))
    if met_toezegging:
        kolommen += [
            ("op_toezegging", "Waarvan uren op een toezegging", 15),
            ("niet_op_toezegging", "Niet op een toezegging: uren zonder enig stuk, en onkosten", 15),
        ]
    kolommen.append(("toelichting", "Toelichting", 90))
    sleutels = [k for k, _, _ in kolommen]
    _rij(ws, [f"Gefactureerd door {inst.bureau} zonder order van {inst.opdrachtgever} - alle medewerkers"])
    _rij(
        ws,
        [
            "Elke factuur (of het deel ervan) dat op geen enkele order staat. Bedragen ex btw. Overgenomen uit blok A "
            "van blad Zonder order per medewerker; de specificatie per dag staat in het overzicht van die medewerker."
        ],
    )
    _rij(ws, [])
    _rij(ws, [kop for _, kop, _ in kolommen])
    eerste = _laatste(ws) + 1
    n = 0
    for o in overzichten:
        for x in o.zonder_order.to_dict("records"):
            _rij(ws, [o.medewerker if k == "medewerker" else x.get(k) for k in sleutels])
            n += 1
    laatste = max(eerste, eerste + n - 1)  # een lege tabel houdt een geldig bereik in de formules
    alle = _Alle(sleutels, eerste, laatste)
    som = ["uren", "bedrag_uren", "onkosten", "zonder_order", "ontvangen", "op_toezegging", "niet_op_toezegging"]
    totaal = ["Totaal"] + [
        _som(alle.kolom[k], eerste, eerste + n - 1) if k in som and k in alle.kolom else None for k in sleutels[1:]
    ]
    _rij(ws, totaal)
    kop_rij, totaal_rij = eerste - 1, _laatste(ws)
    for i, (_, _, breedte) in enumerate(kolommen, start=1):
        ws.column_dimensions[L(i)].width = breedte
    for r in ws.iter_rows():
        for c in r:
            c.font = Font(name=LETTER, bold=c.row in (1, kop_rij, totaal_rij), size=14 if c.row == 1 else 11)
            if c.row >= kop_rij:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                sleutel = sleutels[c.column - 1]
                if c.row > kop_rij and sleutel == "factuurdatum":
                    c.number_format = "DD-MM-YYYY"
                elif c.row > kop_rij and sleutel in som + ["factuur_totaal"]:
                    c.number_format = "#,##0.00"
            if c.row == kop_rij:
                c.fill = KOP_VULLING
    ws.freeze_panes = f"C{eerste}"
    ws.auto_filter.ref = f"A{kop_rij}:{L(len(kolommen))}{max(kop_rij, eerste + n - 1)}"
    return alle


# ---------- Totaal uitgebreid ----------


class _Blad:
    """Het blad Totaal uitgebreid: kolom A is het label, dan een kolom per medewerker, de kolom Totaal en de uitleg."""

    def __init__(self, ws, korte: list[str]):
        self.ws, self.korte, self.n = ws, korte, len(korte)
        self.kolommen = [L(2 + i) for i in range(self.n)]
        self.tot, self.uitleg = L(2 + self.n), L(3 + self.n)
        self.koppen, self.vet, self.tellers = [], [], []

    def kop(self, titel: str, uitleg: str, totaal: str = "Totaal") -> None:
        _rij(self.ws, [titel, *self.korte, totaal, uitleg])
        self.koppen.append(_laatste(self.ws))

    def regel(self, label: str, per_medewerker: list, uitleg: str = "", vet=False, tel=False) -> int:
        """Eén rij: per medewerker een waarde of formule en het totaal erachter (de som)."""
        i = _laatste(self.ws) + 1
        _rij(self.ws, [label, *per_medewerker, Formule(f"=SUM(B{i}:{self.kolommen[-1]}{i})"), uitleg])
        if vet:
            self.vet.append(i)
        if tel:
            self.tellers.append(i)
        return i

    def afgeleid(self, label: str, formule, uitleg: str = "", vet=False) -> int:
        """Een rij waarvan elke kolom, ook het totaal, een formule uit de rijen erboven is: `formule(kolom)`."""
        i = _laatste(self.ws) + 1
        _rij(self.ws, [label, *[Formule(formule(k)) for k in [*self.kolommen, self.tot]], uitleg])
        if vet:
            self.vet.append(i)
        return i

    def leeg(self) -> None:
        _rij(self.ws, [])


def _stand_per_medewerker(t: _Blad, overzichten: list[Overzicht], met_bank: bool) -> dict[str, int]:
    """Blok 1: van urenstaat tot betaling."""
    inst, ovs, r = overzichten[0].inst, overzichten, {}
    t.kop("1. Van urenstaat tot betaling", "Toelichting")
    r["urenstaat"] = t.regel(
        "Gewerkt volgens de goedgekeurde urenstaten", [o.stand["urenstaat"] for o in ovs],
        "Wat volgens de goedgekeurde urenstaten is gewerkt.",
    )  # fmt: skip
    r["gefactureerd"] = t.regel(
        f"Door {inst.bureau} gefactureerd", [o.stand["gefactureerd"] for o in ovs], "Netto, na creditfacturen."
    )
    r["orders"] = t.regel(
        f"Door {inst.opdrachtgever} op orders gezet", [o.stand["op_orders"] for o in ovs],
        f"De orders van {inst.opdrachtgever} met regels van deze medewerker.",
    )  # fmt: skip
    if met_bank:
        r["betaald"] = t.regel(
            f"Door {inst.opdrachtgever} betaald", [o.stand["betaald"] for o in ovs],
            f"Ontvangen op de bank (t/m {ovs[0].aansluiting.bank_tot:%d-%m-%Y}) of in de boekhouding afgeletterd, "
            "teruggerekend naar ex btw.",
        )  # fmt: skip
    r["meer"] = t.afgeleid(
        "Meer gefactureerd dan de goedgekeurde urenstaten",
        lambda k: f"=ROUND({k}{r['gefactureerd']}-{k}{r['urenstaat']},2)",
        "Een negatief bedrag is minder gefactureerd dan de urenstaten. Per week in blad Per week van het overzicht.",
    )
    r["verschil"] = t.afgeleid(
        "Verschil tussen factuur en order (per saldo)",
        lambda k: f"=ROUND({k}{r['gefactureerd']}-{k}{r['orders']},2)",
        "Factuur min order. Per saldo: wat zonder order is gefactureerd min wat wel op een order staat maar niet is "
        "gefactureerd, plus verschillen in tarief en onkosten. Uitgesplitst in blok 2.",
        vet=True,
    )  # fmt: skip
    if met_bank:
        r["open"] = t.afgeleid(
            "Op een order, nog niet betaald",
            lambda k: f"=ROUND({k}{r['orders']}-{k}{r['betaald']},2)",
            "Orders min betaald. Welke orders dat zijn, staat in blad Orders van het overzicht.",
            vet=True,
        )  # fmt: skip
    t.leeg()
    return r


def _oorzaken_en_toewijzingen(t: _Blad, overzichten: list[Overzicht]) -> None:
    """Blok 2 (het verschil per oorzaak, met controle op blok 1) en blok 3 (wie is aan zet)."""
    inst = overzichten[0].inst
    uitleg = model.uitleg(inst)
    per_oorzaak = [o.per_oorzaak().set_index("oorzaak").bedrag for o in overzichten]
    t.kop("2. Waar het verschil tussen factuur en order uit bestaat", "Uitleg")
    eerste = _laatste(t.ws) + 1
    for oorzaak in model.OORZAKEN:
        if any(oorzaak in p.index for p in per_oorzaak):
            t.regel(oorzaak, [float(p.get(oorzaak, 0.0)) for p in per_oorzaak], uitleg[oorzaak])
    i = _laatste(t.ws) + 1
    controle = (
        "Controle: gelijk aan het verschil tussen factuur en order in blok 1"
        if all(o.controles()[model.C_OORZAKEN].sluit for o in overzichten)
        else "LET OP: sluit niet aan op blok 1"
    )
    if _laatste(t.ws) < eerste:  # geen enkel verschil
        _rij(t.ws, ["Totaal", *[0] * (t.n + 1), controle])
    else:
        _rij(t.ws, ["Totaal", *[Formule(f"=SUM({k}{eerste}:{k}{i - 1})") for k in [*t.kolommen, t.tot]], controle])
    t.vet.append(i)
    t.leeg()
    t.kop(
        "3. Wie is aan zet (voorstel; zie blad Verschillen in het overzicht van elke medewerker)",
        "Uitleg. Let op: per toewijzing zijn de bedragen bruto opgeteld (te veel en te weinig vallen niet tegen elkaar "
        "weg). Dit blok sluit daarom niet aan op blok 2.",
    )
    wie = [o.per_toewijzing().set_index("toewijzing").bedrag for o in overzichten]
    for toewijzing, tekst in model.toewijzing_uitleg(inst).items():
        t.regel(toewijzing, [float(w[toewijzing]) for w in wie], tekst)
    saldo = [o.per_saldo() for o in overzichten]
    for label, sleutel in (
        (f"Per saldo door {inst.opdrachtgever} nog op te nemen", "opnemen"),
        (f"Per saldo door {inst.bureau} nog te factureren", "factureren"),
        ("Nog uit te zoeken", "uitzoeken"),
    ):
        t.regel(label, [s[sleutel] for s in saldo], vet=True)
    if any(o.toezegging is not None for o in overzichten):  # wat van "te weinig opgenomen" al op een toezegging staat
        nummers = ", ".join(
            dict.fromkeys(
                n
                for o in overzichten
                if o.toezegging is not None
                for n in o.toezegging_bedragen()["nummers"].split(", ")
            )
        )
        t.regel(
            f"Van '{OW}' staat al op toezegging {nummers}",
            [float(p.get(model.O_TOEZEGGING, 0.0)) for p in per_oorzaak],
            f"Een schriftelijk stuk van {inst.opdrachtgever}, geen order: zie blad Zonder order in het overzicht, "
            "blok D.",
        )
    t.leeg()


def _zonder_order_blok(t: _Blad, overzichten: list[Overzicht], alle: _Alle, met_bank: bool, met_toez: bool) -> dict:
    """Blok 4: wat zonder order is gefactureerd, opgeteld uit het blad Zonder order alle."""
    inst, namen, r = overzichten[0].inst, [o.medewerker for o in overzichten], {}
    t.kop("4. Gefactureerd zonder order", f"Elke factuur staat in blad {ALLE}")
    r["aantal"] = t.regel(
        "Aantal facturen (of delen ervan)", [Formule('=' + alle.aantal(n)) for n in namen],
        "Facturen waarvan dagen op geen enkele order staan.", tel=True,
    )  # fmt: skip
    r["uren"] = t.regel("Uren", [Formule("=" + alle.som("uren", n)) for n in namen], "Het aantal uren, geen bedrag.")
    r["totaal"] = t.regel(
        "Totaal zonder order", [Formule("=" + alle.som("zonder_order", n)) for n in namen],
        "Uren plus onkosten. Meer dan het verschil in blok 1: daar is afgetrokken wat wel op een order staat en "
        "niet is gefactureerd.", vet=True,
    )  # fmt: skip
    if met_toez:
        r["toezegging"] = t.regel(
            "Waarvan uren die wel op een toezegging staan",
            [Formule("=" + alle.som("op_toezegging", n)) for n in namen],
            f"Een schriftelijk stuk van {inst.opdrachtgever} dat geen order is; zie blad Zonder order in het "
            "overzicht, blok D.",
        )  # fmt: skip
    r["geen_stuk"] = t.regel(
        f"Waarvan uren zonder enig stuk van {inst.opdrachtgever}",
        [
            Formule(f"=ROUND({alle.som('bedrag_uren', n)}-{alle.som('op_toezegging', n) if met_toez else 0},2)")
            for n in namen
        ],
        "Geen order en geen toezegging.", vet=True,
    )  # fmt: skip
    r["onkosten"] = t.regel(
        "Waarvan onkosten", [Formule("=" + alle.som("onkosten", n)) for n in namen],
        "Alleen onkosten op dagen zonder order. Onkosten op dagen waarvan de uren wel op een order staan, staan in "
        "blok 2.",
    )  # fmt: skip
    if met_bank:
        r["ontvangen"] = t.regel(
            f"Ontvangen van {inst.opdrachtgever} voor deze facturen",
            [Formule("=" + alle.som("ontvangen", n)) for n in namen],
            "Een betaling hoort bij een order; deze dagen hebben er geen.",
        )  # fmt: skip
    t.leeg()
    return r


def _kaart_blok(t: _Blad, overzichten: list[Overzicht]) -> dict[str, int]:
    """Blok 5: wat echt openstaat volgens de debiteurenkaart, met de controle op het saldo van de kaart."""
    inst, ks = overzichten[0].inst, [o.kaart for o in overzichten]
    namen, heel = [o.medewerker for o in overzichten], ks[0].heel
    t.kop(
        "5. Wat staat er echt open (incl. btw)",
        "Uit blok B van blad Debiteurenkaart in het overzicht van elke medewerker",
    )
    r = {}
    r["open"] = t.regel(
        "Open facturen op de debiteurenkaart", [k.open for k in ks],
        f"Wat in de boekhouding van {inst.bureau} openstaat. Een deel daarvan is wel betaald maar niet afgeletterd.",
    )  # fmt: skip
    r["echt"] = t.regel(
        "Echt open na afletteren", [k.echt_open for k in ks],
        f"Open facturen min wat {inst.opdrachtgever} op de orders van deze medewerker betaalde.", vet=True,
    )  # fmt: skip
    r["saldo"] = t.regel(
        f"Per saldo te vorderen op {inst.opdrachtgever}", [k.per_saldo for k in ks],
        f"Na aftrek van wat {inst.opdrachtgever} betaalde zonder dat {inst.bureau} het factureerde.", vet=True,
    )  # fmt: skip
    ws = t.ws
    for sleutel, label, waarde, tekst in (
        ("buiten", "Open facturen die niet in de uren staan", heel["buiten_uren"],
         "Facturen op de kaart waarvan het nummer bij geen enkele urenregel voorkomt."),
        ("kaart_open", "Hele debiteurenkaart: open facturen", heel["open"],
         "Alle medewerkers, ook facturen die niet in de uren staan."),
        ("kaart_ontvangsten", "Hele debiteurenkaart: ontvangsten, niet afgeletterd", heel["ontvangsten"],
         "Betalingen die nog aan geen factuur zijn gekoppeld (negatief)."),
        ("kaart_saldo", "Hele debiteurenkaart: saldo", heel["saldo"], ""),
    ):  # fmt: skip
        _rij(ws, [label, *[None] * t.n, waarde, tekst])
        r[sleutel] = _laatste(ws)
    zonder_overzicht = [n for n in ks[0].per_medewerker.index if n not in namen]
    controle_label = "Controle: per saldo plus de facturen buiten de uren, min het saldo van de kaart"
    if zonder_overzicht:
        open_ = sum(float(ks[0].per_medewerker[n]) for n in zonder_overzicht)
        _rij(
            ws,
            [
                controle_label, *[None] * t.n, None,
                f"Niet gecontroleerd: de kaart heeft ook open facturen van {', '.join(zonder_overzicht)} "
                f"({nl_bedrag(open_)} incl. btw), die niet in dit bestand staan.",
            ],
        )  # fmt: skip
    else:
        verschil = round(sum(k.per_saldo for k in ks) + heel["buiten_uren"] - heel["saldo"], 2)
        tekst = (
            "Sluit aan op het saldo van de debiteurenkaart"
            if abs(verschil) < model.CONTROLE_VERSCHIL
            else f"LET OP: sluit niet aan op het saldo van de debiteurenkaart, verschil {nl_bedrag(verschil)}"
        )
        controle = f"=ROUND({t.tot}{r['saldo']}+{t.tot}{r['buiten']}-{t.tot}{r['kaart_saldo']},2)"
        _rij(ws, [controle_label, *[None] * t.n, Formule(controle), tekst])
    for o in overzichten:  # een kaart die niet sluit op bank en orders per factuur, per medewerker
        if not o.kaart.sluit:
            verschil = o.controles()[model.C_KAART].verschil
            _rij(
                ws,
                [
                    "Controle: kaart per medewerker", *[None] * t.n, None,
                    f"LET OP: de debiteurenkaart van {o.medewerker} sluit niet aan op bank en orders per factuur, "
                    f"verschil {nl_bedrag(abs(verschil))}.",
                ],
            )  # fmt: skip
    t.leeg()
    return r


def _belangrijk(t: _Blad, opmerkingen: pd.DataFrame, korte: dict[str, str]) -> None:
    """Blok 6: de opmerkingen met belangrijk = ja, het bedrag onder de medewerker (of onder Totaal)."""
    punten = opmerkingen[opmerkingen.belangrijk == "ja"]
    if punten.empty:
        return
    t.kop(
        "6. De belangrijkste punten (de bedragen overlappen: niet optellen)",
        "Wat er aan de hand is en wat te doen (alle punten, met bewijs, in blad Opmerkingen)",
    )
    for x in punten.itertuples():
        voor = korte.get(x.medewerker)
        voor = voor if voor in t.korte else None  # een medewerker die niet in dit bestand staat, telt als Algemeen
        bedrag = None if pd.isna(x.bedrag) else float(x.bedrag)
        naam = f"{voor or ALGEMEEN} {x.nr}. {x.onderwerp}" if x.nr else f"{voor or ALGEMEEN}: {x.onderwerp}"
        waarden = [bedrag if k == voor else None for k in t.korte]
        tekst = f"{x.wat} Te doen: {x.actie}" if x.actie else x.wat
        _rij(t.ws, [naam + (" (incl. btw)" if x.btw == "incl" else ""), *waarden, None if voor else bedrag, tekst])
    t.leeg()


def _opmaak_totaal(t: _Blad) -> None:
    ws = t.ws
    for k, b in zip(["A", *t.kolommen, t.tot, t.uitleg], (62, *[16] * (t.n + 1), 110), strict=True):
        ws.column_dimensions[k].width = b
    for rij in ws.iter_rows():
        for c in rij:
            c.font = Font(name=LETTER, bold=c.row in t.koppen + t.vet or c.row == 1, size=14 if c.row == 1 else 11)
            c.alignment = Alignment(wrap_text=c.column == t.n + 3, vertical="top")
            if 2 <= c.column <= t.n + 2 and c.row not in t.koppen:
                c.number_format = "0" if c.row in t.tellers else EUR
            if c.row in t.koppen:
                c.fill = KOP_VULLING
    ws.freeze_panes = "B5"


def _leeswijzer(t: _Blad) -> None:
    _rij(t.ws, ["Zo lees je dit bestand"])
    t.koppen.append(_laatste(t.ws))
    for blad, tekst in (
        (SAMENVATTING, "De kern op één scherm: de cijfers van dit bestand zonder de uitleg."),
        (ACTIELIJST, "Wat er te doen is, per soort actie."),
        (TOT, "Dit blad: de medewerkers naast elkaar en opgeteld, met uitleg per regel."),
        (
            OPMERKINGEN,
            "Alles wat bij de controle is opgevallen of nog onduidelijk is, met bewijs en wat te doen. Kolom Bron zegt "
            "of een opmerking van jou komt of door de tool is gevonden.",
        ),
        (ALLE, "Elke factuur zonder order van alle medewerkers, met toelichting."),
        (
            "Stand <naam>",
            "Per medewerker: urenstaat, factuur, order en betaling, het verschil per oorzaak en wie aan zet is.",
        ),
        ("Zonder order <naam>", "Per medewerker: de facturen zonder order per factuur en per dag."),
        (
            "Let op",
            "De bladen per medewerker en de opmerkingen zijn overgenomen uit het overzicht van die medewerker. "
            "Verwijzingen naar andere bladen (Per week, Per dag, Verschillen, Debiteurenkaart, Bank, Orders) gelden "
            "voor Overzicht <naam>.xlsx: daar staat de verdieping tot op de dag en de bronregel.",
        ),
    ):
        _rij(t.ws, [blad, *[None] * (t.n + 1), tekst])


# ---------- Samenvatting, Actielijst en Opmerkingen ----------


def _samenvatting(ws, t: _Blad, r1: dict, r4: dict, r5: dict | None, overzichten: list[Overzicht]) -> None:
    inst, bank_tot = overzichten[0].inst, overzichten[0].aansluiting.bank_tot
    toezegging = (
        [
            (
                "Waarvan uren op een toezegging",
                r4["toezegging"],
                f"Een schriftelijk stuk van {inst.opdrachtgever}, geen order.",
            )
        ]
        if "toezegging" in r4
        else []
    )
    verwijs = lambda rij: [Formule(f"='{TOT}'!{k}{rij}") for k in [*t.kolommen, t.tot]]  # noqa: E731
    _rij(ws, ["Aansluiting in het kort"])
    _rij(
        ws,
        [
            f"{_wie_bij_wie(inst)}Gemaakt op {date.today():%d-%m-%Y}. Wat er te doen is, staat in blad {ACTIELIJST}. "
            "De uitleg en de onderbouwing staan in de grijze bladen."
        ],
    )
    _rij(ws, [])
    koppen, vet = [], []
    blokken = [
        (
            "Van urenstaat tot betaling (ex btw)",
            [
                ("Gewerkt volgens de goedgekeurde urenstaten", r1["urenstaat"], ""),
                (f"Door {inst.bureau} gefactureerd", r1["gefactureerd"], ""),
                (f"Door {inst.opdrachtgever} op orders gezet", r1["orders"], ""),
                *(
                    [(f"Door {inst.opdrachtgever} betaald", r1["betaald"], f"Bank t/m {bank_tot:%d-%m-%Y}.")]
                    if "betaald" in r1
                    else []
                ),
            ],
        ),
        (
            "Gefactureerd zonder order (ex btw)",
            [
                ("Totaal zonder order", r4["totaal"], f"Per factuur in blad {ALLE}."),
                *toezegging,
                (
                    f"Waarvan uren zonder enig stuk van {inst.opdrachtgever}",
                    r4["geen_stuk"],
                    "Geen order en geen toezegging.",
                ),
                ("Waarvan onkosten", r4["onkosten"], ""),
            ],
        ),
    ]  # fmt: skip
    if r5 is not None:
        blokken.append(
            (
                "Wat staat er echt open (incl. btw)",
                [
                    ("Open na afletteren", r5["echt"], "Volgens de debiteurenkaart."),
                    (
                        f"Per saldo te vorderen op {inst.opdrachtgever}", r5["saldo"],
                        f"Na aftrek van wat {inst.opdrachtgever} betaalde zonder dat {inst.bureau} het factureerde.",
                    ),
                ],
            )
        )  # fmt: skip
    for i, (kop, regels) in enumerate(blokken):
        if i:
            _rij(ws, [])
        _rij(ws, [kop, *t.korte, "Totaal"])
        koppen.append(_laatste(ws))
        for label, rij, noot in regels:
            _rij(ws, [label, *verwijs(rij), noot])
            if label.startswith(("Totaal zonder", "Per saldo")):
                vet.append(_laatste(ws))
    for i, b in enumerate((46, *[14] * (t.n + 1), 58), start=1):
        ws.column_dimensions[L(i)].width = b
    for rij in ws.iter_rows():
        for c in rij:
            c.font = Font(
                name=LETTER, bold=c.row in koppen + vet or c.row == 1, size=14 if c.row == 1 else 11,
                italic=c.column == t.n + 3, color="595959" if c.column == t.n + 3 else "000000",
            )  # fmt: skip
            if c.row in koppen:
                c.fill = KOP_VULLING
                c.alignment = Alignment(horizontal="right" if c.column > 1 else "left")
            elif 2 <= c.column <= t.n + 2:
                c.number_format = EUR


def _actielijst(ws, acties: list[tuple[str, list[tuple[str, str, str]]]]) -> None:
    _rij(ws, ["Actielijst"])
    _rij(
        ws,
        [
            "Wat er te doen is, per soort actie. Voor wie: de medewerker om wie het gaat, of "
            f"{ALGEMEEN}. De onderbouwing staat in blad {OPMERKINGEN}."
        ],
    )
    _rij(ws, [])
    koppen = []
    for groep, lijst in acties:
        _rij(ws, [groep, "Voor wie", "Toelichting"])
        koppen.append(_laatste(ws))
        for titel, voor_wie, toelichting in lijst:
            _rij(ws, [titel, voor_wie, toelichting])
        _rij(ws, [])
    for k, b in zip("ABC", (58, 16, 120), strict=True):
        ws.column_dimensions[k].width = b
    for rij in ws.iter_rows():
        for c in rij:
            c.font = Font(
                name=LETTER,
                bold=c.row in koppen or c.row == 1 or (c.column == 1 and c.row > 3),
                size=14 if c.row == 1 else 11,
            )
            c.alignment = Alignment(wrap_text=c.row > 3, vertical="top")
            if c.row in koppen:
                c.fill = KOP_VULLING
    ws.freeze_panes = "A4"


def _opmerkingen_blad(ws, opmerkingen: pd.DataFrame) -> None:
    koppen = [
        "Medewerker",
        "Nr",
        "Belangrijk",
        "Onderwerp",
        "Bedrag",
        "Ex of incl. btw",
        "Wat er aan de hand is",
        "Bewijs",
        "Wat te doen",
        "Wie is aan zet",
        "Bron",
        "Bevinding",  # de sleutel in bevindingen.csv; zo is een opmerking daar terug te vinden
    ]
    ws.append(koppen)
    for c, b in zip(ws[1], (22, 5, 10, 40, 12, 9, 80, 60, 60, 12, 12, 40), strict=True):
        c.font, c.fill = Font(name=LETTER, bold=True), KOP_VULLING
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = b
    bron = opmerkingen["bron"] if "bron" in opmerkingen else pd.Series([""] * len(opmerkingen), index=opmerkingen.index)
    for x, b in zip(opmerkingen.itertuples(), bron, strict=True):
        _rij(
            ws,
            [
                x.medewerker or ALGEMEEN,
                x.nr,
                x.belangrijk,
                x.onderwerp,
                x.bedrag,
                x.btw,
                x.wat,
                x.bewijs,
                x.actie,
                x.wie,
                b,
                x.bevinding,
            ],
        )
    for rij in ws.iter_rows(min_row=2):
        for c in rij:
            c.font = Font(name=LETTER)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.column_letter == "E":
                c.number_format = EUR
            elif c.column_letter == "C" and c.value == "ja":
                c.fill = GEEL
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{L(len(koppen))}{max(ws.max_row, 2)}"


# ---------- het bestand ----------


def schrijf(
    overzichten: list[Overzicht],
    opmerkingen: pd.DataFrame,
    acties: list[tuple[str, list[tuple[str, str, str]]]],
    pad: str | Path,
) -> None:
    """Schrijft het totaalbestand. `opmerkingen` zijn alle opmerkingen (handmatig en automatisch, zie
    `bevindingen.opmerkingen`) en `acties` de actielijst (`bevindingen.acties`). De bladnamen volgen `korte_namen`.
    Blokken zonder bron vallen weg: geen bank, dan geen regel
    betaald; geen debiteurenkaart, dan geen blok 5; niemand met een toezegging, dan geen regel daarover. Klopt de
    debiteurenkaart niet met het saldo, dan staat dat als LET OP in blok 5: het bestand komt er wel."""
    if not overzichten:
        raise ValueError("het totaalbestand heeft minstens één medewerker nodig")
    namen = [o.medewerker for o in overzichten]
    kort = korte_namen(namen)
    met_bank = all(o.bank is not None for o in overzichten)
    met_kaart = all(o.kaart is not None for o in overzichten)
    met_toez = any(o.toezegging is not None for o in overzichten)
    wb = Workbook()
    sam = wb.active
    sam.title = SAMENVATTING
    act, alle_ws, tot_ws, opm_ws = (wb.create_sheet(n) for n in (ACTIELIJST, ALLE, TOT, OPMERKINGEN))

    alle = _zonder_order_alle(alle_ws, overzichten, met_bank, met_toez)

    t = _Blad(tot_ws, [kort[n] for n in namen])
    inst = overzichten[0].inst
    _rij(tot_ws, ["Aansluiting: alle medewerkers samen"])
    _rij(
        tot_ws,
        [
            f"{_wie_bij_wie(inst)}{', '.join(namen)}. Bedragen ex btw, tenzij anders vermeld. "
            f"Gemaakt op {date.today():%d-%m-%Y}. De cijfers per medewerker komen uit de bladen erachter; die zijn "
            "overgenomen uit de losse overzichten."
        ],
    )
    _rij(tot_ws, [])
    r1 = _stand_per_medewerker(t, overzichten, met_bank)
    _oorzaken_en_toewijzingen(t, overzichten)
    r4 = _zonder_order_blok(t, overzichten, alle, met_bank, met_toez)
    r5 = _kaart_blok(t, overzichten) if met_kaart else None
    _belangrijk(t, opmerkingen, kort)
    _leeswijzer(t)
    _opmaak_totaal(t)

    _samenvatting(sam, t, r1, r4, r5, overzichten)
    _actielijst(act, acties)
    _opmerkingen_blad(opm_ws, opmerkingen)
    for o in overzichten:
        eigen = opmerkingen[opmerkingen.medewerker == o.medewerker]
        werkboek.stand(wb, o, eigen, f"Stand {kort[o.medewerker]}", als_waarden=True, leeswijzer=False)
        werkboek.zonder_order(wb, o, f"Zonder order {kort[o.medewerker]}", als_waarden=True)
    for ws in wb:
        ws.sheet_properties.tabColor = TAB_HOOFD if ws.title in VOORKANT else TAB_DIEPTE
        ws.sheet_state = "visible"
    wb.active = 0
    pad = Path(pad)
    pad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pad)
