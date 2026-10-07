"""Schrijft een `Overzicht` als Excel-werkboek: vijf hoofdbladen voorop en de verdieping verborgen erachter.

De bladen rekenen zelf door met formules, zodat de gebruiker een keuze kan aanpassen en het resultaat ziet. Alle cijfers
en teksten komen uit het model (`overzicht.maak`); de schrijver rekent niets opnieuw uit. `stand` en `zonder_order`
kunnen ook de uitkomst uit het model schrijven in plaats van formules (`als_waarden`), voor een totaalbestand dat zo'n
blad los toont. Wat de gebruiker zelf koos in een eerder werkboek (de gele kolom op blad Verschillen) leest
`eigen_keuze`; die keuze gaat als invoer naar het model."""

import re
import zipfile
from pathlib import Path
from typing import NamedTuple

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from . import overzicht as model
from .aansluiten import (
    ALLEEN,
    CORR,
    CV,
    CW,
    GEEN,
    GEEN_D,
    GEEN_P,
    KLOPT,
    KM,
    NOG,
    ONK,
    OV,
    OW,
    TOEWIJZINGEN,
    UREN,
    weeksleutels,
)
from .opmaak import nl_bedrag, percentage
from .overzicht import Overzicht

HOOFDBLADEN = ("Stand", "Zonder order", "Debiteurenkaart", "Per periode", "Verschillen")
TAB_HOOFD, TAB_DIEPTE = "1F4E78", "A6A6A6"

LETTER = "Arial"
KOP_VULLING = PatternFill("solid", fgColor="D9E1F2")
GEEL = PatternFill("solid", fgColor="FFFF00")
VET = Font(name=LETTER, bold=True)
ROOD = Font(name=LETTER, bold=True, color="C00000")
EUR, DATUM = '#,##0.00;[Red]-#,##0.00;"-"', "DD-MM-YYYY"
PERIODEKOP = {"week": "Periode (week)", "4 weken": "Periode (4 weken)", "maand": "Periode (maand)"}
PERIODETEKST = {"week": "een week", "4 weken": "4 weken", "maand": "een maand"}
DAGNAAM = ["ma", "di", "wo", "do", "vr", "za", "zo"]

S_KLEUR = {OW: "FFF2CC", OV: "F8CBAD", CW: "DDEBF7", CV: "F8CBAD", NOG: "D9D9D9"}
DAG_KLEUR = {
    KLOPT: "E2EFDA", KM: "FFF2CC", ONK: "FFF2CC", UREN: "F8CBAD", ALLEEN: "F8CBAD", CORR: "F8CBAD",
    GEEN_P: "D9D9D9", GEEN_D: "EDEDED",
}  # fmt: skip
WEEK_KLEUR = {
    model.W_AANSLUIT: "E2EFDA", model.W_ONKOSTEN: "FFF2CC", model.W_AFWIJKT: "FFF2CC", model.W_DEELS: "EDEDED",
    model.W_GEEN: "D9D9D9", model.W_NIET_GEF: "F8CBAD", model.W_NIET_GEWERKT: "FFFFFF",
    model.W_GEEN_URENSTAAT: "F8CBAD",
}  # fmt: skip
PERIODE_KLEUR = {"Sluit aan": "E2EFDA", "Geen order": "D9D9D9", "Verschil": "FFF2CC"}


class Formule(str):
    """Een formule die we zelf opbouwen. Alle andere tekst die het blad in gaat, is invoer en blijft tekst."""


F = Formule  # kort, want bijna elke rij heeft er een paar


def _fx(als_waarden: bool):
    """`fx(formule, waarde)` geeft de formule, of bij `als_waarden` de uitkomst uit het model."""

    def fx(formule: str, waarde):
        return waarde if als_waarden else Formule(formule)

    return fx


PAUZETEKENS = re.compile("[\x0b\x0c]")  # de zachte regelbreuk van Word en de paginawissel: in Excel niet toegestaan


def _schoon(tekst: str) -> str:
    """Tekst voor een cel zonder de stuurtekens die Excel weigert (openpyxl stopt er anders midden in het schrijven
    mee). Een verticale tab of paginawissel is een regelbreuk in de bron en wordt een spatie, de rest valt weg. Een
    tab en een gewoon regeleinde laat Excel toe en blijven staan."""
    return ILLEGAL_CHARACTERS_RE.sub("", PAUZETEKENS.sub(" ", tekst))


def _cel(v) -> object:
    """Een waarde voor een cel: leeg wordt None en een tijdstip een datum. Alle tekst gaat hierlangs en wordt
    opgeschoond (ook voor `totaal.py`)."""
    if isinstance(v, str):
        return _schoon(str(v))
    return None if pd.isna(v) else v.date() if isinstance(v, pd.Timestamp) else v


def _blijf_tekst(cel, v) -> None:
    """Tekst uit de invoer (een bankomschrijving, een project, een notitie, een naam) die met `=` begint, zou in Excel
    een formule worden; alleen een `Formule` mag dat. Zulke tekst wordt als tekst opgeslagen."""
    if isinstance(v, str) and not isinstance(v, Formule) and cel.data_type == "f":
        cel.data_type = "s"


def _laatste(ws: Worksheet) -> int:
    """Het nummer van de laatste rij die is toegevoegd. `ws.max_row` telt een lege rij aan het eind niet mee, en dan
    schuiven de rijen waar formules naar verwijzen."""
    return ws._current_row


def _rij(ws: Worksheet, waarden: list) -> None:
    """Voegt een rij toe. Elke rij gaat hierlangs, zodat tekst uit de invoer nooit een formule wordt."""
    ws.append([_cel(v) for v in waarden])
    for cel, v in zip(ws[_laatste(ws)], waarden, strict=False):
        _blijf_tekst(cel, v)


def _blad(wb: Workbook, naam: str, koppen: list[str], breedtes: list[int]) -> Worksheet:
    ws = wb.create_sheet(naam)
    ws.append(koppen)
    for c, b in zip(ws[1], breedtes, strict=True):
        c.font, c.fill = Font(name=LETTER, bold=True), KOP_VULLING
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = b
    ws.freeze_panes = "B2"
    return ws


def _opmaak(ws: Worksheet, eur: str = "", datum: str = "") -> None:
    """Lettertype, bedragen en datums per kolomletter, en een filter over het hele blad."""
    for rij in ws.iter_rows(min_row=2):
        for c in rij:
            c.font = Font(name=LETTER, bold=c.font.bold)
            if c.column_letter in eur:
                c.number_format = EUR
            if c.column_letter in datum:
                c.number_format = DATUM
    ws.auto_filter.ref = ws.dimensions


def _omloop(ws: Worksheet, kolommen: str) -> None:
    for k in kolommen:
        for c in ws[k]:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def _totaal(ws: Worksheet, n: int, kolommen: str, eur: str) -> None:
    """Vette totaalregel onder de tabel; het filter loopt tot en met de laatste gewone rij."""
    ws.cell(n + 1, 1, "Totaal")
    for k in kolommen:
        ws[f"{k}{n + 1}"] = _som(k, 2, n)
    for c in ws[n + 1]:
        c.font = VET
        if c.column_letter in eur:
            c.number_format = EUR
    ws.auto_filter.ref = f"A1:{ws.cell(1, ws.max_column).column_letter}{n}"


def _som(kolom: str, eerste: int, laatste: int) -> Formule | int:
    """De som van een kolom over de rijen `eerste` t/m `laatste`; zonder rijen een nul. Een bereik van achter naar voren
    telt de rij met de som zelf mee, en dat is een cirkelverwijzing."""
    return F(f"=SUM({kolom}{eerste}:{kolom}{laatste})") if laatste >= eerste else 0


def _n(aantal: int) -> int:
    """De laatste rij van een tabel van `aantal` rijen onder een kopregel, nooit onder 2: een lege tabel houdt zo een
    geldig bereik in de formules."""
    return max(aantal, 1) + 1


def _bereik(blad: str, kolom: str, n: int) -> str:
    """Het bereik van een kolom in een ander blad, voor SUMIFS en COUNTIFS: 'Per dag'!$K$2:$K$80."""
    return f"'{blad}'!${kolom}$2:${kolom}${n}"


def _eerste(naam: str) -> str:
    return naam.split()[0]


def _geen_bank(o: Overzicht) -> bool:
    return o.bank_los is None


# ---------- wat de gebruiker zelf koos ----------


KEUZE_KOLOM = 14  # kolom N van blad Verschillen (1 is kolom A): de gele kolom "Eigen keuze"
WEEK_KOLOM, OORZAAK_KOLOM = 2, 3  # kolom B en C van blad Verschillen


def eigen_keuze(pad: Path) -> dict[tuple[str, str], str]:
    """De keuzes die de gebruiker in de gele kolom van blad Verschillen invulde, als {(week, oorzaak): toewijzing}.
    Is er geen werkboek op `pad`, dan is het resultaat leeg. Een keuze die geen toewijzing is, geeft een `ValueError`
    met het bestand, het blad en de cel."""
    pad = Path(pad)
    if not pad.exists():
        return {}
    try:
        wb = load_workbook(pad, data_only=True)
    except (OSError, KeyError, zipfile.BadZipFile, InvalidFileException) as fout:
        raise ValueError(
            f"{pad}: het bestaande overzicht is niet te lezen ({fout}). Sluit het in Excel of haal het weg."
        ) from None
    if "Verschillen" not in wb.sheetnames:
        return {}
    rijen = wb["Verschillen"].iter_rows(values_only=True)
    kop = next(rijen, ())
    k = KEUZE_KOLOM - 1  # plaats in een rij
    if len(kop) < KEUZE_KOLOM or not str(kop[k]).startswith("Eigen keuze"):
        return {}
    keuzes = {}
    for nr, r in enumerate(rijen, start=2):
        if len(r) < KEUZE_KOLOM or r[k] in (None, "") or not r[WEEK_KOLOM - 1] or not r[OORZAAK_KOLOM - 1]:
            continue
        keuze = str(r[k]).strip()
        if keuze not in TOEWIJZINGEN:
            raise ValueError(
                f"{pad}: blad Verschillen, cel {L(KEUZE_KOLOM)}{nr}: eigen keuze '{keuze}' is geen van: "
                + ", ".join(TOEWIJZINGEN)
            )
        keuzes[(r[WEEK_KOLOM - 1], r[OORZAAK_KOLOM - 1])] = keuze
    return keuzes


class _Orderkolommen(NamedTuple):
    """De nummers van de kolommen van blad Orders die andere bladen nodig hebben. Er is een kolom per rekening, dus de
    plaats van de rest hangt van de instellingen af."""

    incl: int
    eerste_rekening: int
    totaal: int
    open: int
    datum: int
    betaling: int
    ex: int
    gefactureerd: int
    verschil: int
    bank: int


def _orderkolommen(inst) -> _Orderkolommen:
    na = 10 + len(inst.rekeningen)  # de rekeningen beginnen na 'Order incl. btw' (kolom 9)
    return _Orderkolommen(9, 10, na, na + 1, na + 2, na + 3, na + 4, na + 5, na + 6, na + 7)


# ---------- de vijf hoofdbladen ----------


def stand(
    wb: Workbook,
    o: Overzicht,
    opmerkingen: pd.DataFrame,
    titel: str = "Stand",
    als_waarden: bool = False,
    leeswijzer: bool = True,
) -> Worksheet:
    """Blad Stand: van urenstaat tot betaling, het verschil per oorzaak, wie aan zet is, de bank, de belangrijkste
    opmerkingen en een leeswijzer. Met `als_waarden` staan er getallen en tekst uit het model in plaats van formules.
    Zonder `leeswijzer` blijft die weg: een blad dat los in een ander bestand staat, heeft de bladen waar hij naar
    verwijst niet."""
    fx, inst, vn = _fx(als_waarden), o.inst, _eerste(o.medewerker)
    met_bank, met_kaart, met_opmerkingen = not _geen_bank(o), o.kaart is not None, len(opmerkingen) > 0
    week = lambda k: _bereik("Per week", k, _n(len(o.weken)))  # noqa: E731
    ver = lambda k: _bereik("Verschillen", k, _n(len(o.verschillen)))  # noqa: E731
    orders = lambda k: _bereik("Orders", k, _n(len(o.orders)))  # noqa: E731
    kolom = _orderkolommen(inst)
    ws = wb.create_sheet(titel)
    for k, b in zip("ABCD", (62, 16, 16, 110), strict=True):
        ws.column_dimensions[k].width = b
    koppen, vet = [], []

    def kop(*cellen) -> None:
        _rij(ws, list(cellen))
        koppen.append(_laatste(ws))

    st, stappen, controles = o.stand, o.stappen(), o.controles()
    per = o.per_dag
    _rij(ws, [f"{o.medewerker}: stand van de aansluiting"])
    tekst = (
        f"Periode {per.datum.min():%d-%m-%Y} t/m {per.datum.max():%d-%m-%Y}. Bedragen ex btw, tenzij anders vermeld."
    )
    if met_bank:
        tekst += f" Bank t/m {o.aansluiting.bank_tot:%d-%m-%Y}."
    _rij(ws, [tekst])
    ontbreekt = []  # wat er niet is aangeleverd: één regel in het rood op rij 3
    if not met_bank:
        ontbreekt.append("Er is geen bank aangeleverd: wat betaald is, is niet te zien.")
        ontbreekt.append("Een debiteurenkaart wordt tegen de bank gelegd en blijft daarom ook weg.")
    elif not met_kaart:
        ontbreekt.append("Er is geen debiteurenkaart aangeleverd: wat in de boekhouding openstaat, is niet te zien.")
    _rij(ws, [" ".join(ontbreekt)] if ontbreekt else [])

    # 1. van urenstaat tot betaling
    kop("1. Van urenstaat tot betaling", "Bedrag", "Verschil met de regel erboven", "Toelichting")
    r1 = _laatste(ws)
    weken_met_bedrag = int((~o.weken.status.isin((model.W_NIET_GEWERKT, model.W_GEEN_URENSTAAT))).sum())
    credits = {x for t in o.facturen.creditfacturen for x in t.split(", ") if x}
    _rij(
        ws,
        [
            "Gewerkt volgens de goedgekeurde urenstaten",
            fx(f"=SUM({week('G')})", st["urenstaat"]),
            None,
            f"{weken_met_bedrag} weken met een bedrag; zie blad Per week",
        ],
    )
    _rij(
        ws,
        [
            f"Door {inst.bureau} gefactureerd (netto, na credits)",
            fx(f"=SUM({week('J')})", st["gefactureerd"]),
            fx(f"=ROUND(B{r1 + 1}-B{r1 + 2},2)", stappen["urenstaat_gefactureerd"]),
            f"{len(o.facturen)} facturen en {len(credits)} creditfacturen; zie blad Facturen",
        ],
    )
    _rij(
        ws,
        [
            f"Door {inst.opdrachtgever} op orders gezet",
            fx(f"=SUM({week('M')})", st["op_orders"]),
            fx(f"=ROUND(B{r1 + 2}-B{r1 + 3},2)", stappen["gefactureerd_orders"]),
            f"{len(o.orders)} orders met regels van {vn}; zie blad Orders",
        ],
    )
    if met_bank:
        _rij(
            ws,
            [
                f"Door {inst.opdrachtgever} betaald (het deel van {vn})",
                fx(f"=SUM({orders(L(kolom.ex))})", st["betaald"]),
                fx(f"=ROUND(B{r1 + 3}-B{r1 + 4},2)", stappen["orders_betaald"]),
                "Wat er van de orders op de bank is ontvangen of in de boekhouding is afgeletterd, teruggerekend "
                "naar ex btw; zie blok 4",
            ],
        )
    _rij(ws, [])

    # 2. waar het verschil uit bestaat
    kop(
        "2. Waar het verschil tussen factuur en order uit bestaat",
        "Bedrag",
        "Regels",
        "Uitleg (elke regel staat in blad Verschillen)",
    )
    r2 = _laatste(ws)
    uitleg = model.uitleg(inst)
    oorzaken = o.per_oorzaak()
    for x in oorzaken.itertuples():
        i = _laatste(ws) + 1
        _rij(
            ws,
            [
                x.oorzaak,
                fx(f"=SUMIFS({ver('G')},{ver('C')},A{i})", x.bedrag),
                fx(f"=COUNTIFS({ver('C')},A{i})", x.regels),
                uitleg[x.oorzaak],
            ],
        )
    eerste, laatste = (r2 + 1, _laatste(ws)) if len(oorzaken) else (r2, r2)
    tot = _laatste(ws) + 1
    _rij(
        ws,
        [
            "Totaal",
            fx(f"=SUM(B{eerste}:B{laatste})", round(oorzaken.bedrag.sum(), 2)),
            fx(f"=SUM(C{eerste}:C{laatste})", int(oorzaken.regels.sum())),
            fx(
                f'=IF(ROUND(B{tot}-C{r1 + 3},2)=0,"Controle: gelijk aan het verschil tussen factuur en order in '
                f'blok 1","LET OP: sluit niet aan op blok 1")',
                "Controle: gelijk aan het verschil tussen factuur en order in blok 1"
                if controles[model.C_OORZAKEN].sluit
                else "LET OP: sluit niet aan op blok 1",
            ),
        ],
    )
    r2t = _laatste(ws)
    _rij(ws, [])

    # 3. wie is aan zet
    kop(
        "3. Wie is aan zet (voorstel)",
        "Bedrag",
        "Regels",
        "Uitleg (kolom Toewijzing in blad Verschillen; daar is de keuze aan te passen)",
    )
    r3 = _laatste(ws)
    wie, saldo = o.per_toewijzing().set_index("toewijzing"), o.per_saldo()
    for t, tekst in model.toewijzing_uitleg(inst).items():
        i = _laatste(ws) + 1
        _rij(
            ws,
            [
                t,
                fx(f"=SUMIFS({ver('H')},{ver('I')},A{i})", wie.bedrag[t]),
                fx(f"=COUNTIFS({ver('I')},A{i})", int(wie.regels[t])),
                tekst,
            ],
        )
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=S_KLEUR[t])
    _rij(
        ws,
        [
            f"Per saldo door {inst.opdrachtgever} nog op te nemen",
            fx(f"=ROUND(B{r3 + 1}-B{r3 + 2},2)", saldo["opnemen"]),
        ],
    )
    _rij(
        ws,
        [
            f"Per saldo door {inst.bureau} nog te factureren",
            fx(f"=ROUND(B{r3 + 3}-B{r3 + 4},2)", saldo["factureren"]),
        ],
    )
    _rij(ws, ["Nog uit te zoeken", fx(f"=B{r3 + 5}", saldo["uitzoeken"])])
    vet += [r2t, _laatste(ws) - 2, _laatste(ws) - 1, _laatste(ws)]
    if o.toezegging is not None:
        _rij(
            ws,
            [
                f"Van '{OW}' staat al op toezegging {o.toezegging_bedragen()['nummers']}",
                fx(
                    f'=SUMIFS({ver("G")},{ver("C")},"{model.O_TOEZEGGING}")',
                    float(oorzaken.set_index("oorzaak").bedrag.get(model.O_TOEZEGGING, 0.0)),
                ),
                None,
                f"Een schriftelijk stuk van {inst.opdrachtgever}, geen order: zie blad Zonder order, blok D.",
            ],
        )
    _rij(ws, [])

    # 4. bank
    if met_bank:
        kop("4. Bank (incl. btw)", "Bedrag", None, "Toelichting")
        r4 = _laatste(ws)
        totalen = o.bank_totalen()
        open_ = [
            f"{k.order} ({model.ORDER_KORT[k.soort]})" for k in o.orders.itertuples() if k.soort not in model.BETAALD
        ]
        verdeling = "Verdeeld over de rekeningen: " + ", ".join(f"{r.aandeel:.0%} {r.naam}" for r in inst.rekeningen)
        _rij(
            ws,
            [
                f"Orders waar {vn} op staat (hele order)",
                fx(f"=SUM({orders(L(kolom.incl))})", totalen["incl"]),
                None,
                "Inclusief de regels van andere medewerkers op dezelfde order",
            ],
        )
        _rij(
            ws,
            [
                "Daarvan ontvangen",
                fx(f"=SUM({orders(L(kolom.totaal))})", totalen["ontvangen"]),
                None,
                f"{verdeling}. Per order in blad Orders, ook waar een creditorder is verrekend of de "
                "betaling alleen uit de boekhouding blijkt",
            ],
        )
        _rij(
            ws,
            [
                "Nog open",
                fx(f"=ROUND(B{r4 + 1}-B{r4 + 2},2)", totalen["open"]),
                None,
                "; ".join(open_) or "Alles ontvangen",
            ],
        )
        vet.append(_laatste(ws))

    # 5. open punten
    if met_opmerkingen:
        _rij(ws, [])
        kop(
            "5. Open punten",
            "Bedrag",
            "Wie is aan zet",
            "Wat er aan de hand is en wat te doen (alle punten met bewijs in blad Opmerkingen)",
        )
        for x in opmerkingen[opmerkingen.belangrijk == "ja"].itertuples():
            naam = f"{x.nr}. {x.onderwerp}" if x.nr else x.onderwerp
            tekst = f"{x.wat} Te doen: {x.actie}" if x.actie else x.wat
            _rij(ws, [naam + (" (incl. btw)" if x.btw == "incl" else ""), x.bedrag, x.wie, tekst])

    # leeswijzer: alleen bladen die er zijn
    if leeswijzer:
        _rij(ws, [])
        kop("Zo lees je dit bestand")
        uit_woorden = (
            f"Welke facturen geen order van {inst.opdrachtgever} hebben, per factuur en per dag, met toelichting."
        )
        bladen = [
            ("Stand", "Dit blad."),
            ("Zonder order", uit_woorden + (" En of er toch iets voor betaald is." if met_bank else "")),
        ]
        if met_kaart:
            bladen.append(
                (
                    "Debiteurenkaart",
                    f"Wat in de boekhouding van {inst.bureau} nog openstaat, per factuur naast wat "
                    f"{inst.opdrachtgever} volgens de bank heeft betaald.",
                )
            )
        bladen += [
            (
                "Per periode",
                f"Eén rij per blok van {PERIODETEKST[inst.blok]}, zoals {inst.opdrachtgever} zijn orders "
                "plaatst. Het snelste overzicht.",
            ),
            (
                "Verschillen",
                "Elk verschil één regel, met oorzaak en wie aan zet is. In de gele kolom kun je zelf een "
                "andere toewijzing kiezen.",
            ),
            (
                "Verborgen bladen",
                "De bladen hieronder zijn verborgen om het bestand overzichtelijk te houden. Terughalen: "
                "rechtsklik op een tabblad onderaan en kies Zichtbaar maken (Unhide).",
            ),
        ]
        if met_opmerkingen:
            bladen.append(("Opmerkingen", "Alles wat bij de controle is opgevallen, met bewijs en wat te doen."))
        bladen += [
            ("Per week", "Urenstaat, factuur en order per week. Ook de weken zonder werk."),
            ("Per dag", "Hetzelfde per dag. Hier komen alle totalen vandaan."),
            ("Orders", "Elke order met wat erop staat en wat ervan betaald is."),
            ("Facturen", "Elke factuur met credit en wat ervan zonder order is."),
        ]
        if met_bank:
            bladen.append(("Bank", f"Elke ontvangst van {inst.opdrachtgever}, met de order waar hij bij hoort."))
        bladen.append(("Urenregels, Orderregels", "De bronregels uit de uren en uit de orders."))
        for naam, tekst in bladen:
            _rij(ws, [naam, None, None, tekst])

    for rij in ws.iter_rows():
        for c in rij:
            c.font = Font(name=LETTER, bold=c.row in koppen + vet or c.row == 1, size=14 if c.row == 1 else 11)
            c.alignment = Alignment(wrap_text=c.column == 4, vertical="top")
            if c.column in (2, 3) and c.row not in koppen:
                c.number_format = "0" if c.column == 3 and r2 < c.row <= r3 + 5 else EUR
            if c.row in koppen:
                c.fill = KOP_VULLING
    ws["A3"].font, ws["A3"].alignment = ROOD, Alignment(wrap_text=False)
    ws.freeze_panes = "A4"
    return ws


def zonder_order(wb: Workbook, o: Overzicht, titel: str = "Zonder order", als_waarden: bool = False) -> Worksheet:
    """Blad Zonder order: wat het bureau factureerde zonder dat het op een order staat. A per factuur, B per dag, C de
    bankboekingen die niet precies bij een order passen (alleen met bank) en D de toezegging regel voor regel
    (alleen met een toezegging). Met `als_waarden` staan er getallen uit het model
    in de cellen in plaats van formules."""
    fx, inst, vn = _fx(als_waarden), o.inst, _eerste(o.medewerker)
    A, B, bank, toez = o.zonder_order, o.zonder_order_dag, o.bank_los, o.toezegging
    met_toez = toez is not None
    ws = wb.create_sheet(titel)
    for k, b in zip("ABCDEFGHIJKLMNO", (15, 12, 13, 26, 22, 8, 12, 11, 13, 12, 46, 12, 50, 40, 90), strict=True):
        ws.column_dimensions[k].width = b
    urenregels = lambda k: _bereik("Urenregels", k, _n(len(o.aansluiting.uren)))  # noqa: E731
    n_fact = _n(len(o.facturen))
    kop_rijen, vet_rijen, eur_cellen = [], [], []
    ra = 6  # kopregel van blok A
    rb = ra + len(A) + 3  # kopregel van blok B
    b1, bn = rb + 1, rb + len(B)
    bcol = lambda k: f"${k}${b1}:${k}${bn}"  # noqa: E731
    _rij(ws, [f"Gefactureerd door {inst.bureau} zonder order van {inst.opdrachtgever} - {o.medewerker}"])
    _rij(ws, [f"Deze facturen (of delen ervan) staan op geen enkele order van {inst.opdrachtgever}. Bedragen ex btw."])
    if bank is None:
        _rij(ws, ["Er is geen bank aangeleverd: of hier iets voor is ontvangen, is niet te zien."])
    else:
        _rij(
            ws,
            [
                f"De bank is per order gekoppeld. Voor deze dagen is er geen order, dus is er tot en met "
                f"{o.aansluiting.bank_tot:%d-%m-%Y} niets aan gekoppeld; zie blok C onderaan."
            ],
        )
    _rij(
        ws,
        [
            "Blok A: per factuur, met toelichting. Blok B: de specificatie per dag."
            + (" Blok C: de bankboekingen die niet precies bij een order passen." if bank is not None else "")
            + (" Blok D: de toezegging regel voor regel." if met_toez else "")
        ],
    )
    if met_toez:
        b = o.toezegging_bedragen()
        _rij(
            ws,
            [
                f"Let op: voor {nl_bedrag(b['op_toezegging'])} hiervan staat een toezegging van {inst.opdrachtgever} "
                f"({b['nummers']}; kolom P en blok D). De rest, {nl_bedrag(b['rest'])}, staat er niet op (kolom Q): "
                f"{nl_bedrag(b['uren_rest'])} aan uren zonder enig stuk van {inst.opdrachtgever} en "
                f"{nl_bedrag(b['onkosten'])} aan onkosten."
            ],
        )
    else:
        _rij(ws, [])
    koppen_a = [
        "A. Per factuur",
        "Factuurdatum",
        "Week",
        "Dagen zonder order",
        "Urenstaat",
        "Uren",
        "Bedrag uren",
        "Onkosten",
        "Zonder order",
        "Factuur totaal",
        f"Order van {inst.opdrachtgever}",
        f"Ontvangen van {inst.opdrachtgever}",
        f"Ander schriftelijk stuk van {inst.opdrachtgever}",
        f"Kilometers die {vn} die week bij {inst.opdrachtgever} boekte",
        "Toelichting",
    ]
    if met_toez:
        koppen_a += ["Waarvan uren op de toezegging", "Niet op de toezegging: uren zonder enig stuk, en onkosten"]
    _rij(ws, koppen_a)
    ws.column_dimensions["P"].width = ws.column_dimensions["Q"].width = 15
    kop_rijen.append(_laatste(ws))
    kolommen_a = "FGHIJL" + ("PQ" if met_toez else "")
    for x in A.itertuples():
        i = _laatste(ws) + 1
        rij = [
            x.factuur,
            x.factuurdatum,
            x.weken,
            x.dagen,
            x.urenstaten,
            fx(f"=SUMPRODUCT(({bcol('A')}=A{i})*{bcol('F')})", x.uren),
            fx(f"=SUMPRODUCT(({bcol('A')}=A{i})*{bcol('G')})", x.bedrag_uren),
            fx(f"=SUMPRODUCT(({bcol('A')}=A{i})*{bcol('H')})", x.onkosten),
            fx(f"=G{i}+H{i}", x.zonder_order),
            fx(f"=SUMPRODUCT(({urenregels('H')}=A{i})*{urenregels('M')})", x.factuur_totaal),
            x.order_bij,
            x.ontvangen,
            x.ander_stuk,
            x.km_bij_opdrachtgever,
            x.toelichting,
        ]
        if met_toez:
            rij += [
                fx(f"=SUMPRODUCT(({bcol('A')}=A{i})*{bcol('N')})", x.op_toezegging),
                fx(f"=ROUND(I{i}-P{i},2)", x.niet_op_toezegging),
            ]
        _rij(ws, rij)
        eur_cellen += [(i, k) for k in kolommen_a]
    i = _laatste(ws) + 1

    def som(kolom: str, veld: str):
        """De som van kolom `kolom` van blok A: een formule, of de waarde uit het model."""
        return round(A[veld].sum(), 2) if als_waarden else _som(kolom, ra + 1, i - 1)

    totaal = ["Totaal", None, None, None, None]
    totaal += [som(k, c) for k, c in zip("FGHI", ("uren", "bedrag_uren", "onkosten", "zonder_order"), strict=True)]
    totaal += [
        None,
        fx(
            f"=IF(ROUND(I{i}-SUM({_bereik('Facturen', 'I', n_fact)}),2)=0,"
            '"Controle: gelijk aan kolom I in blad Facturen","LET OP: sluit niet aan op blad Facturen")',
            "Controle: gelijk aan kolom I in blad Facturen"
            if o.controles()[model.C_ZONDER_ORDER].sluit
            else "LET OP: sluit niet aan op blad Facturen",
        ),
        som("L", "ontvangen"),
    ]
    if met_toez:
        totaal += [None, None, None, som("P", "op_toezegging"), som("Q", "niet_op_toezegging")]
    _rij(ws, totaal)
    vet_rijen.append(i)
    eur_cellen += [(i, k) for k in "FGHIL" + ("PQ" if met_toez else "")]
    _rij(ws, [])

    # blok B: per dag
    koppen_b = [
        "B. Specificatie per dag",
        "Datum",
        "Week",
        "Dag",
        "Urenstaat",
        "Uren",
        "Bedrag uren",
        "Onkosten",
        "Totaal",
        "Uurtarief",
        "Soort onkosten",
        f"Km bij {inst.opdrachtgever} geboekt",
        "Rij in de uren",
    ]
    if met_toez:
        koppen_b.append("Waarvan uren op de toezegging")
    _rij(ws, koppen_b)
    kop_rijen.append(_laatste(ws))
    assert _laatste(ws) == rb, "blok B staat niet waar de formules van blok A naar verwijzen"
    for x in B.itertuples():
        i = _laatste(ws) + 1
        rij = [
            x.factuur,
            x.datum,
            x.week,
            DAGNAAM[x.datum.weekday()],
            x.urenstaten,
            x.uren,
            x.bedrag_uren,
            x.onkosten,
            fx(f"=G{i}+H{i}", x.totaal),
            x.uurtarief,
            x.soort_onkosten,
            x.km_bij_opdrachtgever,
            x.rij,
        ]
        if met_toez:
            rij.append(x.op_toezegging)
        _rij(ws, rij)
        eur_cellen += [(i, k) for k in "FGHIJ" + ("N" if met_toez else "")]
    i = _laatste(ws) + 1
    _rij(
        ws,
        ["Totaal", None, None, None, None]
        + [
            round(B[c].sum(), 2) if als_waarden else _som(k, b1, bn)
            for k, c in zip("FGHI", ("uren", "bedrag_uren", "onkosten", "totaal"), strict=True)
        ],
    )
    vet_rijen.append(i)
    eur_cellen += [(i, k) for k in "FGHI"]

    # blok C: is er toch iets voor ontvangen?
    if bank is not None:
        _rij(ws, [])
        _rij(ws, ["C. Is er toch iets voor ontvangen?"])
        kop_rijen.append(_laatste(ws))
        eerste_dag, telling = o.aansluiting.uren.datum.min(), o.bank_telling()
        slot = (
            "Er zijn geen boekingen die niet precies bij een order passen."
            if bank.empty
            else "Hieronder de ene boeking die niet precies bij een order past."
            if len(bank) == 1
            else f"Hieronder de {len(bank)} boekingen die niet precies bij een order passen."
        )
        _rij(
            ws,
            [
                f"In blad Bank staan {telling['totaal']} ontvangsten van {inst.opdrachtgever} vanaf "
                f"{eerste_dag:%d-%m-%Y}: {telling['bij_order']} horen bij een order en {telling['zonder_order']} niet. "
                f"{slot}"
            ],
        )
        _rij(ws, ["Zie per regel de opmerking."] if len(bank) else [None])
        _rij(
            ws,
            [
                "Datum",
                "Rekening",
                None,
                "Omschrijving bank",
                None,
                None,
                "Bedrag",
                None,
                None,
                None,
                "Gekoppeld aan order",
                None,
                f"Betreft {vn}",
                None,
                "Opmerking",
            ],
        )
        kop_rijen.append(_laatste(ws))
        for r in bank.itertuples():
            _rij(
                ws,
                [
                    r.datum,
                    r.rekening,
                    None,
                    r.omschrijving,
                    None,
                    None,
                    r.bedrag,
                    None,
                    None,
                    None,
                    r.orders,
                    None,
                    r.betreft,
                    None,
                    r.opmerking,
                ],
            )
            eur_cellen.append((_laatste(ws), "G"))

    # blok D: de toezegging regel voor regel
    if met_toez:
        for nr, t in toez.groupby("toezegging", sort=False):
            datum = t.datum.iloc[0]
            _rij(ws, [])
            _rij(
                ws,
                [
                    f"D. Toezegging {nr}" + (f" van {datum:%d-%m-%Y}" if pd.notna(datum) else ""),
                    "Week of maand",
                    "Soort",
                    "Omschrijving op de toezegging",
                    "Eenheid",
                    "Aantal",
                    "Prijs",
                    None,
                    "Bedrag",
                    None,
                    "Ook op order",
                    "Uren daar opgenomen",
                    f"Door {inst.bureau} gefactureerd over deze week of maand",
                    None,
                    "Wat dit betekent",
                ],
            )
            kop_rijen.append(_laatste(ws))
            p1 = _laatste(ws) + 1
            for x in t.itertuples():
                _rij(
                    ws,
                    [
                        nr,
                        x.week_of_maand,
                        x.soort,
                        x.omschrijving,
                        x.eenheid,
                        x.aantal,
                        x.prijs,
                        None,
                        x.bedrag,
                        None,
                        x.op_order,
                        x.uren_op_order,
                        x.gefactureerd,
                        None,
                        x.betekenis,
                    ],
                )
                eur_cellen += [(_laatste(ws), k) for k in "FGIL"]
            pn = _laatste(ws)
            _rij(
                ws,
                [
                    "Totaal toezegging",
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    fx(f"=SUM(I{p1}:I{pn})", round(t.bedrag.sum(), 2)),
                    None,
                    None,
                    None,
                    "Ex btw. Een toezegging is geen order: er is geen betaling aan gekoppeld.",
                ],
            )
            vet_rijen.append(_laatste(ws))
            eur_cellen.append((_laatste(ws), "I"))
            for s in (model.T_ALLEEN, model.T_DUBBEL, model.T_REIS):
                _rij(
                    ws,
                    [
                        f"Waarvan: {s[0].lower()}{s[1:]}",
                        None,
                        None,
                        None,
                        None,
                        None,
                        None,
                        None,
                        fx(f'=SUMIFS(I{p1}:I{pn},C{p1}:C{pn},"{s}")', round(t.bedrag[t.soort == s].sum(), 2)),
                    ],
                )
                eur_cellen.append((_laatste(ws), "I"))

    for rij in ws.iter_rows():
        for c in rij:
            c.font = Font(name=LETTER, bold=c.row in kop_rijen + vet_rijen or c.row == 1, size=14 if c.row == 1 else 11)
            if c.row >= ra:
                c.alignment = Alignment(wrap_text=c.column > 1, vertical="top")  # kolom A niet: daar staan losse zinnen
            if c.row in kop_rijen:
                c.fill = KOP_VULLING
            if c.column_letter == "B" or (c.column_letter == "A" and c.row > bn + 3):
                c.number_format = DATUM
    for i, k in eur_cellen:
        ws[f"{k}{i}"].number_format = "#,##0.00" if k == "L" else EUR  # ontvangen: een nul moet als 0,00 te lezen zijn
    if met_toez:
        ws["A5"].font, ws["A5"].alignment = VET, Alignment(wrap_text=False)
    if bank is None:
        ws["A3"].font, ws["A3"].alignment = ROOD, Alignment(wrap_text=False)
    ws.freeze_panes = "B5"
    return ws


def _debiteurenkaart(wb: Workbook, o: Overzicht) -> Worksheet:
    """Blad Debiteurenkaart: A de hele kaart, B wat er voor deze medewerker echt openstaat na afletteren (met een
    controleregel), C elke factuur naast wat ervoor is betaald en D de ontvangsten op de kaart."""
    k, inst, vn = o.kaart, o.inst, _eerste(o.medewerker)
    ws = wb.create_sheet("Debiteurenkaart")
    for kol, b in zip("ABCDEFGHIJ", (17, 12, 44, 14, 32, 14, 14, 14, 14, 100), strict=True):
        ws.column_dimensions[kol].width = b
    kop_rijen, vet_rijen, eur_cellen = [], [], []
    fk, ko = k.facturen, k.ontvangsten
    nb = len(k.regels)
    n_eraf = len(k.eraf) + 1 if k.eraf else 0

    _rij(ws, [f"Debiteurenkaart {inst.opdrachtgever} naast de bank en de orders - {o.medewerker}"])
    _rij(
        ws,
        [
            f"Bron: de debiteurenkaart. Bedragen incl. btw. De kaart toont facturen die in de boekhouding van "
            f"{inst.bureau} nog openstaan, en ontvangsten die nog niet aan een factuur zijn gekoppeld "
            "(afgeletterd)."
        ],
    )
    _rij(
        ws,
        [
            "Let op: niet-afgeletterde ontvangsten staan op de kaart als negatief bedrag. Wie alleen de open "
            "facturen optelt, ziet te veel openstaan."
        ],
    )
    _rij(ws, [])
    _rij(ws, ["A. De hele kaart (alle medewerkers)", "Aantal", None, "Bedrag"])
    kop_rijen.append(_laatste(ws))
    _rij(ws, ["Open facturen", k.aantallen["open"], None, k.heel["open"]])
    _rij(ws, ["Ontvangsten, niet afgeletterd", k.aantallen["ontvangsten"], None, k.heel["ontvangsten"]])
    _rij(ws, ["Saldo van de kaart", None, None, F(f"=D{_laatste(ws) - 1}+D{_laatste(ws)}")])
    vet_rijen.append(_laatste(ws))
    for naam in k.volgorde:  # zoals ze voor het eerst op de kaart staan
        if naam:
            _rij(ws, [f"Open facturen {naam}", int(k.per_medewerker_aantal[naam]), None, k.per_medewerker[naam]])
        else:
            _rij(
                ws,
                [
                    "Open facturen die niet in de uren staan",
                    k.aantallen["buiten_uren"],
                    ", ".join(k.buiten_nummers),
                    k.heel["buiten_uren"],
                ],
            )
    eur_cellen += [(i, "D") for i in range(6, _laatste(ws) + 1)]
    _rij(ws, [])
    rb_, n_f, n_o = _laatste(ws) + 1, len(fk), len(ko)
    rc = rb_ + nb + 5 + n_eraf  # kopregel van blok C
    c1, cn = rc + 1, rc + n_f
    rd = cn + 4  # kopregel van blok D
    d1, dn = rd + 1, rd + n_o

    _rij(ws, [f"B. {vn}: wat staat er echt open", None, None, "Bedrag"])
    kop_rijen.append(_laatste(ws))
    for sleutel, (label, bedrag, uitleg) in zip(k.regel_sleutels, k.regels, strict=True):
        if sleutel == model.R_OPEN:
            bedrag = _som("H", c1, cn)
        elif sleutel == model.R_AF_ONTVANGEN and n_o:
            bedrag = F(
                f"=-SUM(F{d1}:F{dn})"
                + (f"-{k.bij_credit:.2f}" if k.bij_credit else "")
                + (f"+{k.al_afgeletterd:.2f}" if k.al_afgeletterd else "")
            )
        _rij(ws, [label, None, None, bedrag] + ([None] * 5 + [uitleg] if uitleg else []))
    _rij(ws, ["Echt open na afletteren", None, None, F(f"=SUM(D{rb_ + 1}:D{rb_ + nb})")])
    vet_rijen.append(_laatste(ws))
    # FIXED (zonder duizendtallen) in plaats van TEXT(x,"0.00"): de opmaakcode van TEXT hangt af van de taal van Excel
    verschil = F(
        f'=IF(ABS(D{rb_ + nb + 1}-D{rb_ + nb + 2})<{model.CONTROLE_VERSCHIL},"Sluit aan",'
        f'"LET OP: sluit niet aan, verschil "&FIXED(D{rb_ + nb + 1}-D{rb_ + nb + 2},2,TRUE))'
    )
    _rij(
        ws,
        [
            "Ter controle: echt open volgens bank en orders (blok C, facturen op de kaart)"
            + (" plus de teruggehaalde creditorder" if k.bij_credit else ""),
            None,
            None,
            F(f'=SUMIF(H{c1}:H{cn},">0",G{c1}:G{cn})' + (f"+{k.bij_credit:.2f}" if k.bij_credit else "")),
            None,
            None,
            None,
            None,
            None,
            verschil,
        ],
    )
    for label, bedrag, uitleg in k.eraf:
        _rij(ws, [label, None, None, -bedrag, None, None, None, None, None, uitleg])
    if k.eraf:
        _rij(
            ws,
            [
                f"Per saldo te vorderen op {inst.opdrachtgever}",
                None,
                None,
                F(f"=D{rb_ + nb + 1}+SUM(D{rb_ + nb + 3}:D{rb_ + nb + 2 + len(k.eraf)})"),
            ],
        )
        vet_rijen.append(_laatste(ws))
    eur_cellen += [(i, "D") for i in range(rb_ + 1, rb_ + nb + 3 + n_eraf)]
    _rij(ws, [])
    _rij(ws, [])

    _rij(
        ws,
        [
            f"C. Elke factuur van {vn}",
            "Factuurdatum",
            "Week(en)",
            "Factuur incl. btw",
            "Order(s)",
            f"Ontvangen van {inst.opdrachtgever}",
            "Echt open",
            "Open op de kaart",
            "Te veel open op de kaart",
            "Oordeel",
        ],
    )
    kop_rijen.append(_laatste(ws))
    assert _laatste(ws) == rc, "blok C staat niet waar de formules naar verwijzen"
    for x in fk.itertuples():
        i = _laatste(ws) + 1
        _rij(
            ws,
            [
                x.factuur,
                x.factuurdatum,
                x.weken,
                x.incl,
                x.orders,
                x.ontvangen,
                F(f"=ROUND(D{i}-F{i},2)"),
                x.kaart,
                F(f"=ROUND(H{i}-G{i},2)"),
                x.oordeel,
            ],
        )
        eur_cellen += [(i, c) for c in "DFGHI"]
    i = _laatste(ws) + 1
    _rij(ws, ["Totaal", None, None] + [_som(c, c1, cn) if c != "E" else None for c in "DEFGHI"])
    vet_rijen.append(i)
    eur_cellen += [(i, c) for c in "DFGHI"]
    _rij(ws, [])
    _rij(ws, [])

    _rij(
        ws,
        [
            "D. Ontvangsten op de kaart",
            "Rekening",
            "Omschrijving",
            "Bedrag",
            "Order(s) volgens de bank",
            f"Waarvan voor {vn}",
            None,
            None,
            None,
            "Opmerking",
        ],
    )
    kop_rijen.append(_laatste(ws))
    assert _laatste(ws) == rd, "blok D staat niet waar de formules naar verwijzen"
    for x in ko.itertuples():
        _rij(ws, [x.datum, x.rekening, x.omschrijving, x.bedrag, x.orders, x.waarvan, None, None, None, x.opmerking])
        eur_cellen += [(_laatste(ws), c) for c in "DF"]
    i = _laatste(ws) + 1
    _rij(ws, ["Totaal", None, None, _som("D", d1, dn), None, _som("F", d1, dn)])
    vet_rijen.append(i)
    eur_cellen += [(i, c) for c in "DF"]

    for rij in ws.iter_rows():
        for c in rij:
            c.font = Font(name=LETTER, bold=c.row in kop_rijen + vet_rijen or c.row == 1, size=14 if c.row == 1 else 11)
            if c.row >= rc:
                c.alignment = Alignment(wrap_text=c.column > 1, vertical="top")
            if c.row in kop_rijen:
                c.fill = KOP_VULLING
            if c.column_letter == "B" and c.row > rc or c.column_letter == "A" and c.row > rd:
                c.number_format = DATUM
    for i, c in eur_cellen:
        ws[f"{c}{i}"].number_format = "#,##0.00"
    ws["A3"].font = ROOD  # de waarschuwing over niet-afgeletterde ontvangsten (rij 3), niet de lege rij eronder
    ws.freeze_panes = "B5"
    return ws


def _per_periode(wb: Workbook, o: Overzicht) -> None:
    p, n_week, n_per = o.perioden, _n(len(o.weken)), len(o.perioden) + 1
    week = lambda k: _bereik("Per week", k, n_week)  # noqa: E731
    ws = _blad(
        wb,
        "Per periode",
        [
            PERIODEKOP[o.inst.blok],
            "Van",
            "T/m",
            "Uren gewerkt",
            "Urenstaat (goedgekeurd)",
            "Gefactureerd",
            "Verschil urenstaat - factuur",
            "Op order",
            "Verschil factuur - order",
            "Order(s)",
            "Betaling",
            "Status",
            "Waar het verschil uit bestaat",
        ],
        [18, 12, 12, 9, 13, 13, 12, 13, 13, 34, 44, 20, 90],
    )
    for i, x in enumerate(p.itertuples(), start=2):
        crit = f"{week('B')},A{i}"
        _rij(
            ws,
            [
                x.periode,
                x.van,
                x.tm,
                F(f"=SUMIFS({week('E')},{crit})"),
                F(f"=SUMIFS({week('G')},{crit})"),
                F(f"=SUMIFS({week('J')},{crit})"),
                F(f"=ROUND(E{i}-F{i},2)"),
                F(f"=SUMIFS({week('M')},{crit})"),
                F(f"=ROUND(F{i}-H{i},2)"),
                x.orders,
                x.betaling,
                x.status,
                x.oorzaken,
            ],
        )
        ws.cell(i, 12).fill = PatternFill("solid", fgColor=PERIODE_KLEUR[x.status])
    _opmaak(ws, eur="DEFGHI", datum="BC")
    _totaal(ws, n_per, "DEFGHI", "DEFGHI")
    _omloop(ws, "JKM")


def _verschillen(wb: Workbook, o: Overzicht) -> None:
    v = o.verschillen
    n_v = len(v) + 1
    ws = _blad(
        wb,
        "Verschillen",
        [
            PERIODEKOP[o.inst.blok],
            "Week",
            "Oorzaak",
            "Dagen",
            "Gefactureerd",
            "Op order",
            "Verschil",
            "Bedrag",
            "Toewijzing",
            "Toelichting",
            "Factuur",
            "Order",
            "Voorstel",
            "Eigen keuze (alleen invullen als je het anders ziet)",
        ],
        [18, 13, 46, 34, 13, 13, 12, 12, 30, 80, 30, 22, 30, 30],
    )
    for i, x in enumerate(v.itertuples(), start=2):
        _rij(
            ws,
            [
                x.periode,
                x.week,
                x.oorzaak,
                x.dagen,
                x.fact,
                x.order,
                F(f"=ROUND(E{i}-F{i},2)"),
                F(f"=ABS(G{i})"),
                F(f'=IF(N{i}="",M{i},N{i})'),
                x.toelichting,
                x.facturen,
                x.orders,
                x.voorstel,
                x.eigen or None,
            ],
        )
        ws.cell(i, 9).fill = PatternFill("solid", fgColor=S_KLEUR.get(x.eigen or x.voorstel, "FFFFFF"))
        ws.cell(i, 14).fill = GEEL
    keuze = DataValidation(type="list", formula1='"' + ",".join(TOEWIJZINGEN) + '"', allow_blank=True)
    ws.add_data_validation(keuze)
    keuze.add(f"N2:N{max(n_v, 2)}")
    _opmaak(ws, eur="EFGH")
    _totaal(ws, n_v, "EFG", "EFG")
    _omloop(ws, "DJ")


def _opmerkingen(wb: Workbook, opmerkingen: pd.DataFrame) -> None:
    ws = _blad(
        wb,
        "Opmerkingen",
        [
            "Nr",
            "Onderwerp",
            "Bedrag",
            "Ex of incl. btw",
            "Wat er aan de hand is",
            "Bewijs",
            "Wat te doen",
            "Wie is aan zet",
        ],
        [5, 40, 12, 9, 80, 60, 60, 12],
    )
    for x in opmerkingen.itertuples():
        _rij(ws, [x.nr, x.onderwerp, x.bedrag, x.btw, x.wat, x.bewijs, x.actie, x.wie])
    _opmaak(ws, eur="C")
    _omloop(ws, "BEFG")


def _per_week(wb: Workbook, o: Overzicht) -> None:
    n_dag = _n(len(o.per_dag))
    dag = lambda k: _bereik("Per dag", k, n_dag)  # noqa: E731
    ws = _blad(
        wb,
        "Per week",
        [
            "Week",
            PERIODEKOP[o.inst.blok],
            "Van",
            "T/m",
            "Uren gewerkt",
            "Urenstaat/urenstaten",
            "Urenstaat (goedgekeurd)",
            "Factuur",
            "Factuurdatum",
            "Gefactureerd",
            "Verschil urenstaat - factuur",
            "Order(s)",
            "Op order",
            "Verschil factuur - order",
            "Status",
            "Opmerking",
        ],
        [13, 18, 12, 12, 9, 24, 13, 26, 22, 13, 12, 24, 13, 13, 32, 50],
    )
    for i, x in enumerate(o.weken.itertuples(), start=2):
        crit = f"{dag('B')},A{i}"
        _rij(
            ws,
            [
                x.week,
                x.periode,
                x.van,
                x.tm,
                x.uren,
                x.urenstaten,
                x.urenstaat_bedrag,
                x.facturen,
                x.factuurdatums,
                F(f"=SUMIFS({dag('K')},{crit})"),
                F(f"=ROUND(G{i}-J{i},2)"),
                x.orders,
                F(f"=SUMIFS({dag('N')},{crit})"),
                F(f"=ROUND(J{i}-M{i},2)"),
                x.status,
                x.opmerking,
            ],
        )
        ws.cell(i, 15).fill = PatternFill("solid", fgColor=WEEK_KLEUR[x.status])
    n = len(o.weken) + 1
    _opmaak(ws, eur="EGJKMN", datum="CD")
    _totaal(ws, n, "EGJKMN", "EGJKMN")


def _per_dag(wb: Workbook, o: Overzicht) -> None:
    ws = _blad(
        wb,
        "Per dag",
        [
            "Datum",
            "Week",
            PERIODEKOP[o.inst.blok],
            "Order",
            "Urenstaat/urenstaten",
            "Factuur",
            "Uren gefactureerd",
            "Uren op order",
            "Gefactureerd uren",
            "Gefactureerd onkosten",
            "Gefactureerd totaal",
            "Order uren",
            "Order onkosten",
            "Order totaal",
            "Verschil",
            "Status",
            "Toelichting",
        ],
        [12, 13, 18, 13, 22, 24, 9, 9, 12, 11, 12, 12, 11, 12, 12, 42, 60],
    )
    for i, r in enumerate(o.per_dag.itertuples(), start=2):
        _rij(
            ws,
            [
                r.datum,
                r.week,
                r.periode,
                r.order,
                r.urenstaten,
                r.facturen,
                r.fact_uren,
                r.order_uren,
                r.fact_bedrag_uren,
                r.fact_onkosten,
                F(f"=I{i}+J{i}"),
                r.order_bedrag_uren,
                r.order_onkosten,
                F(f"=L{i}+M{i}"),
                F(f"=ROUND(K{i}-N{i},2)"),
                r.oordeel,
                r.toelichting,
            ],
        )
        ws.cell(i, 16).fill = PatternFill("solid", fgColor=DAG_KLEUR[r.status])
    _opmaak(ws, eur="GHIJKLMNO", datum="A")
    _totaal(ws, len(o.per_dag) + 1, "GHIJKLMNO", "GHIJKLMNO")


def _orders(wb: Workbook, o: Overzicht) -> None:
    inst, vn = o.inst, _eerste(o.medewerker)
    n_dag = _n(len(o.per_dag))
    dag = lambda k: _bereik("Per dag", k, n_dag)  # noqa: E731
    rek, kol = inst.rekeningen, _orderkolommen(inst)
    k0, totaal_k, open_k, datum_k, betaling_k = kol.eerste_rekening, kol.totaal, kol.open, kol.datum, kol.betaling
    ex_k, fact_k, versch_k, bank_k = kol.ex, kol.gefactureerd, kol.verschil, kol.bank
    ws = _blad(
        wb,
        "Orders",
        [
            "Order",
            "Entiteit",
            "Factuurdatum",
            f"Regels {vn} van",
            "T/m",
            "Order ex btw (totaal)",
            f"Waarvan {o.medewerker}",
            "Anderen op deze order",
            "Order incl. btw",
            *[f"Ontvangen {r.naam} ({r.aandeel:.0%})" for r in rek],
            "Ontvangen totaal",
            "Nog open",
            "Datum ontvangst",
            "Betaling",
            f"Deel {vn} betaald, ex btw",
            "Gefactureerd op deze dagen",
            "Verschil factuur - order",
            "Omschrijving bank",
        ],
        [13, 8, 12, 12, 12, 13, 13, 13, 13, *[13] * len(rek), 13, 12, 12, 46, 13, 13, 13, 70],
    )
    for i, x in enumerate(o.orders.itertuples(), start=2):
        r = o.orders.iloc[i - 2]
        ontvangen = [r[f"ontvangen {z.naam}"] for z in rek]
        _rij(
            ws,
            [
                x.order,
                x.entiteit,
                x.factuurdatum,
                x.van,
                x.tm,
                x.excl,
                F(f"=SUMIFS({dag('N')},{dag('D')},A{i})"),
                F(f"=ROUND(F{i}-G{i},2)"),
                x.incl,
                *ontvangen,
                F(f"=SUM({L(k0)}{i}:{L(totaal_k - 1)}{i})"),
                F(f"=ROUND(I{i}-{L(totaal_k)}{i},2)"),
                x.betaald_op,
                x.oordeel,
                F(f"=IF(I{i}=0,0,ROUND(G{i}*{L(totaal_k)}{i}/I{i},2))"),
                F(f"=SUMIFS({dag('K')},{dag('D')},A{i})"),
                F(f"=ROUND({L(fact_k)}{i}-G{i},2)"),
                x.bank_omschrijving,
            ],
        )
        kleur = (
            "E2EFDA"
            if x.soort in model.BETAALD
            else "FFF2CC"
            if x.soort == model.ORDER_NOG_NIET
            else "D9D9D9"
            if x.soort == model.ORDER_GEEN_BANK
            else "F8CBAD"
        )
        ws.cell(i, betaling_k).fill = PatternFill("solid", fgColor=kleur)
    n_io = len(o.orders) + 1
    eur_kolommen = [6, 7, 8, 9, *range(k0, open_k + 1), ex_k, fact_k, versch_k]
    _opmaak(ws, eur="".join(L(c) for c in eur_kolommen), datum="CDE" + L(datum_k))
    _totaal(ws, n_io, "".join(L(c) for c in eur_kolommen), "".join(L(c) for c in eur_kolommen))
    verv = o.vervallen_orders()
    if len(verv):  # onder het totaal: orders die niet meetellen omdat ze door een andere order zijn vervangen
        _rij(ws, [])
        kop = [None] * bank_k
        kop[0], kop[1], kop[2], kop[5], kop[6], kop[8] = (
            "Vervallen, telt niet mee",
            "Entiteit",
            "Factuurdatum",
            "Order ex btw (totaal)",
            f"Waarvan {o.medewerker}",
            "Order incl. btw",
        )
        kop[betaling_k - 1], kop[bank_k - 1] = "Vervangen door", "Toelichting"
        _rij(ws, kop)
        for c in ws[_laatste(ws)]:
            c.font, c.fill = VET, KOP_VULLING
        for x in verv.itertuples():
            rij = [None] * bank_k
            rij[0], rij[1], rij[2], rij[5], rij[6], rij[8] = (
                x.order,
                x.entiteit,
                x.factuurdatum,
                x.excl,
                x.waarvan_medewerker,
                x.incl,
            )
            rij[betaling_k - 1], rij[bank_k - 1] = x.vervangen_door, x.toelichting
            _rij(ws, rij)
            for c in ws[_laatste(ws)]:
                c.font = Font(name=LETTER)
                c.number_format = (
                    DATUM if c.column_letter == "C" else EUR if c.column_letter in "FGI" else c.number_format
                )


def _facturen(wb: Workbook, o: Overzicht) -> None:
    n_uren = _n(len(o.aansluiting.uren))
    uren = lambda k: _bereik("Urenregels", k, n_uren)  # noqa: E731
    ws = _blad(
        wb,
        "Facturen",
        [
            "Factuur",
            "Factuurdatum",
            "Week(en)",
            "Urenstaat/urenstaten",
            "Omzet in de uren",
            "Credit in de uren",
            "Creditfactuur",
            "Netto (telt mee)",
            "Waarvan zonder order",
            "Dagen na de laatste werkdag",
            "Opmerking",
        ],
        [15, 12, 34, 40, 13, 13, 15, 13, 13, 10, 70],
    )
    for i, x in enumerate(o.facturen.itertuples(), start=2):
        # SUMPRODUCT en geen SUMIFS: een factuurnummer met cijfers en een streepje wordt als criterium niet
        # betrouwbaar als tekst gelezen
        _rij(
            ws,
            [
                x.factuur,
                x.factuurdatum,
                x.weken,
                x.urenstaten,
                F(f"=SUMPRODUCT(({uren('H')}=A{i})*{uren('J')})"),
                F(f"=SUMPRODUCT(({uren('H')}=A{i})*{uren('L')})"),
                x.creditfacturen,
                F(f"=E{i}+F{i}"),
                F(f'=SUMPRODUCT(({uren("H")}=A{i})*({uren("Q")}="{GEEN}")*{uren("M")})'),
                x.dagen_na_werkdag,
                x.opmerking,
            ],
        )
    n = len(o.facturen) + 1
    _opmaak(ws, eur="EFHI", datum="B")
    _totaal(ws, n, "EFHI", "EFHI")


def _bank(wb: Workbook, o: Overzicht) -> None:
    vn, btw = _eerste(o.medewerker), o.inst.btw
    ws = _blad(
        wb,
        "Bank",
        [
            "Datum",
            "Rekening",
            "Aandeel",
            "Bedrag",
            "Omschrijving",
            "Order",
            "Medewerker(s) op die order",
            f"Betreft {vn}",
            "Hele betaling incl. btw (bedrag / aandeel)",
            f"Ex btw bij {percentage(btw)}",
            "Opmerking",
        ],
        [12, 16, 8, 12, 70, 34, 30, 16, 14, 14, 110],
    )
    for i, r in enumerate(o.bank.itertuples(), start=2):
        _rij(
            ws,
            [
                r.datum,
                r.rekening,
                r.aandeel,
                r.bedrag,
                r.omschrijving,
                ", ".join(r.orders),
                r.medewerkers,
                r.betreft,
                F(f"=ROUND(D{i}/C{i},2)"),
                F(f"=ROUND(I{i}/{round(1 + btw, 6)},2)"),
                r.opmerking,
            ],
        )
    _opmaak(ws, eur="DIJ", datum="A")
    for c in ws["C"][1:]:
        c.number_format = "0%"


def _bronbladen(wb: Workbook, o: Overzicht) -> None:
    a = o.aansluiting
    kolommen = (
        "rij urenstaat datum uren soort tarief status factuurnummer factuurdatum bedrag "
        "creditnummer creditbedrag netto factuurbedrag post telt_mee order"
    ).split()
    ws = _blad(
        wb,
        "Urenregels",
        [
            "Rij in bron",
            "Urenstaat",
            "Datum",
            "Uren",
            "Tariefregel",
            "Tarief",
            "Status urenstaat",
            "Factuur",
            "Factuurdatum",
            "Bedrag",
            "Creditfactuur",
            "Creditbedrag",
            "Netto",
            "Factuurbedrag",
            "Soort",
            "Telt mee",
            "Toegewezen order",
            "Week",
        ],
        [9, 12, 12, 7, 34, 9, 13, 14, 12, 10, 14, 10, 10, 12, 10, 8, 14, 13],
    )
    u = a.uren.sort_values(["datum", "rij"])
    weken = weeksleutels(u.datum, a.inst)
    for r, week in zip(u[kolommen].itertuples(index=False), weken, strict=True):
        _rij(ws, [("ja" if v else "nee") if isinstance(v, bool) else v for v in r] + [week])
    _opmaak(ws, eur="DFJLMN", datum="CI")
    ws = _blad(
        wb,
        "Orderregels",
        [
            "Order",
            "Entiteit",
            "Medewerker",
            "Project",
            "Datum",
            "Omschrijving",
            "Eenheid",
            "Aantal",
            "Tarief",
            "Bedrag",
            "Soort",
            f"Telt mee ({o.medewerker})",
            "Bestand",
        ],
        [13, 8, 18, 34, 12, 34, 11, 9, 8, 10, 10, 10, 46],
    )
    for r in a.order_regels.sort_values(["order", "datum"]).itertuples():
        _rij(
            ws,
            [
                r.order,
                r.entiteit,
                r.medewerker,
                r.project,
                r.datum,
                r.omschrijving,
                r.eenheid,
                r.aantal,
                r.tarief,
                r.bedrag,
                r.post,
                "ja" if r.telt_mee else "nee",
                r.bestand,
            ],
        )
    _opmaak(ws, eur="HIJ", datum="E")


def schrijf(o: Overzicht, opmerkingen: pd.DataFrame, pad: str | Path) -> None:
    """Schrijft het werkboek van één medewerker. `opmerkingen` zijn de opmerkingen van deze medewerker (de kolommen van
    `inlezen.lees_opmerkingen`); de belangrijke komen ook op blad Stand. De keuze die de gebruiker in een eerder
    werkboek maakte, zit al in `o` (zie `eigen_keuze`)."""
    pad = Path(pad)
    wb = Workbook()
    stand(wb, o, opmerkingen)
    zonder_order(wb, o)
    if o.kaart is not None:
        _debiteurenkaart(wb, o)
    _per_periode(wb, o)
    _verschillen(wb, o)
    if len(opmerkingen):
        _opmerkingen(wb, opmerkingen)
    _per_week(wb, o)
    _per_dag(wb, o)
    _orders(wb, o)
    _facturen(wb, o)
    if o.bank is not None:
        _bank(wb, o)
    _bronbladen(wb, o)
    del wb["Sheet"]
    for ws in wb:  # de vijf hoofdbladen blijven zichtbaar, de verdieping is verborgen (in Excel terug te halen)
        hoofd = ws.title in HOOFDBLADEN
        ws.sheet_properties.tabColor = TAB_HOOFD if hoofd else TAB_DIEPTE
        ws.sheet_state = "visible" if hoofd else "hidden"
    pad.parent.mkdir(parents=True, exist_ok=True)
    wb.save(pad)
