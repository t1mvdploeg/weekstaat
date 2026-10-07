"""De opdrachtregel op de voorbeeldset."""

import os
import shutil

import pandas as pd
import pytest
from conftest import MEDEWERKERS, VOORBEELD
from openpyxl import load_workbook

from weekstaat import dossier, inlezen
from weekstaat.cli import _in_naam, _toon, main


def test_aansluiten_schrijft_een_overzicht_en_drukt_de_samenvatting_af(tmp_path, capsys):
    naam = MEDEWERKERS[0]
    assert main(["aansluiten", str(VOORBEELD), "--uit", str(tmp_path), "--medewerker", naam]) == 0
    assert (tmp_path / f"Overzicht {naam}.xlsx").exists()
    uit = capsys.readouterr().out
    assert "5 orders, 154 urenregels, 18 bankregels" in uit
    assert "Klopt" in uit and "Totaal" in uit


def test_een_onbekende_medewerker_geeft_een_melding_en_geen_stacktrace(tmp_path, capsys):
    assert main(["aansluiten", str(VOORBEELD), "--uit", str(tmp_path), "--medewerker", "Bestaat Niet"]) == 1
    assert "geen uren gevonden voor Bestaat Niet" in capsys.readouterr().err


def test_een_map_zonder_instellingen_geeft_een_melding(tmp_path, capsys):
    assert main(["aansluiten", str(tmp_path)]) == 1
    assert "instellingen.toml" in capsys.readouterr().err


def test_omzetten_bouwt_de_ruwe_exporten_om_naar_de_sjablonen(tmp_path):
    ruw = VOORBEELD / "ruw"
    assert main(["omzetten-uren", str(ruw / "uren-export.xlsx"), str(tmp_path / "uren.csv")]) == 0
    assert len(inlezen.lees_uren(tmp_path / "uren.csv")) == len(inlezen.lees_uren(VOORBEELD / "uren.csv"))
    opdracht = ["omzetten-bank", str(tmp_path / "bank.csv"), str(ruw / "bank-g-rekening.csv")]
    opdracht += [str(ruw / "bank-gewone-rekening.csv"), "--instellingen", str(VOORBEELD / "instellingen.toml")]
    assert main(opdracht + ["--debiteurenkaart", str(ruw / "debiteurenkaart.xlsx")]) == 0
    assert (
        inlezen.lees_bank(tmp_path / "bank.csv")
        .drop(columns="referentie")
        .equals(inlezen.lees_bank(VOORBEELD / "bank.csv"))
    )


def test_omzetten_kaart_bouwt_de_debiteurenkaart_om_naar_het_sjabloon(tmp_path, capsys):
    naar = tmp_path / "debiteurenkaart.csv"
    assert main(["omzetten-kaart", str(VOORBEELD / "ruw" / "debiteurenkaart.xlsx"), str(naar)]) == 0
    assert f"geschreven: {naar}" in capsys.readouterr().out
    kaart = inlezen.lees_kaart(naar)
    assert set(kaart.soort) == {"factuur", "ontvangst"}
    assert main(["omzetten-kaart", str(tmp_path / "bestaat-niet.xlsx"), str(naar)]) == 1


def test_aansluiten_met_orders_uit_een_sjabloon_in_plaats_van_pdfs(tmp_path, capsys):
    map_ = tmp_path / "map"
    shutil.copytree(VOORBEELD, map_, ignore=shutil.ignore_patterns("orders", "mail", "ruw", "verwacht"))
    kop = pd.read_csv(VOORBEELD / "verwacht" / "orders-kop.csv", parse_dates=["factuurdatum"]).fillna(
        {"referentie": ""}
    )
    regels = pd.read_csv(VOORBEELD / "verwacht" / "orders-regels.csv", parse_dates=["datum"])
    dossier.schrijf(map_, kop=kop, regels=regels)
    assert main(["aansluiten", str(map_), "--uit", str(tmp_path / "uit"), "--medewerker", MEDEWERKERS[0]]) == 0
    assert "5 orders, 154 urenregels, 18 bankregels" in capsys.readouterr().out


def test_een_naam_met_een_pad_erin_blijft_binnen_de_uitvoermap(tmp_path):
    naam = _in_naam("../../buiten")
    assert "/" not in naam and (tmp_path / f"Overzicht {naam}.xlsx").resolve().parent == tmp_path.resolve()


def test_stuurtekens_in_een_naam_bereiken_de_terminal_niet(tmp_path, capsys):
    assert _toon("naam\x1b[2Jverder\x07") == "naam?[2Jverder?"
    assert _toon("naam\nKlopt  999") == "naam?Klopt  999"  # een regeleinde kan geen eigen regel in de uitvoer maken
    assert main(["aansluiten", str(VOORBEELD), "--uit", str(tmp_path), "--medewerker", "x\x1b[31m"]) == 1
    assert "\x1b" not in capsys.readouterr().err


def test_aansluiten_schrijft_overzichten_een_totaalbestand_en_bevindingen(tmp_path, capsys):
    uit = tmp_path / "uit"
    assert main(["aansluiten", str(VOORBEELD), "--uit", str(uit)]) == 0
    for naam in MEDEWERKERS:
        assert (uit / f"Overzicht {naam}.xlsx").exists()
    wb = load_workbook(uit / "Overzicht totaal.xlsx")
    assert wb.sheetnames[:5] == ["Samenvatting", "Actielijst", "Zonder order alle", "Totaal uitgebreid", "Opmerkingen"]
    assert len(wb.sheetnames) == 5 + 2 * len(MEDEWERKERS)
    bevindingen = pd.read_csv(uit / "bevindingen.csv", sep=";", decimal=",", keep_default_na=False)
    assert list(bevindingen.columns) == [
        "sleutel", "soort", "medewerker", "onderwerp", "bedrag", "btw", "wat", "bewijs", "actie", "wie", "groep"
    ]  # fmt: skip
    assert len(bevindingen) > 0 and bevindingen.sleutel.is_unique
    for naam in MEDEWERKERS:  # de bladen die het overzicht van een medewerker met kaart en opmerkingen heeft
        eigen = load_workbook(uit / f"Overzicht {naam}.xlsx").sheetnames
        assert "Debiteurenkaart" in eigen and "Opmerkingen" in eigen
    regels = capsys.readouterr().out.rstrip().splitlines()
    assert regels[-2].startswith("Totaal van alle medewerkers  ->  ") and regels[-2].endswith("Overzicht totaal.xlsx")
    assert regels[-1].startswith(f"{len(bevindingen)} bevindingen  ->  ") and regels[-1].endswith("bevindingen.csv")


def test_opmerkingen_in_het_overzicht_zijn_alleen_die_van_de_medewerker(tmp_path):
    map_ = tmp_path / "map"
    shutil.copytree(VOORBEELD, map_, ignore=shutil.ignore_patterns("ruw", "verwacht"))
    kolommen = inlezen.OPMERKINGEN
    rij = lambda m, nr, onderwerp: [m, nr, "ja", onderwerp, None, "", "Wat", "Bewijs", "Actie", "Het bureau", ""]  # noqa: E731
    handmatig = pd.DataFrame(
        [rij(MEDEWERKERS[0], "1", "Eerste"), rij(MEDEWERKERS[1], "1", "Tweede"), rij("", "2", "Voor iedereen")],
        columns=kolommen,
    )
    dossier.schrijf(map_, opmerkingen=handmatig.astype({"bedrag": float}))
    uit = tmp_path / "uit"
    assert main(["aansluiten", str(map_), "--uit", str(uit)]) == 0
    eigen = load_workbook(uit / f"Overzicht {MEDEWERKERS[0]}.xlsx")["Opmerkingen"]
    onderwerpen = [r[1] for r in eigen.iter_rows(min_row=2, values_only=True)]
    assert "Eerste" in onderwerpen and "Tweede" not in onderwerpen and "Voor iedereen" not in onderwerpen
    alle = load_workbook(uit / "Overzicht totaal.xlsx")["Opmerkingen"]
    assert {"Eerste", "Tweede", "Voor iedereen"} <= {r[3] for r in alle.iter_rows(min_row=2, values_only=True)}


def test_een_onbekende_sleutel_in_de_opmerkingen_komt_als_let_op_regel(tmp_path, capsys):
    map_ = tmp_path / "map"
    shutil.copytree(VOORBEELD, map_, ignore=shutil.ignore_patterns("ruw", "verwacht"))
    rij = [
        MEDEWERKERS[0],
        "7",
        "nee",
        "Typfout",
        None,
        "",
        "Wat",
        "Bewijs",
        "Actie",
        "Het bureau",
        "uren-zonder-ordr|x|y",
    ]
    dossier.schrijf(map_, opmerkingen=pd.DataFrame([rij], columns=inlezen.OPMERKINGEN).astype({"bedrag": float}))
    assert main(["aansluiten", str(map_), "--uit", str(tmp_path / "uit")]) == 0
    fout = capsys.readouterr().err
    assert "let op: let op: opmerking 7" not in fout
    assert "let op: " in fout and "opmerking 7" in fout and "uren-zonder-ordr|x|y" in fout


# ---------- schrijven mislukt en een bestand van de verkeerde soort ----------


def test_een_totaalbestand_dat_niet_te_schrijven_is_geeft_een_melding_met_het_pad(tmp_path, capsys):
    uit = tmp_path / "uit"
    (uit / "Overzicht totaal.xlsx").mkdir(parents=True)  # een map waar het bestand moet komen
    assert main(["aansluiten", str(VOORBEELD), "--uit", str(uit)]) == 1
    fout = capsys.readouterr().err
    assert fout.startswith("weekstaat: kan ") and "Overzicht totaal.xlsx niet schrijven" in fout
    assert "staat het bestand open in Excel?" in fout and "Errno" not in fout and "Traceback" not in fout


def test_bevindingen_die_niet_te_schrijven_zijn_geven_een_melding_met_het_pad(tmp_path, capsys):
    uit = tmp_path / "uit"
    (uit / "bevindingen.csv").mkdir(parents=True)
    assert main(["aansluiten", str(VOORBEELD), "--uit", str(uit)]) == 1
    fout = capsys.readouterr().err
    assert "kan " in fout and "bevindingen.csv niet schrijven: staat het bestand open in Excel?" in fout


@pytest.mark.skipif(not hasattr(os, "geteuid") or os.geteuid() == 0, reason="root mag ook in een alleen-lezen map")
def test_een_overzicht_dat_alleen_lezen_is_geeft_een_melding_met_het_pad(tmp_path, capsys):
    naam = MEDEWERKERS[0]
    assert main(["aansluiten", str(VOORBEELD), "--uit", str(tmp_path), "--medewerker", naam]) == 0
    pad = tmp_path / f"Overzicht {naam}.xlsx"
    pad.chmod(0o444)  # zo staat een bestand dat Excel (Windows) open heeft voor het programma
    capsys.readouterr()
    try:
        assert main(["aansluiten", str(VOORBEELD), "--uit", str(tmp_path), "--medewerker", naam]) == 1
    finally:
        pad.chmod(0o644)
    fout = capsys.readouterr().err
    assert f"kan {pad} niet schrijven: staat het bestand open in Excel?" in fout and "Errno" not in fout


def test_omzetten_uren_op_een_csv_zegt_welk_bestand_en_wat_er_verwacht_wordt(tmp_path, capsys):
    csv = tmp_path / "uren.csv"
    csv.write_text("a;b\n1;2\n", encoding="utf-8")
    assert main(["omzetten-uren", str(csv), str(tmp_path / "uit.csv")]) == 1
    fout = capsys.readouterr().err
    assert f"{csv}:" in fout and ".xlsx" in fout and "Salesforce" in fout
    assert "cannot be determined" not in fout and not (tmp_path / "uit.csv").exists()


def test_omzetten_kaart_en_bank_met_een_bestand_van_de_verkeerde_soort(tmp_path, capsys):
    csv = tmp_path / "kaart.csv"
    csv.write_text("a;b\n", encoding="utf-8")
    assert main(["omzetten-kaart", str(csv), str(tmp_path / "uit.csv")]) == 1
    assert f"{csv}:" in capsys.readouterr().err
    opdracht = ["omzetten-bank", str(tmp_path / "bank.csv"), str(VOORBEELD / "ruw" / "bank-g-rekening.csv")]
    opdracht += ["--instellingen", str(VOORBEELD / "instellingen.toml"), "--debiteurenkaart", str(csv)]
    assert main(opdracht) == 1
    assert f"{csv}:" in capsys.readouterr().err
    assert main(["omzetten-bank", str(tmp_path / "bank.csv"), str(VOORBEELD / "ruw" / "uren-export.xlsx"),
                 "--instellingen", str(VOORBEELD / "instellingen.toml")]) == 1  # fmt: skip
    assert "uren-export.xlsx" in capsys.readouterr().err


def test_een_bestand_dat_niet_bestaat_geeft_een_melding_in_het_nederlands(tmp_path, capsys):
    assert main(["omzetten-uren", str(tmp_path / "weg.xlsx"), str(tmp_path / "uit.csv")]) == 1
    fout = capsys.readouterr().err
    assert "weg.xlsx" in fout and "niet gevonden" in fout and "Errno" not in fout
