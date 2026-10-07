"""Gedeelde hulpjes voor de tests van het model, de bevindingen en de werkboeken: kleine, verzonnen dossiers (een
medewerker met facturen, orders, bank en debiteurenkaart), het overzicht van een medewerker daaruit en het doorrekenen
van een werkboek."""

import tempfile
from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from conftest import VOORBEELD
from openpyxl import load_workbook

from weekstaat import aansluiten, dossier, overzicht
from weekstaat.overzicht import maak


def aansluiting(d: dossier.Dossier, naam: str, uren: pd.DataFrame | None = None) -> aansluiten.Aansluiting:
    uren = d.uren if uren is None else uren
    return aansluiten.sluit_aan(uren, d.kop, d.regels, naam, d.inst, d.bank, d.vervallen, d.saldo(naam))


def overzicht_van(d: dossier.Dossier, naam: str, **eigen) -> overzicht.Overzicht:
    return maak(aansluiting(d, naam), d, **eigen)


def doorgerekend(pad: Path) -> dict[str, list[list]]:
    """Rekent een werkboek door met `formulas` (alleen in de tests aanwezig); blad (in hoofdletters) -> rijen. Zonder
    `formulas` wordt de test overgeslagen."""
    formulas = pytest.importorskip("formulas")
    with tempfile.TemporaryDirectory() as tijdelijk:
        model = formulas.ExcelModel().loads(str(pad)).finish()
        model.calculate()
        model.write(dirpath=tijdelijk)
        wb = load_workbook(next(Path(tijdelijk).glob("*")), data_only=True)
        return {ws.title.upper(): [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}


def zonder_betalingen(a: aansluiten.Aansluiting) -> aansluiten.Aansluiting:
    """Dezelfde aansluiting waarin geen order is gekoppeld aan een bankboeking: er is geen betaling gevonden. De
    boekingen blijven staan, zonder orders, zodat `betaald` en `bankregels` elkaar niet tegenspreken."""
    boekingen = a.bankregels.assign(orders=[[] for _ in range(len(a.bankregels))])
    return replace(a, betaald={}, bankregels=boekingen)


FEMKE = "Femke Jansma"
DAAN = "Daan Postma"
IBAN_G, IBAN_GEWOON = "NL00RABO0991000001", "NL00RABO0123456789"
NAMEN = 'bureau = "Studio Noord"\nopdrachtgever = "Oeverland"\n\n'


def dagregels(maandag: str, dagen: int, tarief: float) -> list[tuple[str, float, float]]:
    """Regels (datum, uren, tarief) van `dagen` achtereenvolgende dagen van acht uur vanaf `maandag`."""
    return [(str(pd.Timestamp(maandag) + pd.Timedelta(days=i))[:10], 8.0, tarief) for i in range(dagen)]


def mini_dossier(tmp_path, facturen, orders, bank, kaart, ander=True, namen=NAMEN) -> dossier.Dossier:
    """Een dossier met een medewerker (Femke) en, met `ander`, een tweede (Daan) met één factuur.

    `facturen` zijn (factuurnummer, factuurdatum, regels), `orders` (ordernummer, orderdatum, medewerker, regels) met
    regels als (datum, uren, tarief), `bank` (datum, iban, bedrag) en `kaart` (datum, soort, nummer, iban, bedrag).
    Alles is verzonnen; de btw is 21%."""
    uren = []
    if ander:
        facturen = [*facturen, ("F-401", "2025-03-31", [("2025-03-10", 8.0, 31.0)])]
    for nr, factuurdatum, regels in facturen:
        wie = DAAN if nr == "F-401" else FEMKE
        for dag, aantal, tarief in regels:
            bedrag = round(aantal * tarief, 2)
            uren.append(
                dict(
                    rij=len(uren) + 2, medewerker=wie, urenstaat="U-1", datum=pd.Timestamp(dag), uren=aantal,
                    soort="Normale Uren", tarief=tarief, status="Goedgekeurd", factuurnummer=nr,
                    factuurdatum=pd.Timestamp(factuurdatum), bedrag=bedrag, creditnummer="", creditbedrag=0.0,
                    factuurbedrag=bedrag,
                )
            )  # fmt: skip
    kop, regels = [], []
    for nr, datum, wie, rijen in orders:
        excl = round(sum(a * t for _, a, t in rijen), 2)
        bestand = f"{nr}.pdf"
        kop.append(
            dict(bestand=bestand, order=nr, entiteit="OI", factuurdatum=pd.Timestamp(datum), excl=excl,
                 incl=round(excl * 1.21, 2), referentie="")
        )  # fmt: skip
        for dag, aantal, tarief in rijen:
            regels.append(
                dict(bestand=bestand, order=nr, entiteit="OI", medewerker=wie, project="P-1",
                     datum=pd.Timestamp(dag), omschrijving="Uren", eenheid="Uren", aantal=aantal, tarief=tarief,
                     bedrag=round(aantal * tarief, 2))
            )  # fmt: skip
    banken = pd.DataFrame(
        [(pd.Timestamp(d), i, b, "Oeverland infra", "betaling") for d, i, b in bank],
        columns=["datum", "rekening", "bedrag", "tegenpartij", "omschrijving"],
    )
    kaarten = pd.DataFrame(
        [(pd.Timestamp(d), s, n, "", i, b) for d, s, n, i, b in kaart],
        columns=["datum", "soort", "nummer", "omschrijving", "rekening", "bedrag"],
    )
    map_ = tmp_path / "mini"
    dossier.schrijf(
        map_, instellingen=namen + (VOORBEELD / "instellingen.toml").read_text(encoding="utf-8"),
        uren=pd.DataFrame(uren), kop=pd.DataFrame(kop), regels=pd.DataFrame(regels), bank=banken, kaart=kaarten,
    )  # fmt: skip
    return dossier.lees(map_)


# De basis: drie facturen van Femke (week 11, 12 en 13), twee orders (week 11 en 12), vier bankregels.
FACTUREN = [
    ("F-101", "2025-03-17", dagregels("2025-03-10", 5, 40.0)),
    ("F-102", "2025-03-24", dagregels("2025-03-17", 4, 40.0)),
    ("F-103", "2025-03-31", dagregels("2025-03-24", 5, 40.0)),
]


def orders_basis(tarief_a: float = 40.0, tarief_b: float = 40.0):
    return [
        ("I01250001", "2025-03-17", FEMKE, dagregels("2025-03-10", 5, tarief_a)),
        ("I01250002", "2025-03-24", FEMKE, dagregels("2025-03-17", 4, tarief_b)),
    ]


BANK_B = [("2025-04-25", IBAN_G, 464.64), ("2025-04-25", IBAN_GEWOON, 1084.16)]  # order 2: 1548,80 incl
KAART_B = [  # order 2 is betaald maar nog niet afgeletterd; de facturen van week 12 en 13 staan open
    ("2025-03-24", "factuur", "F-102", "", 1548.80), ("2025-03-31", "factuur", "F-103", "", 1936.00),
    ("2025-03-31", "factuur", "F-401", "", 300.08), ("2025-03-31", "factuur", "F-900", "", 121.00),
    ("2025-04-25", "ontvangst", "", IBAN_G, -464.64), ("2025-04-25", "ontvangst", "", IBAN_GEWOON, -1084.16),
]  # fmt: skip


def kaart_dossier(tmp_path, tarief_a=40.0, bank_a=None, kaart=None, orders=None, bank_b=None):
    """De basis, met bij order 1 een eigen tarief, eigen bankregels en eigen kaartregels."""
    bank_a = bank_a if bank_a is not None else [("2025-04-14", IBAN_G, 580.80), ("2025-04-14", IBAN_GEWOON, 1355.20)]
    return mini_dossier(
        tmp_path,
        FACTUREN,
        orders or orders_basis(tarief_a),
        bank_a + (BANK_B if bank_b is None else bank_b),
        KAART_B if kaart is None else kaart,
    )
