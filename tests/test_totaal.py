"""Tests voor het totaalbestand (`totaal.schrijf`): de medewerkers naast elkaar en opgeteld, de samenvatting, de
actielijst, de opmerkingen en de bladen per medewerker. Alles op kleine, verzonnen dossiers uit `hulp`."""

from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from hulp import (
    DAAN,
    FACTUREN,
    FEMKE,
    IBAN_G,
    IBAN_GEWOON,
    KAART_B,
    doorgerekend,
    kaart_dossier,
    mini_dossier,
    orders_basis,
    overzicht_van,
)
from openpyxl import load_workbook

from weekstaat import bevindingen, dossier, inlezen, totaal
from weekstaat import overzicht as model

BUREAU, OPDRACHTGEVER = "Studio Noord", "Oeverland"
TOT = "Totaal uitgebreid"
VOORKANT = ["Samenvatting", "Actielijst", "Zonder order alle", TOT, "Opmerkingen"]


def maak(d: dossier.Dossier, tmp_path: Path, *namen: str, korte=None, naam: str = "totaal.xlsx") -> Path:
    """Leidt alles af zoals de opdrachtregel en schrijft het totaalbestand."""
    overzichten = [overzicht_van(d, n) for n in namen or (FEMKE, DAAN)]
    bs = bevindingen.leid_af(overzichten, d)
    opm = bevindingen.opmerkingen(bs, d)
    act = bevindingen.acties(bs, d, korte)
    pad = tmp_path / naam
    totaal.schrijf(overzichten, opm, act, pad)
    return pad


def rijen(ws) -> list[list]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def zoek(ws_rijen: list[list], begin: str, kolom: int = 0) -> list:
    """De eerste rij waarvan kolom A met `begin` begint."""
    return next(r for r in ws_rijen if isinstance(r[kolom], str) and r[kolom].startswith(begin))


# ---------- korte namen ----------


def test_korte_bladnamen():
    k = totaal.korte_namen(["Thijmen Lucas Van Dijk", "Thijmen Bakker", "Sanne/Bakker"])
    assert len(set(k.values())) == 3
    assert all(len("Zonder order " + v) <= 31 and not set(v) & set("/\\?*[]:") for v in k.values())
    assert k["Thijmen Lucas Van Dijk"] == "Thijmen D" and k["Thijmen Bakker"] == "Thijmen B"
    assert k["Sanne/Bakker"] == "Sanne_Bakker"


def test_korte_namen_zijn_voornamen_zolang_die_verschillen():
    assert totaal.korte_namen(["Femke Jansma", "Daan Postma"]) == {"Femke Jansma": "Femke", "Daan Postma": "Daan"}


def test_korte_namen_blijven_uniek_na_afkappen_en_bij_gelijke_beginletters():
    namen = [
        "Jan Dijk",
        "Jan De Vries",
        "Jan Dekker",
        "Bartholomeus-Alexander-Maximiliaan Visser",
        "Bartholomeus-Alexander-Maximiliaan Vos",
    ]
    k = totaal.korte_namen(namen)
    assert len({v.lower() for v in k.values()}) == len(namen)  # Excel kent bladnamen niet hoofdlettergevoelig
    assert all(len("Zonder order " + v) <= 31 for v in k.values())


# ---------- het bestand als geheel ----------


def test_bladen_volgorde_kleur_en_zichtbaarheid(tmp_path):
    wb = load_workbook(maak(kaart_dossier(tmp_path), tmp_path))
    assert wb.sheetnames == [*VOORKANT, "Stand Femke", "Zonder order Femke", "Stand Daan", "Zonder order Daan"]
    assert all(ws.sheet_state == "visible" for ws in wb)
    kleuren = {ws.title: ws.sheet_properties.tabColor.rgb[-6:] for ws in wb}
    assert [t for t, k in kleuren.items() if k == "1F4E78"] == VOORKANT[:3]
    assert all(k == "A6A6A6" for t, k in kleuren.items() if t not in VOORKANT[:3])


def test_totaal_is_de_som_van_de_medewerkers(tmp_path):
    d = kaart_dossier(tmp_path)
    pad = maak(d, tmp_path)
    ovs = [overzicht_van(d, n) for n in (FEMKE, DAAN)]
    rekenblad = doorgerekend(pad)
    t = rekenblad[TOT.upper()]
    for begin, sleutel in (
        ("Gewerkt volgens", "urenstaat"), (f"Door {BUREAU} gefactureerd", "gefactureerd"),
        (f"Door {OPDRACHTGEVER} op orders", "op_orders"), (f"Door {OPDRACHTGEVER} betaald", "betaald"),
    ):  # fmt: skip
        r = zoek(t, begin)
        assert r[1] == pytest.approx(ovs[0].stand[sleutel]) and r[2] == pytest.approx(ovs[1].stand[sleutel])
        assert r[3] == pytest.approx(r[1] + r[2])
    for r in t:
        if isinstance(r[0], str) and all(isinstance(x, int | float) for x in r[1:4]) and r[0] != "Totaal":
            assert r[3] == pytest.approx(r[1] + r[2], abs=0.005), r[0]
    fouten = [(b, c) for b, rr in rekenblad.items() for r in rr for c in r if isinstance(c, str) and c.startswith("#")]
    assert not fouten
    # blok 2 sluit op blok 1 en de samenvatting verwijst naar dezelfde cijfers
    assert zoek(t, "Totaal")[-1].startswith("Controle:")
    s = rekenblad["SAMENVATTING"]
    for label, begin in (("Gewerkt volgens", "Gewerkt volgens"), (f"Door {BUREAU} gefactureerd", f"Door {BUREAU}")):
        assert zoek(s, label)[1:4] == zoek(t, begin)[1:4]


def test_blok_5_sluit_op_het_saldo_van_de_kaart(tmp_path):
    d = kaart_dossier(tmp_path)
    t = doorgerekend(maak(d, tmp_path))[TOT.upper()]
    k = overzicht_van(d, FEMKE).kaart
    assert zoek(t, "Hele debiteurenkaart")[3] == pytest.approx(k.heel["open"])
    assert zoek(t, "Open facturen die niet in de uren staan")[3] == pytest.approx(k.heel["buiten_uren"])
    controle = zoek(t, "Controle: per saldo")
    assert (
        controle[3] == pytest.approx(0.0, abs=0.05) and controle[-1] == "Sluit aan op het saldo van de debiteurenkaart"
    )


def test_blok_5_volgt_de_tolerantie_van_het_model(tmp_path, monkeypatch):
    """Eén tolerantie voor de kaart: `overzicht.CONTROLE_VERSCHIL`. Met een negatieve waarde sluit niets meer."""
    d = kaart_dossier(tmp_path)
    monkeypatch.setattr(model, "CONTROLE_VERSCHIL", -1.0)
    t = rijen(load_workbook(maak(d, tmp_path))[TOT])
    assert zoek(t, "Controle: per saldo")[-1].startswith("LET OP: sluit niet aan op het saldo van de debiteurenkaart")


def test_een_kaart_die_niet_sluit_stopt_niet_maar_zegt_het(tmp_path):
    kaart = [("2025-03-24", "factuur", "F-102", "", 1000.00), *KAART_B[1:]]  # 548,80 te laag
    d = kaart_dossier(tmp_path, kaart=kaart)
    t = doorgerekend(maak(d, tmp_path))[TOT.upper()]
    teksten = [c for r in t for c in r if isinstance(c, str) and "sluit niet aan" in c]
    van_femke = f"LET OP: de debiteurenkaart van {FEMKE} sluit niet aan"
    assert any(c.startswith(van_femke) and "548,80" in c for c in teksten)
    assert zoek(t, "Controle: per saldo")[-1] == "Sluit aan op het saldo van de debiteurenkaart"  # de kaart zelf klopt


def test_saldo_dat_niet_sluit_geeft_een_let_op_met_het_verschil(tmp_path):
    d = kaart_dossier(tmp_path)
    ovs = [overzicht_van(d, n) for n in (FEMKE, DAAN)]
    ovs[0].kaart.per_saldo += 10.0  # een fout in het model, om de controle te zien
    pad = tmp_path / "t.xlsx"
    totaal.schrijf(ovs, bevindingen.opmerkingen([], d), bevindingen.acties([], d), pad)
    controle = zoek(doorgerekend(pad)[TOT.upper()], "Controle: per saldo")
    assert controle[3] == pytest.approx(10.0)
    assert controle[-1] == "LET OP: sluit niet aan op het saldo van de debiteurenkaart, verschil 10,00"


def test_een_medewerker(tmp_path):
    d = kaart_dossier(tmp_path)
    pad = maak(d, tmp_path, FEMKE)
    wb = load_workbook(pad)
    assert wb.sheetnames == [*VOORKANT, "Stand Femke", "Zonder order Femke"]
    t = rijen(wb[TOT])
    kop = zoek(t, "1. Van urenstaat")
    assert kop[1:3] == ["Femke", "Totaal"]  # de kolom Totaal blijft staan
    rekenblad = doorgerekend(pad)[TOT.upper()]
    r = zoek(rekenblad, "Gewerkt volgens")
    assert r[1] == r[2] == pytest.approx(overzicht_van(d, FEMKE).stand["urenstaat"])
    # de kaart bevat ook facturen van een medewerker die niet in dit bestand staat: dat staat er als feit, geen LET OP
    controle = zoek(rekenblad, "Controle: per saldo")
    assert DAAN in controle[-1] and "LET OP" not in controle[-1]


def test_samenvatting_zonder_kaart_heeft_geen_blok_echt_open(tmp_path):
    d = replace(kaart_dossier(tmp_path), kaart=None)
    wb = load_workbook(maak(d, tmp_path))
    s, t = rijen(wb["Samenvatting"]), rijen(wb[TOT])
    assert not [r for r in s if isinstance(r[0], str) and r[0].startswith("Wat staat er echt open")]
    assert not [
        r for r in t if isinstance(r[0], str) and r[0].startswith(("5.", "Open facturen op de debiteurenkaart"))
    ]
    assert zoek(s, "Door " + OPDRACHTGEVER + " betaald")  # de bank is er wel
    assert zoek(t, "6.")  # en de volgende blokken lopen door


def test_zonder_bank_geen_regel_betaald(tmp_path):
    d = replace(kaart_dossier(tmp_path), bank=None)
    wb = load_workbook(maak(d, tmp_path))
    for blad in ("Samenvatting", TOT):
        r = rijen(wb[blad])
        assert not [x for x in r if isinstance(x[0], str) and x[0].startswith(f"Door {OPDRACHTGEVER} betaald")]
        assert not [x for x in r if isinstance(x[0], str) and x[0].startswith("Op een order, nog niet betaald")]
        assert not [x for x in r if isinstance(x[0], str) and "debiteurenkaart" in x[0].lower()]
    assert zoek(rijen(wb["Samenvatting"]), "Door " + OPDRACHTGEVER + " op orders gezet")


def test_regel_over_een_toezegging_alleen_als_iemand_er_een_heeft(tmp_path):
    zonder = load_workbook(maak(kaart_dossier(tmp_path / "a"), tmp_path / "a"))
    assert not [r for r in rijen(zonder["Samenvatting"]) if isinstance(r[0], str) and "toezegging" in r[0]]
    kolommen = ["medewerker", "nr", "datum", "week", "maand", "omschrijving", "eenheid", "aantal", "prijs", "bedrag"]
    t = pd.DataFrame(
        [[FEMKE, "T-007", pd.Timestamp("2025-04-01"), "2025 wk 13", "", "Week 13", "Uren", 40.0, 40.0, 1600.0]],
        columns=kolommen,
    )
    d = replace(kaart_dossier(tmp_path / "b"), toezegging=t)
    met = load_workbook(maak(d, tmp_path / "b"))
    r = zoek(rijen(met["Samenvatting"]), "Waarvan uren op een toezegging")
    assert r[0] and "T-007" not in r[0]
    van = zoek(
        doorgerekend(tmp_path / "b" / "totaal.xlsx")[TOT.upper()],
        "Van 'Opdrachtgever te weinig opgenomen' staat al op toezegging T-007",
    )
    assert van[1:4] == [1600.0, 0.0, 1600.0]
    uit = doorgerekend(tmp_path / "b" / "totaal.xlsx")["SAMENVATTING"]
    assert zoek(uit, "Waarvan uren op een toezegging")[1:4] == [1600.0, 0.0, 1600.0]


def test_zonder_order_alle_heeft_elke_factuur_van_alle_medewerkers(tmp_path):
    d = kaart_dossier(tmp_path)
    ovs = [overzicht_van(d, n) for n in (FEMKE, DAAN)]
    pad = maak(d, tmp_path)
    r = rijen(load_workbook(pad)["Zonder order alle"])
    kop = zoek(r, "Medewerker")
    data = [x for x in r if x[0] in (FEMKE, DAAN)]
    assert len(data) == sum(len(o.zonder_order) for o in ovs) and len(data) > 0
    assert {x[0] for x in data} == {FEMKE, DAAN}
    i = kop.index("Zonder order")
    rek = doorgerekend(pad)["ZONDER ORDER ALLE"]
    assert zoek(rek, "Totaal")[i] == pytest.approx(sum(o.zonder_order.zonder_order.sum() for o in ovs))
    blok4 = zoek(doorgerekend(pad)[TOT.upper()], "Totaal zonder order")
    assert blok4[1] == pytest.approx(ovs[0].zonder_order.zonder_order.sum())
    assert blok4[2] == pytest.approx(ovs[1].zonder_order.zonder_order.sum())


# ---------- bladen per medewerker ----------


def test_stand_per_medewerker_zonder_leeswijzer_en_met_waarden(tmp_path):
    d = kaart_dossier(tmp_path)
    wb = load_workbook(maak(d, tmp_path))
    for blad in ("Stand Femke", "Zonder order Femke"):
        ws = wb[blad]
        assert not [c for r in ws.iter_rows() for c in r if c.data_type == "f"]  # alleen waarden
        assert not [r for r in rijen(ws) if isinstance(r[0], str) and r[0].startswith("Zo lees je dit bestand")]
    t = rijen(wb[TOT])
    assert zoek(t, "Zo lees je dit bestand")
    wijzer = [r[-1] for r in t if r[-1] and "Overzicht <naam>.xlsx" in str(r[-1])]
    assert wijzer, "de leeswijzer zegt dat verwijzingen naar andere bladen gelden voor Overzicht <naam>.xlsx"


# ---------- opmerkingen en actielijst ----------


def handmatige_opmerkingen(d, regels: list[list]) -> dossier.Dossier:
    return replace(d, opmerkingen=pd.DataFrame(regels, columns=inlezen.OPMERKINGEN).astype({"bedrag": float}))


def test_opmerkingen_blad_heeft_alle_opmerkingen_met_bron(tmp_path):
    d = kaart_dossier(tmp_path)
    sleutel = f"uren-zonder-order|{FEMKE}|2025 wk 13-16"
    d = handmatige_opmerkingen(
        d, [[FEMKE, "1", "ja", "Eigen punt", 12.5, "ex", "Wat", "Bewijs", "Actie", BUREAU, sleutel],
            ["", "2", "nee", "Voor iedereen", None, "", "Wat", "Bewijs", "Actie", BUREAU, ""]],
    )  # fmt: skip
    overzichten = [overzicht_van(d, n) for n in (FEMKE, DAAN)]
    bs = bevindingen.leid_af(overzichten, d)
    opm = bevindingen.opmerkingen(bs, d)
    pad = maak(d, tmp_path)
    ws = load_workbook(pad)["Opmerkingen"]
    kop = [c.value for c in ws[1]]
    assert kop[0] == "Medewerker" and "Bron" in kop and "Belangrijk" in kop
    assert ws.max_row == 1 + len(opm)
    assert ws.freeze_panes == "E2" and ws.auto_filter.ref
    bron = kop.index("Bron")
    assert [r[bron] for r in ws.iter_rows(min_row=2, values_only=True)].count("handmatig") == 2
    # kolom Bevinding: de sleutel van de bevinding (automatisch) of wat in de csv staat (handmatig)
    assert kop[-1] == "Bevinding"
    sleutels = [r[-1] for r in ws.iter_rows(min_row=2, values_only=True)]
    automatisch = [r[-1] for r in ws.iter_rows(min_row=2, values_only=True) if r[bron] == "automatisch"]
    assert automatisch and all(s in {b.sleutel for b in bs} for s in automatisch)
    assert len(set(automatisch)) == len(automatisch)  # elke sleutel een keer
    handmatig = [
        s for s, r in zip(sleutels, ws.iter_rows(min_row=2, values_only=True), strict=True) if r[bron] == "handmatig"
    ]
    assert handmatig == [sleutel, None]  # wat in kolom `bevinding` van de csv staat
    assert sleutel not in automatisch  # de handmatige opmerking vervangt de automatische
    kolom = kop.index("Belangrijk") + 1
    ja = [c for c in ws.iter_rows(min_row=2, min_col=kolom, max_col=kolom) for c in c if c.value == "ja"]
    assert ja and all(c.fill.fgColor.rgb[-6:] == "FFF2CC" for c in ja)


def test_tekst_uit_een_bron_wordt_nooit_een_formule(tmp_path):
    d = kaart_dossier(tmp_path)
    d = handmatige_opmerkingen(
        d, [[FEMKE, "=1+1", "ja", "=SUM(A1)", None, "", "=HYPERLINK(1)", "=1", "=2", "=3", ""]]
    )  # fmt: skip
    todo = [("=Groep", [("=Titel", "=Toelichting van Femke")])]
    d = replace(d, todo=todo)
    wb = load_workbook(maak(d, tmp_path, korte={FEMKE: "Femke", DAAN: "Daan"}))
    for blad in ("Opmerkingen", "Actielijst"):
        cellen = [c for r in wb[blad].iter_rows() for c in r if c.data_type == "f"]
        assert not cellen, f"{blad}: {[c.coordinate for c in cellen]}"
    assert any(c.value == "=Titel" for r in wb["Actielijst"].iter_rows() for c in r)


def test_actielijst_toont_de_acties_per_groep_met_voor_wie(tmp_path):
    d = kaart_dossier(tmp_path)
    overzichten = [overzicht_van(d, n) for n in (FEMKE, DAAN)]
    bs = bevindingen.leid_af(overzichten, d)
    korte = {FEMKE: "Femke", DAAN: "Daan"}
    act = bevindingen.acties(bs, d, korte)
    ws = load_workbook(maak(d, tmp_path, korte=korte))["Actielijst"]
    r = rijen(ws)
    koppen = [x for x in r if x[1:3] == ["Voor wie", "Toelichting"]]
    assert [x[0] for x in koppen] == [groep for groep, _ in act]
    titels = [x[0] for x in r if x[0] and x[1:3] != ["Voor wie", "Toelichting"]][2:]  # na titel en inleiding
    assert titels == [t for _, lijst in act for t, _, _ in lijst]
    assert ws.freeze_panes == "A4"


def test_alles_sluit_aan_geeft_een_actielijst_zonder_acties(tmp_path):
    d = kaart_dossier(tmp_path)
    ovs = [overzicht_van(d, FEMKE)]
    act = bevindingen.acties([], d)
    pad = tmp_path / "t.xlsx"
    totaal.schrijf(ovs, bevindingen.opmerkingen([], d), act, pad)
    ws = load_workbook(pad)["Actielijst"]
    assert [x[0] for x in rijen(ws) if x[0] == "Er zijn geen acties: alles sluit aan."]


def test_korte_namen_met_een_vaste_basis_en_een_volgnummer():
    k = totaal.korte_namen(["Jan Dijk", "Jan Dorp", "Jan Dekker", "Jan Dam", "Jan Dolk"])
    assert list(k.values()) == ["Jan D", "Jan D 2", "Jan D 3", "Jan D 4", "Jan D 5"]


# ---------- titels, lange namen ----------


def titels(pad: Path) -> list[str]:
    wb = load_workbook(pad)
    return [wb[blad]["A1"].value for blad in ("Samenvatting", TOT)]


def test_titels_zonder_naam_en_de_namen_in_de_regel_eronder_als_ze_zijn_ingesteld(tmp_path):
    eigen = load_workbook(maak(kaart_dossier(tmp_path / "eigen"), tmp_path / "eigen"))
    assert [eigen[b]["A1"].value for b in ("Samenvatting", TOT)] == [
        "Aansluiting in het kort",
        "Aansluiting: alle medewerkers samen",
    ]
    assert f"Uren van {BUREAU} bij {OPDRACHTGEVER}." in eigen["Samenvatting"]["A2"].value
    assert f"Uren van {BUREAU} bij {OPDRACHTGEVER}." in eigen[TOT]["A2"].value
    bank = [("2025-04-14", IBAN_G, 580.80), ("2025-04-14", IBAN_GEWOON, 1355.20)]
    standaard = mini_dossier(tmp_path / "st", FACTUREN[:1], orders_basis()[:1], bank, [], ander=False, namen="")
    wb = load_workbook(maak(standaard, tmp_path / "st", FEMKE))
    for blad in ("Samenvatting", TOT):
        a1, a2 = wb[blad]["A1"].value, wb[blad]["A2"].value
        assert a1[:1].isupper() and "het bureau-" not in a1 and "de opdrachtgever" not in a2 + a1
        assert "Uren van" not in a2  # bij de standaardnamen staat die zin er niet
    for blad in (wb["Samenvatting"], wb[TOT], wb["Actielijst"], wb["Zonder order alle"]):
        for r in blad.iter_rows(values_only=True):  # een tekst die met een naam uit de instellingen begint
            for c in r:
                if isinstance(c, str):
                    assert c[:1].isupper() or c[:1].isdigit() or not c[:1].isalpha(), (blad.title, c)


def test_lange_naam_met_verboden_tekens_geeft_korte_geldige_bladnamen(tmp_path):
    d = kaart_dossier(tmp_path)
    lang = "Bartholomeus/Alexander-Maximiliaan [x] Jansma"
    ander = "Daan: Postma?"
    hernoemd = {FEMKE: lang, DAAN: ander}
    uren = d.uren.assign(medewerker=d.uren.medewerker.replace(hernoemd))
    regels = d.regels.assign(medewerker=d.regels.medewerker.replace(hernoemd))
    d = replace(d, uren=uren, regels=regels)
    pad = maak(d, tmp_path, lang, ander)
    wb = load_workbook(pad)  # te openen
    kort = totaal.korte_namen([lang, ander])
    bladen = [f"{voor} {k}" for k in kort.values() for voor in ("Stand", "Zonder order")]
    assert sorted(wb.sheetnames[5:]) == sorted(bladen)
    assert all(len(n) <= 31 and not set(n) & set("/\\?*[]:") for n in wb.sheetnames)
    assert rijen(wb[f"Stand {kort[lang]}"])[0][0].startswith(lang)  # de volledige naam staat in het blad zelf
