"""Tests voor `werkboek`: de bladen, wat zichtbaar is, wat de gebruiker zelf invult, tekst die geen formule mag worden,
en dat de formules na het doorrekenen de cijfers van het model geven."""

import datetime
import itertools
import re
import shutil
from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from conftest import VOORBEELD
from hulp import KAART_B, doorgerekend, kaart_dossier, overzicht_van, zonder_betalingen
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from test_overzicht import SANNE, THIJMEN, toezegging

from weekstaat import aansluiten, dossier, inlezen, overzicht, totaal, werkboek
from weekstaat.cli import main

LEEG = pd.DataFrame(columns=inlezen.OPMERKINGEN)
MET_BANK = [
    "Stand", "Zonder order", "Per periode", "Verschillen", "Opmerkingen", "Per week", "Per dag", "Orders", "Facturen",
    "Bank", "Urenregels", "Orderregels",
]  # fmt: skip


def opmerkingen(*rijen: tuple) -> pd.DataFrame:
    """Opmerkingen met de kolommen van `inlezen.lees_opmerkingen`: (nr, belangrijk, onderwerp, wat, actie)."""
    tabel = pd.DataFrame(
        [
            (SANNE, nr, belangrijk, onderwerp, None, "ex", wat, "bewijs", actie, "bureau", "")
            for nr, belangrijk, onderwerp, wat, actie in rijen
        ],
        columns=inlezen.OPMERKINGEN,
    )
    return tabel.fillna({"wat": "", "actie": ""})


OPM = opmerkingen(
    ("1", "ja", "Eerste punt", "Er is iets aan de hand.", "Bel het bureau."),
    ("2", "nee", "Tweede punt", "Niet belangrijk.", "Niets."),
)


@pytest.fixture(scope="module")
def voorbeeld() -> dossier.Dossier:
    return dossier.lees(VOORBEELD)


@pytest.fixture(scope="module")
def overzicht_sanne(voorbeeld) -> overzicht.Overzicht:
    return overzicht_van(voorbeeld, SANNE)


@pytest.fixture(scope="module")
def sanne_zonder_kaart(voorbeeld) -> overzicht.Overzicht:
    """De voorbeeldset met bank, maar zonder debiteurenkaart (die heeft de voorbeeldset sinds de uitbreiding wel)."""
    return overzicht_van(replace(voorbeeld, kaart=None), SANNE)


@pytest.fixture(scope="module")
def met_kaart(tmp_path_factory) -> overzicht.Overzicht:
    """Een overzicht met bank en debiteurenkaart (een kleine verzonnen set)."""
    d = kaart_dossier(tmp_path_factory.mktemp("kaart"))
    return overzicht_van(d, "Femke Jansma")


def schrijf(o: overzicht.Overzicht, pad: Path, opm: pd.DataFrame = OPM) -> Path:
    werkboek.schrijf(o, opm, pad)
    return pad


# ---------- bladen ----------


def test_bladen_zichtbaar_en_verborgen(tmp_path, sanne_zonder_kaart):
    wb = load_workbook(schrijf(sanne_zonder_kaart, tmp_path / "a.xlsx"))
    assert wb.sheetnames == MET_BANK
    for ws in wb:
        hoofd = ws.title in werkboek.HOOFDBLADEN
        assert ws.sheet_state == ("visible" if hoofd else "hidden"), ws.title
        assert ws.sheet_properties.tabColor.rgb.endswith("1F4E78" if hoofd else "A6A6A6"), ws.title
    assert [ws.title for ws in wb if ws.sheet_state == "visible"] == [
        "Stand",
        "Zonder order",
        "Per periode",
        "Verschillen",
    ]


def test_met_kaart_staat_de_debiteurenkaart_derde_en_zichtbaar(tmp_path, met_kaart):
    wb = load_workbook(schrijf(met_kaart, tmp_path / "a.xlsx", LEEG))
    assert wb.sheetnames[:5] == list(werkboek.HOOFDBLADEN)
    assert wb["Debiteurenkaart"].sheet_state == "visible"
    assert "Opmerkingen" not in wb.sheetnames  # zonder opmerkingen geen blad


def test_leeswijzer_noemt_alleen_bladen_die_er_zijn_en_zegt_hoe_je_ze_terughaalt(
    tmp_path, sanne_zonder_kaart, met_kaart
):
    def leeswijzer(o, opm) -> list[str]:
        ws = load_workbook(schrijf(o, tmp_path / "a.xlsx", opm))["Stand"]
        rij = next(c.row for c in ws["A"] if c.value == "Zo lees je dit bestand")
        return [r[0] for r in ws.iter_rows(min_row=rij + 1, max_col=1, values_only=True) if r[0]]

    zonder = leeswijzer(sanne_zonder_kaart, LEEG)
    assert "Debiteurenkaart" not in zonder and "Opmerkingen" not in zonder and "Bank" in zonder
    assert "Verborgen bladen" in zonder
    assert "Debiteurenkaart" in leeswijzer(met_kaart, OPM)
    tekst = " ".join(
        str(c.value) for rij in load_workbook(tmp_path / "a.xlsx")["Stand"].iter_rows() for c in rij if c.value
    )
    assert "Zichtbaar maken (Unhide)" in tekst


def test_zonder_bank_en_zonder_kaart_ontbreken_de_bladen(tmp_path, voorbeeld):
    o = overzicht_van(replace(voorbeeld, bank=None), SANNE)
    wb = load_workbook(schrijf(o, tmp_path / "a.xlsx", LEEG))
    assert wb.sheetnames == [b for b in MET_BANK if b not in ("Bank", "Opmerkingen")]
    stand = [c.value for rij in wb["Stand"].iter_rows() for c in rij if c.value]
    assert "Er is geen bank aangeleverd: wat betaald is, is niet te zien." in " ".join(stand)
    assert not any(str(v).startswith("4. Bank") for v in stand)
    assert not any(str(v).startswith("Door de opdrachtgever betaald") for v in stand)
    assert not any(v in ("Bank", "Debiteurenkaart") for v in stand)  # ook niet in de leeswijzer
    zonder_order = [c.value for rij in wb["Zonder order"].iter_rows() for c in rij if c.value]
    assert not any(str(v).startswith("C. ") for v in zonder_order)
    assert not any("Blok C" in str(v) for v in zonder_order)


def test_zonder_kaart_zegt_de_stand_dat_de_kaart_ontbreekt(tmp_path, sanne_zonder_kaart):
    wb = load_workbook(schrijf(sanne_zonder_kaart, tmp_path / "a.xlsx"))
    assert "Debiteurenkaart" not in wb.sheetnames
    assert "debiteurenkaart" in wb["Stand"]["A3"].value


def test_een_medewerker_zonder_verschillen_en_zonder_facturen_zonder_order_geeft_een_werkboek(tmp_path, voorbeeld):
    o = overzicht_van(voorbeeld, THIJMEN)
    assert o.verschillen.empty and o.zonder_order.empty
    wb = load_workbook(schrijf(o, tmp_path / "a.xlsx", LEEG))
    assert "Verschillen" in wb.sheetnames and wb["Verschillen"].max_row >= 2
    assert wb["Zonder order"]["A6"].value.startswith("A. Per factuur")


def test_een_bestaande_map_en_een_nieuwe_map_geven_hetzelfde_bestand(tmp_path, overzicht_sanne):
    pad = tmp_path / "nieuw" / "diep" / "a.xlsx"
    schrijf(overzicht_sanne, pad)
    assert pad.exists()


# ---------- eigen invoer van de gebruiker ----------


def test_eigen_keuze_is_leeg_zonder_werkboek(tmp_path):
    assert werkboek.eigen_keuze(tmp_path / "bestaat-niet.xlsx") == {}


def test_eigen_keuze_blijft_staan(tmp_path, voorbeeld, overzicht_sanne):
    pad = schrijf(overzicht_sanne, tmp_path / "a.xlsx")
    wb = load_workbook(pad)
    ws = wb["Verschillen"]
    week, oorzaak, voorstel = ws["B2"].value, ws["C2"].value, ws["M2"].value
    keuze = next(t for t in aansluiten.TOEWIJZINGEN if t != voorstel)
    ws["N2"] = keuze
    wb.save(pad)
    assert werkboek.eigen_keuze(pad) == {(week, oorzaak): keuze}
    # opnieuw maken met die keuze: de gele kolom en de toewijzing houden haar
    opnieuw = overzicht_van(voorbeeld, SANNE, eigen=werkboek.eigen_keuze(pad))
    schrijf(opnieuw, pad)
    ws = load_workbook(pad)["Verschillen"]
    assert (ws["B2"].value, ws["C2"].value, ws["N2"].value, ws["M2"].value) == (week, oorzaak, keuze, voorstel)
    assert ws["N3"].value is None  # alleen de regel die de gebruiker koos
    assert ws["I2"].value == '=IF(N2="",M2,N2)'  # de toewijzing blijft een formule: de eigen keuze telt


def test_een_onleesbaar_overzicht_geeft_een_melding(tmp_path):
    pad = tmp_path / "a.xlsx"
    pad.write_text("dit is geen werkboek")
    with pytest.raises(ValueError, match="niet te lezen"):
        werkboek.eigen_keuze(pad)


def test_een_ongeldige_eigen_keuze_noemt_bestand_blad_en_cel(tmp_path, overzicht_sanne):
    pad = schrijf(overzicht_sanne, tmp_path / "a.xlsx")
    wb = load_workbook(pad)
    wb["Verschillen"]["N3"] = "Iets anders"
    wb.save(pad)
    with pytest.raises(ValueError) as fout:
        werkboek.eigen_keuze(pad)
    melding = str(fout.value)
    assert str(pad) in melding and "Verschillen" in melding and "N3" in melding and "Iets anders" in melding
    assert aansluiten.NOG in melding  # de geldige keuzes staan erbij


def test_een_overzicht_van_een_andere_vorm_geeft_geen_eigen_keuze(tmp_path):
    wb = Workbook()
    wb.active.title = "Verschillen"
    wb.active.append(["a", "b", "c"])
    wb.save(tmp_path / "a.xlsx")
    assert werkboek.eigen_keuze(tmp_path / "a.xlsx") == {}


def test_tekst_wordt_geen_formule(tmp_path, voorbeeld):
    """Order-omschrijving, project, bankomschrijving, opmerking en saldo-toelichting beginnen met een `=`."""
    map_ = tmp_path / "map"
    shutil.copytree(VOORBEELD, map_, ignore=shutil.ignore_patterns("ruw", "verwacht"))
    saldo = inlezen.lees_saldo(VOORBEELD / f"saldo {SANNE}.csv").assign(toelichting="=INJECT saldo")
    dossier.schrijf(map_, saldo={SANNE: saldo})
    d = replace(
        dossier.lees(map_),
        regels=voorbeeld.regels.assign(project="=INJECT project", omschrijving="=INJECT orderregel"),
        bank=voorbeeld.bank.assign(omschrijving="=INJECT bank"),
    )
    opm = opmerkingen(("1", "ja", "=INJECT onderwerp", "=INJECT wat", "=INJECT actie"))
    wb = load_workbook(schrijf(overzicht_van(d, SANNE), tmp_path / "a.xlsx", opm))
    gevonden = {}
    for ws in wb:
        for rij in ws.iter_rows():
            for c in rij:
                if isinstance(c.value, str) and "INJECT" in c.value:
                    assert c.data_type != "f", f"{ws.title}!{c.coordinate} is een formule: {c.value}"
                    gevonden.setdefault(ws.title, c.value)
    assert {"Stand", "Verschillen", "Opmerkingen", "Orders", "Bank", "Orderregels"} <= gevonden.keys()
    assert wb["Bank"]["E2"].data_type == "s"
    assert wb["Bank"]["I2"].data_type == "f"  # de eigen formules in hetzelfde blad blijven formules


@pytest.mark.parametrize("begin", ["=", "+", "-", "@"])
def test_tekst_met_een_formuleteken_blijft_tekst_ook_voor_notities_en_opmerkingen(tmp_path, voorbeeld, begin):
    """Een bankopmerking en een opmerking komen uit de invoer; in Excel mag geen van beide een formule worden."""
    notities = pd.DataFrame(
        [["nabetaling Kim", "Sanne", f"{begin}HYPERLINK(1)"]], columns=["tekst", "betreft", "opmerking"]
    )
    o = overzicht_van(replace(voorbeeld, bank_notities=notities), SANNE)
    opm = opmerkingen(("1", "ja", f"{begin}onderwerp", f"{begin}wat", f"{begin}actie"))
    wb = load_workbook(schrijf(o, tmp_path / "a.xlsx", opm))
    gevonden = [
        c
        for ws in wb
        for rij in ws.iter_rows()
        for c in rij
        if str(c.value).startswith((f"{begin}HYPERLINK", f"{begin}onderwerp", f"{begin}wat", f"{begin}actie"))
    ]  # de formules van het werkboek zelf beginnen ook met een =
    assert {c.parent.title for c in gevonden} >= {"Zonder order", "Bank", "Opmerkingen", "Stand"}
    assert all(c.data_type != "f" for c in gevonden)


STUURTEKENS = "a\x0bb\x00c\x07d\x0ce"  # een zachte regelbreuk van Word, een nulteken, een belteken, een paginawissel
SCHOON = "a bcd e"  # \x0b en \x0c worden een spatie, de rest valt weg


def illegale_tekens(wb) -> list[str]:
    return [
        f"{ws.title}!{c.coordinate}"
        for ws in wb
        for rij in ws.iter_rows()
        for c in rij
        if isinstance(c.value, str) and ILLEGAL_CHARACTERS_RE.search(c.value)
    ]


def test_stuurtekens_in_brontekst_geven_geen_fout_en_worden_opgeschoond(tmp_path, voorbeeld):
    """Een verticale tab (de zachte regelbreuk van Word) of een nulteken in een bankomschrijving, order of opmerking
    liet het schrijven eerder stoppen met een `IllegalCharacterError` van openpyxl."""
    d = replace(
        voorbeeld,
        regels=voorbeeld.regels.assign(project=f"P{STUURTEKENS}", omschrijving=f"O{STUURTEKENS}"),
        bank=voorbeeld.bank.assign(omschrijving=f"B{STUURTEKENS}"),
    )
    opm = opmerkingen(("1", "ja", f"Onderwerp{STUURTEKENS}", f"Wat{STUURTEKENS}", f"Actie{STUURTEKENS}"))
    wb = load_workbook(schrijf(overzicht_van(d, SANNE), tmp_path / "a.xlsx", opm))
    assert not illegale_tekens(wb)
    gevonden = {c.value for ws in wb for rij in ws.iter_rows() for c in rij if isinstance(c.value, str)}
    for kop in ("B", "O", "Onderwerp", "Wat", "Actie"):
        assert kop + SCHOON in gevonden, kop


def test_een_tab_en_een_regeleinde_in_brontekst_blijven_staan(tmp_path, voorbeeld):
    """Dat zijn tekens die Excel wel toelaat; alleen wat Excel weigert wordt aangepast."""
    d = replace(voorbeeld, bank=voorbeeld.bank.assign(omschrijving="eerste\tregel\ntweede"))
    wb = load_workbook(schrijf(overzicht_van(d, SANNE), tmp_path / "a.xlsx"))
    assert wb["Bank"]["E2"].value == "eerste\tregel\ntweede"


def test_het_totaalbestand_schoont_stuurtekens_ook_op(tmp_path, voorbeeld):
    d = replace(voorbeeld, bank=voorbeeld.bank.assign(omschrijving=f"B{STUURTEKENS}"))
    o = overzicht_van(d, SANNE)
    opm = opmerkingen(("1", "ja", f"Onderwerp{STUURTEKENS}", "Wat", "Actie"))
    pad = tmp_path / "totaal.xlsx"
    acties = [(f"Groep{STUURTEKENS}", [(f"Titel{STUURTEKENS}", "Algemeen", f"Toelichting{STUURTEKENS}")])]
    totaal.schrijf([o], opm, acties, pad)
    wb = load_workbook(pad)
    assert not illegale_tekens(wb)
    assert f"Titel{SCHOON}" in {c.value for rij in wb["Actielijst"].iter_rows() for c in rij}


# ---------- de verdiepingsbladen naast het model ----------


def _leeg(x):
    return "" if x is None else x


def test_elk_verdiepingsblad_heeft_evenveel_rijen_en_dezelfde_status_als_het_model(tmp_path, overzicht_sanne):
    o, a = overzicht_sanne, overzicht_sanne.aansluiting
    wb = load_workbook(schrijf(o, tmp_path / "a.xlsx"))
    kolom = werkboek._orderkolommen(o.inst)

    def rijen(blad: str, totaal: bool = False) -> list[tuple]:
        alle = list(wb[blad].iter_rows(min_row=2, values_only=True))
        return alle[:-1] if totaal else alle  # een blad met totaalregel

    per_week = rijen("Per week", totaal=True)
    assert [r[0] for r in per_week] == list(o.weken.week) and [r[14] for r in per_week] == list(o.weken.status)
    per_dag = rijen("Per dag", totaal=True)
    assert [r[3] for r in per_dag] == list(o.per_dag.order) and [r[15] for r in per_dag] == list(o.per_dag.oordeel)
    orders = [r for r in rijen("Orders", totaal=True) if r[0] in set(o.orders.order)]
    assert [r[0] for r in orders] == list(o.orders.order)
    assert [r[kolom.betaling - 1] for r in orders] == list(o.orders.oordeel)
    facturen = rijen("Facturen", totaal=True)
    assert [r[0] for r in facturen] == list(o.facturen.factuur)
    assert [_leeg(r[10]) for r in facturen] == list(o.facturen.opmerking)
    bank = rijen("Bank")
    assert len(bank) == len(o.bank)
    assert [_leeg(r[10]) for r in bank] == list(o.bank.opmerking) and [r[7] for r in bank] == list(o.bank.betreft)
    uren = rijen("Urenregels")
    assert len(uren) == len(a.uren)
    assert sorted(r[6] for r in uren) == sorted(a.uren.status)
    orderregels = rijen("Orderregels")
    assert len(orderregels) == len(a.order_regels)
    assert sorted(r[11] for r in orderregels) == sorted("ja" if t else "nee" for t in a.order_regels.telt_mee)


@pytest.mark.parametrize("blok", ["week", "4 weken", "maand"])
def test_week_en_periode_volgen_de_instellingen(tmp_path, voorbeeld, blok):
    d = replace(voorbeeld, inst=replace(voorbeeld.inst, blok=blok))
    o = overzicht_van(d, SANNE)  # opnieuw aansluiten: de periode van elke dag volgt uit de instellingen
    wb = load_workbook(schrijf(o, tmp_path / "a.xlsx"))
    weken = [r for r in wb["Per week"].iter_rows(min_row=2, values_only=True) if r[0] != "Totaal"]
    assert [r[0] for r in weken] == list(o.weken.week)
    assert all(re.fullmatch(r"\d{4} wk \d{2}", r[0]) for r in weken)  # de weeksleutel blijft tekst
    patroon = {"week": r"\d{4} wk \d{2}", "4 weken": r"\d{4} wk \d{2}-\d{2}", "maand": r"\d{4}-\d{2}"}[blok]
    assert all(re.fullmatch(patroon, r[1]) for r in weken)
    perioden = [r[0] for r in wb["Per periode"].iter_rows(min_row=2, values_only=True) if r[0] != "Totaal"]
    assert perioden == list(o.perioden.periode) and len(set(perioden)) == len(perioden)


# ---------- instellingen ----------


def _orders_oordeel(o: overzicht.Overzicht, order: str, pad: Path) -> str:
    ws = load_workbook(schrijf(o, pad))["Orders"]
    kolom = werkboek._orderkolommen(o.inst).betaling
    return next(r[kolom - 1] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == order)


def test_de_betaaltermijn_uit_de_instellingen_bepaalt_wanneer_een_order_niet_ontvangen_is(tmp_path, voorbeeld):
    """Een order die 37 dagen voor het einde van de bank is gefactureerd en er niet in voorkomt."""
    order = "I01250491"
    voorbeeld = replace(voorbeeld, kaart=None)  # met de kaart van de voorbeeldset zijn de facturen afgeletterd
    a = overzicht_van(voorbeeld, SANNE).aansluiting
    zonder = zonder_betalingen(a)
    assert zonder.bank_tot - a.kop.set_index("order").factuurdatum[order] == pd.Timedelta(days=37)
    kort = overzicht.maak(replace(zonder, inst=replace(a.inst, betaaltermijn=30)), voorbeeld)
    lang = overzicht.maak(replace(zonder, inst=replace(a.inst, betaaltermijn=45)), voorbeeld)
    assert _orders_oordeel(kort, order, tmp_path / "k.xlsx") == "Niet ontvangen"
    assert _orders_oordeel(lang, order, tmp_path / "l.xlsx").startswith("Nog niet vervallen bij einde bankexport")


def test_het_btw_tarief_uit_de_instellingen_komt_in_het_blad_bank(tmp_path, overzicht_sanne, voorbeeld):
    for btw, tekst in ((0.21, "21%"), (0.09, "9%"), (0.065, "6,5%")):
        a = replace(overzicht_sanne.aansluiting, inst=replace(voorbeeld.inst, btw=btw))
        ws = load_workbook(schrijf(overzicht.maak(a, voorbeeld), tmp_path / "a.xlsx"))["Bank"]
        assert ws["J1"].value == f"Ex btw bij {tekst}"
        assert ws["J2"].value == f"=ROUND(I2/{round(1 + btw, 6)},2)"


def test_de_koppen_en_uitleg_noemen_het_blok_uit_de_instellingen(tmp_path, voorbeeld, overzicht_sanne):
    for blok, kop, uitleg in (
        ("4 weken", "Periode (4 weken)", "Eén rij per blok van 4 weken"),
        ("week", "Periode (week)", "Eén rij per blok van een week"),
        ("maand", "Periode (maand)", "Eén rij per blok van een maand"),
    ):
        a = replace(overzicht_sanne.aansluiting, inst=replace(voorbeeld.inst, blok=blok))
        wb = load_workbook(schrijf(overzicht.maak(a, voorbeeld), tmp_path / "a.xlsx"))
        assert (wb["Per periode"]["A1"].value, wb["Per week"]["B1"].value, wb["Per dag"]["C1"].value) == (kop,) * 3
        assert wb["Verschillen"]["A1"].value == kop
        tekst = " ".join(str(c.value) for rij in wb["Stand"].iter_rows() for c in rij if c.value)
        assert uitleg in tekst


def test_de_namen_uit_de_instellingen_staan_in_de_teksten(tmp_path, voorbeeld, overzicht_sanne):
    inst = replace(voorbeeld.inst, bureau="Studio Noord", opdrachtgever="Oeverland")
    o = overzicht.maak(replace(overzicht_sanne.aansluiting, inst=inst), voorbeeld)
    wb = load_workbook(schrijf(o, tmp_path / "a.xlsx", LEEG))
    for blad in ("Stand", "Zonder order"):
        tekst = " ".join(str(c.value) for rij in wb[blad].iter_rows() for c in rij if c.value)
        assert "Studio Noord" in tekst and "Oeverland" in tekst  # de namen uit de instellingen
        assert "het bureau" not in tekst and "de opdrachtgever" not in tekst.lower()
        assert not any(w in tekst.lower() for w in ("inleenorder", "timesheet"))


def test_opmerkingen_staan_op_een_verborgen_blad_en_de_belangrijke_ook_op_stand(tmp_path, overzicht_sanne):
    wb = load_workbook(schrijf(overzicht_sanne, tmp_path / "a.xlsx"))
    onderwerpen = [r[1] for r in wb["Opmerkingen"].iter_rows(min_row=2, values_only=True)]
    assert onderwerpen == ["Eerste punt", "Tweede punt"]
    stand = [c.value for rij in wb["Stand"].iter_rows() for c in rij if c.value]
    assert "1. Eerste punt" in stand and "2. Tweede punt" not in stand
    assert any(str(v).startswith("Er is iets aan de hand. Te doen: Bel het bureau.") for v in stand)


# ---------- formules en waarden ----------


def waarde_naast(rijen: list[list], label: str, kolom: int = 1):
    """De waarde in `kolom` van de eerste rij waarvan kolom A met `label` begint."""
    return next(r[kolom] for r in rijen if isinstance(r[0], str) and r[0].startswith(label))


def test_formules_geven_de_cijfers_van_het_model(tmp_path, overzicht_sanne):
    pytest.importorskip("formulas")
    o = overzicht_sanne
    bladen = doorgerekend(schrijf(o, tmp_path / "a.xlsx"))
    cellen = [(naam, c) for naam, rijen in bladen.items() for r in rijen for c in r]
    assert not [c for _, c in cellen if isinstance(c, str) and c.startswith("#")]
    assert not [c for _, c in cellen if isinstance(c, str) and "sluit niet aan" in c.lower()]
    stand = bladen["STAND"]
    assert waarde_naast(stand, "Gewerkt volgens") == pytest.approx(o.stand["urenstaat"])
    assert waarde_naast(stand, "Door het bureau gefactureerd") == pytest.approx(o.stand["gefactureerd"])
    assert waarde_naast(stand, "Door de opdrachtgever op orders gezet") == pytest.approx(o.stand["op_orders"])
    assert waarde_naast(stand, "Door de opdrachtgever betaald") == pytest.approx(o.stand["betaald"])
    assert any(isinstance(r[3], str) and r[3].startswith("Controle: gelijk aan het verschil") for r in stand)
    totaal = next(r for r in bladen["VERSCHILLEN"] if r[0] == "Totaal")
    assert totaal[6] == pytest.approx(round(o.verschillen.verschil.sum(), 2))
    for oorzaak in o.per_oorzaak().itertuples():
        assert waarde_naast(stand, oorzaak.oorzaak) == pytest.approx(oorzaak.bedrag)
    wie = o.per_toewijzing().set_index("toewijzing")
    for naam, rij in wie.iterrows():
        assert waarde_naast(stand, naam) == pytest.approx(rij.bedrag)
    zonder = bladen["ZONDER ORDER"]
    assert waarde_naast(zonder, "Totaal", 8) == pytest.approx(round(o.zonder_order.zonder_order.sum(), 2))
    assert any(isinstance(r[10], str) and r[10].startswith("Controle: gelijk aan kolom I") for r in zonder)


def test_wie_is_aan_zet_op_stand_verandert_mee_met_een_eigen_keuze(tmp_path, overzicht_sanne):
    pytest.importorskip("formulas")
    pad = schrijf(overzicht_sanne, tmp_path / "a.xlsx")
    voor = waarde_naast(doorgerekend(pad)["STAND"], aansluiten.NOG)
    wb = load_workbook(pad)
    ws = wb["Verschillen"]
    assert ws["M2"].value != aansluiten.NOG
    ws["N2"] = aansluiten.NOG
    wb.save(pad)
    na = waarde_naast(doorgerekend(pad)["STAND"], aansluiten.NOG)
    assert na == pytest.approx(voor + abs(overzicht_sanne.verschillen.verschil.iloc[0]))


def test_formules_op_de_debiteurenkaart_geven_de_cijfers_van_het_model(tmp_path, met_kaart):
    pytest.importorskip("formulas")
    k = met_kaart.kaart
    bladen = doorgerekend(schrijf(met_kaart, tmp_path / "a.xlsx", LEEG))
    kaart = bladen["DEBITEURENKAART"]
    assert not [c for r in kaart for c in r if isinstance(c, str) and c.startswith("#")]
    assert waarde_naast(kaart, "Echt open na afletteren", 3) == pytest.approx(k.echt_open)
    assert waarde_naast(kaart, "Ter controle", 3) == pytest.approx(k.controle)
    assert waarde_naast(kaart, "Ter controle", 9) == "Sluit aan"
    assert waarde_naast(kaart, "Saldo van de kaart", 3) == pytest.approx(k.heel["saldo"])
    for label, bedrag, _ in k.regels:
        assert waarde_naast(kaart, label, 3) == pytest.approx(bedrag)
    assert not [
        c
        for naam, rijen in bladen.items()
        for r in rijen
        for c in r
        if isinstance(c, str) and "sluit niet aan" in c.lower()
    ]


def test_de_waarschuwing_op_de_debiteurenkaart_staat_in_het_rood(tmp_path, met_kaart):
    ws = load_workbook(schrijf(met_kaart, tmp_path / "a.xlsx", LEEG))["Debiteurenkaart"]
    assert ws["A3"].value.startswith("Let op: niet-afgeletterde ontvangsten")
    assert ws["A3"].font.bold and ws["A3"].font.color.rgb.endswith("C00000")
    assert ws["A4"].value is None and not ws["A4"].font.bold  # de lege regel eronder is gewoon leeg gebleven


def test_de_controle_op_de_kaart_volgt_de_tolerantie_van_het_model_en_hangt_niet_van_de_taal_af(
    tmp_path, met_kaart, monkeypatch
):
    def formule(o) -> str:
        ws = load_workbook(schrijf(o, tmp_path / "t.xlsx", LEEG))["Debiteurenkaart"]
        return next(str(c.value) for rij in ws.iter_rows() for c in rij if "sluit niet aan, verschil" in str(c.value))

    f = formule(met_kaart)
    assert f"<{overzicht.CONTROLE_VERSCHIL}," in f
    assert "TEXT(" not in f and "FIXED(" in f  # TEXT(x,"0.00") hangt af van de taal van Excel
    monkeypatch.setattr(overzicht, "CONTROLE_VERSCHIL", 0.5)
    assert "<0.5," in formule(met_kaart)


def test_een_kaart_die_niet_sluit_zegt_het_met_het_verschil_en_een_die_sluit_zegt_sluit_aan(tmp_path, met_kaart):
    """De controlecel is een formule; doorgerekend staat er "LET OP: sluit niet aan, verschil ..." of "Sluit aan"."""
    pytest.importorskip("formulas")
    kaart = [("2025-03-24", "factuur", "F-102", "", 1000.00), *KAART_B[1:]]  # 548,80 te laag
    scheef = overzicht_van(kaart_dossier(tmp_path / "scheef", kaart=kaart), "Femke Jansma")
    controle = scheef.controles()[overzicht.C_KAART]
    assert controle.sluit is False and controle.verschil == pytest.approx(-548.80)
    cel = waarde_naast(doorgerekend(schrijf(scheef, tmp_path / "a.xlsx", LEEG))["DEBITEURENKAART"], "Ter controle", 9)
    assert cel.startswith("LET OP: sluit niet aan, verschil ") and float(cel.rsplit(" ", 1)[1]) == pytest.approx(
        -548.80
    )
    assert met_kaart.controles()[overzicht.C_KAART].sluit is True
    sluit = waarde_naast(
        doorgerekend(schrijf(met_kaart, tmp_path / "b.xlsx", LEEG))["DEBITEURENKAART"], "Ter controle", 9
    )
    assert sluit == "Sluit aan"


@pytest.fixture(scope="module")
def met_toezegging(voorbeeld) -> overzicht.Overzicht:
    t = toezegging(
        ("T-001", "2025 wk 33", "", "Sanne week 33", "Uren", 32.0, 37.31, 1193.92),  # geen order voor deze week
        ("T-001", "2025 wk 29", "", "Sanne week 29", "Uren", 32.0, 37.31, 1193.92),  # staat ook op een order
        ("T-001", "2025 wk 34", "", "Sanne week 34", "Uren", None, 37.31, None),  # zonder bedrag
        ("T-001", "", "2025-08", "Reiskosten augustus - 4 dagen", "Kilometers", 80.0, 0.19, 15.2),
    )
    return overzicht_van(replace(voorbeeld, toezegging=t), SANNE)


def test_met_een_toezegging_heeft_zonder_order_blok_d_en_de_kolommen_p_en_q(tmp_path, met_toezegging):
    pytest.importorskip("formulas")
    o = met_toezegging
    pad = schrijf(o, tmp_path / "a.xlsx", LEEG)
    ws = load_workbook(pad)["Zonder order"]
    assert (ws["P6"].value, ws["Q6"].value) == (
        "Waarvan uren op de toezegging",
        "Niet op de toezegging: uren zonder enig stuk, en onkosten",
    )
    kop = next(r for r in ws.iter_rows() if str(r[0].value).startswith("D. Toezegging T-001 van 10-11-2025"))
    soorten = [ws.cell(kop[0].row + i, 3).value for i in range(1, 5)]
    assert soorten == [overzicht.T_ALLEEN, overzicht.T_DUBBEL, overzicht.T_LEEG, overzicht.T_REIS]
    stand = [c.value for rij in load_workbook(pad)["Stand"].iter_rows() for c in rij if c.value]
    assert f"Van '{aansluiten.OW}' staat al op toezegging T-001" in stand
    bladen = doorgerekend(pad)
    assert not [c for r in bladen["ZONDER ORDER"] for c in r if isinstance(c, str) and c.startswith("#")]
    p = waarde_naast(bladen["ZONDER ORDER"], "Totaal", 15)
    assert p == pytest.approx(round(o.zonder_order.op_toezegging.sum(), 2))
    for soort in (overzicht.T_ALLEEN, overzicht.T_DUBBEL, overzicht.T_REIS):
        label = f"Waarvan: {soort[0].lower()}{soort[1:]}"
        verwacht = o.toezegging.bedrag[o.toezegging.soort == soort].sum()
        assert waarde_naast(bladen["ZONDER ORDER"], label, 8) == pytest.approx(verwacht)
    assert waarde_naast(bladen["STAND"], f"Van '{aansluiten.OW}'") == pytest.approx(
        o.verschillen.verschil[o.verschillen.oorzaak == overzicht.O_TOEZEGGING].sum()
    )


def test_met_een_toezegging_als_waarden_schrijft_geen_formules(met_toezegging):
    wb = Workbook()
    assert geen_formules(werkboek.zonder_order(wb, met_toezegging, als_waarden=True)) == []
    assert geen_formules(werkboek.stand(wb, met_toezegging, LEEG, als_waarden=True)) == []


def geen_formules(ws) -> list[str]:
    return [c.coordinate for rij in ws.iter_rows() for c in rij if c.data_type == "f"]


def test_als_waarden_schrijft_geen_formules(tmp_path, overzicht_sanne, met_kaart):
    for o in (overzicht_sanne, met_kaart):
        wb = Workbook()
        s = werkboek.stand(wb, o, OPM, titel="Stand Sanne", als_waarden=True)
        z = werkboek.zonder_order(wb, o, titel="Zonder order Sanne", als_waarden=True)
        assert (s.title, z.title) == ("Stand Sanne", "Zonder order Sanne")
        assert geen_formules(s) == [] and geen_formules(z) == []
        wb.save(tmp_path / "a.xlsx")  # een blad los, zonder de bladen waar formules naar zouden verwijzen
    # en zonder die vlag zijn het wel formules
    wb = Workbook()
    assert geen_formules(werkboek.stand(wb, overzicht_sanne, OPM)) != []
    assert geen_formules(werkboek.zonder_order(wb, overzicht_sanne)) != []


def test_als_waarden_heeft_dezelfde_indeling_en_dezelfde_getallen_als_de_formules(tmp_path, overzicht_sanne):
    pytest.importorskip("formulas")
    wb = Workbook()
    werkboek.stand(wb, overzicht_sanne, OPM, als_waarden=True)
    werkboek.zonder_order(wb, overzicht_sanne, als_waarden=True)
    del wb["Sheet"]
    wb.save(tmp_path / "waarden.xlsx")
    bladen = doorgerekend(schrijf(overzicht_sanne, tmp_path / "formules.xlsx"))
    waarden = {
        ws.title.upper(): [list(r) for r in ws.iter_rows(values_only=True)]
        for ws in load_workbook(tmp_path / "waarden.xlsx")
    }
    for blad in ("STAND", "ZONDER ORDER"):
        assert len(waarden[blad]) == len(bladen[blad])
        for rij_w, rij_f in zip(waarden[blad], bladen[blad], strict=True):
            for w, f in itertools.zip_longest(rij_w, rij_f):
                w, f = (None if x == "" else x for x in (w, f))  # een lege tekst is een lege cel
                if isinstance(w, datetime.date):
                    w = (pd.Timestamp(w) - pd.Timestamp("1899-12-30")).days
                assert soort_waarde(w) == soort_waarde(f), (blad, rij_w[0], w, f)
                if isinstance(w, int | float):
                    assert w == pytest.approx(f, abs=0.005), (blad, rij_w[0])
                else:
                    assert w == f, (blad, rij_w[0])


def soort_waarde(w) -> str:
    """De soort cel: `formulas` geeft een datum als dagnummer terug, dus een datum telt als getal."""
    if isinstance(w, datetime.date):
        return "getal"
    return "getal" if isinstance(w, int | float) and not isinstance(w, bool) else type(w).__name__


# ---------- de opdrachtregel ----------


def test_aansluiten_schrijft_een_overzicht_en_geeft_de_eigen_keuze_door(tmp_path, capsys):
    kopie = tmp_path / "map"
    shutil.copytree(VOORBEELD, kopie, ignore=shutil.ignore_patterns("ruw", "verwacht"))
    uit = tmp_path / "uit"
    assert main(["aansluiten", str(kopie), "--uit", str(uit), "--medewerker", SANNE]) == 0
    pad = uit / f"Overzicht {SANNE}.xlsx"
    assert pad.exists() and not (uit / f"Aansluiting {SANNE}.xlsx").exists()
    wb = load_workbook(pad)
    wb["Verschillen"]["N2"] = aansluiten.NOG
    wb.save(pad)
    assert main(["aansluiten", str(kopie), "--uit", str(uit), "--medewerker", SANNE]) == 0
    assert load_workbook(pad)["Verschillen"]["N2"].value == aansluiten.NOG


def test_een_medewerker_zonder_orders_geeft_een_werkboek(tmp_path, voorbeeld):
    eigen = ["I01250324", "I02250348"]  # alleen orders van de andere medewerker
    d = replace(
        voorbeeld,
        kop=voorbeeld.kop[voorbeeld.kop.order.isin(eigen)],
        regels=voorbeeld.regels[voorbeeld.regels.order.isin(eigen)],
        vervallen=None,
    )
    o = overzicht_van(d, SANNE)
    assert o.orders.empty
    wb = load_workbook(schrijf(o, tmp_path / "a.xlsx", LEEG))
    assert wb["Orders"].max_row == 2  # kopregel en totaalregel
    assert "Orders" in wb.sheetnames and "Stand" in wb.sheetnames


def test_lege_blokken_geven_geen_cirkelverwijzing_of_foutwaarde(tmp_path, voorbeeld):
    """Zonder verschillen en zonder facturen zonder order, en een kaart zonder ontvangsten: een som over nul rijen is
    een nul, geen bereik dat de somrij zelf meetelt."""
    pytest.importorskip("formulas")
    leeg = overzicht_van(voorbeeld, THIJMEN)
    assert leeg.verschillen.empty and leeg.zonder_order.empty
    d = kaart_dossier(tmp_path / "d", kaart=[("2025-03-31", "factuur", "F-103", "", 1936.00)])
    zonder_ontvangsten = overzicht_van(d, "Femke Jansma")
    assert zonder_ontvangsten.kaart.ontvangsten.empty
    for naam, o in (("leeg", leeg), ("kaart", zonder_ontvangsten)):
        pad = schrijf(o, tmp_path / f"{naam}.xlsx", LEEG)
        wb = load_workbook(pad)
        for ws in wb:
            for rij in ws.iter_rows():
                for c in rij:
                    m = c.data_type == "f" and re.fullmatch(r"=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", c.value)
                    if m:  # een bereik van achter naar voren telt de somrij zelf mee
                        assert int(m.group(4)) >= int(m.group(2)), f"{ws.title}!{c.coordinate}: {c.value}"
        bladen = doorgerekend(pad)
        assert not [c for r in bladen["STAND"] for c in r if isinstance(c, str) and c.startswith("#")]
    kaart = doorgerekend(tmp_path / "kaart.xlsx")["DEBITEURENKAART"]
    assert waarde_naast(kaart, "Echt open na afletteren", 3) == pytest.approx(zonder_ontvangsten.kaart.echt_open)
    assert waarde_naast(kaart, "Ter controle", 9) == "Sluit aan"
