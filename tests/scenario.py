"""Bouwt verzonnen dossiers in sjabloonvorm voor de scenario-tests.

`bouw` schrijft een map met alles wat `dossier.lees` nodig heeft. Zonder afwijkingen sluit alles aan; elke afwijking
verandert één ding. De bouwer geeft naast de map terug wat hij veranderde (`Scenario`), zodat een test het verwachte
bedrag kent zonder de tool na te rekenen.

Alles is verzonnen en elke bouw geeft dezelfde bestanden: de dagen liggen in een vast jaar (2029 begint op een
maandag, dus week 1 is 1 t/m 7 januari) en er is geen willekeur.

De wereld van de bouwer, zonder afwijkingen:

- Elke medewerker werkt elke week van maandag t/m donderdag 8 uur tegen een vast uurtarief (per medewerker anders).
- Het bureau factureert per medewerker per blok één factuur, de dag na de laatste werkdag van het blok.
- De opdrachtgever zet per medewerker per blok één order met dezelfde regels, vier dagen na de laatste werkdag.
- Elke order wordt in één keer betaald, drie dagen plus de plaats van de medewerker in de lijst na de orderdatum,
  per rekening volgens het aandeel van die rekening. Zo staat een betaling altijd vóór de volgende order en liggen
  de betalingen van twee medewerkers niet op dezelfde dag.
- De boekhouding heeft alles afgeletterd: de debiteurenkaart is een bestand met alleen een kopregel. Dat is voor
  `dossier.lees` een lege kaart (niets open) en geeft dus een kaartblok; `kaart=False` laat het bestand weg. Wat een
  afwijking op de kaart verandert, staat bij die afwijking.

Wat de afwijkingen doen (`AFWIJKINGEN`). Elke afwijking krijgt een eigen week; de afwijkingen met dezelfde
medewerker liggen vijf weken uit elkaar, zodat ze in geen enkele blokgrootte hetzelfde blok raken en alle te
combineren zijn. Het `bedrag` van `Afwijking` is ex btw en heeft het teken van factuur min order, tenzij anders
vermeld.

- `blok_zonder_order`: de order van een heel blok ontbreekt. Het bedrag is wat het bureau over dat blok
  factureerde. De facturen van het blok staan open op de kaart.
- `dag_zonder_order`: één dag (een donderdag) staat niet op de order van zijn blok. Het deel dat de order niet
  dekt staat open op de kaart.
- `order_niet_betaald`: de opdrachtgever betaalt één order niet. Het bedrag is de order (ex btw); de facturen
  staan open.
- `niet_gefactureerd`: één dag (een donderdag) staat op de order en is betaald, maar het bureau factureerde hem
  niet. Het bedrag is negatief. Het teveel betaalde is bij het afletteren weggeboekt: de kaart toont niets.
- `ander_tarief`: een week staat op de order tegen een lager tarief dan gefactureerd. Het deel dat de order niet
  dekt staat open op de kaart.
- `dagvergoeding_niet_op_order`: het bureau factureert twee weken lang een vaste vergoeding per dag die op geen
  order staat. Het deel dat de order niet dekt staat open op de kaart.
- `dubbel_gefactureerd`: één dag staat op een tweede factuur. De goedgekeurde urenstaat van die week is het bedrag
  zonder die tweede factuur (`urenstaat-afwijkend.csv`). De tweede factuur staat open op de kaart.
- `creditorder_later_verrekend`: een creditorder (een tariefcorrectie voor één dag van de eerste medewerker) wordt
  verrekend met de betaling voor de order van de volgende medewerker in hetzelfde blok. Het bedrag is het verschil
  dat de creditorder maakt: factuur min orders, positief. De factuur van de betaler staat open en de twee
  ontvangsten staan ongelet op de kaart. Vraagt minstens twee medewerkers.
- `vervangen_order`: een order is door een andere vervangen zonder gecrediteerd te zijn (`vervallen-orders.csv`);
  alleen de nieuwe wordt betaald. Het bedrag is de vervangen order (ex btw): dat telt nergens mee.
- `dubbele_order`: dezelfde order komt twee keer binnen. Een sjabloon kan dat niet uitdrukken (een order twee keer
  in `orders.csv` telt als dubbele regels), een pdf wel; `Scenario.lees` zet de kopie er daarom in het ingelezen
  dossier bij, zoals een pdf onder een tweede naam. Het bedrag is 0.
- `overlappende_bankexport`: de tweede helft van de betalingen staat in twee exporten. Het bedrag is 0; `aantal`
  van de afwijking is het aantal dubbele boekingen.
- `restant_op_kaart`: een week staat op de order tegen een hoger tarief dan gefactureerd en de order is helemaal
  betaald. Wat te veel is betaald staat als ontvangst op de kaart (op de gewone rekening, of de enige). Het bedrag
  is negatief.
- `meer_betaald`: hetzelfde, maar bij het afletteren is het teveel weggeboekt: de kaart toont niets.

Twee opties veranderen de wereld voor een scenario waar het om de bank en de kaart draait. `gelijk_betaald`: de
opdrachtgever betaalt de orders van alle medewerkers in een blok op dezelfde dag (anders drie dagen plus de plaats
van de medewerker). `kaart_ongelet`: de boekhouding heeft de betalingen nog niet afgeletterd, dus de facturen van
betaalde orders staan helemaal open en de ontvangsten staan los op de kaart. Met `tarieven` krijgen medewerkers
hetzelfde tarief en dus even grote orders, en `betaaltermijn` is die van de instellingen.

Een medewerker in `zonder_orders` staat op geen enkele order: de opdrachtgever betaalt niets en de facturen staan
open.
"""

import datetime
import re
import tempfile
from dataclasses import dataclass, field, replace
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from weekstaat import aansluiten, bevindingen, dossier, inlezen, overzicht, totaal
from weekstaat.cli import main

JAAR = 2029
EERSTE_MAANDAG = pd.Timestamp(JAAR, 1, 1)
UREN_PER_DAG = 8.0
DAGEN_PER_WEEK = 4
TARIEVEN = (41.20, 38.60, 43.10, 39.90)
DAGVERGOEDING = 9.50
IBAN_G, IBAN_GEWOON = "NL00RABO0991000001", "NL00RABO0123456789"
ENTITEITEN = (("OI", "Oeverland Infra"), ("OIAM", "Oeverland Infra Asset Management"))
BLOKKEN = ("week", "4 weken", "maand")

AFWIJKINGEN = ("blok_zonder_order", "dag_zonder_order", "order_niet_betaald", "niet_gefactureerd", "ander_tarief",
               "dagvergoeding_niet_op_order", "dubbel_gefactureerd", "creditorder_later_verrekend", "vervangen_order",
               "dubbele_order", "overlappende_bankexport", "restant_op_kaart", "meer_betaald")  # fmt: skip

# Bij welke medewerker (de plaats in de lijst; bij te weinig medewerkers de laatste) een afwijking hoort.
_VAN = {
    "blok_zonder_order": 0, "dag_zonder_order": 0, "order_niet_betaald": 0, "niet_gefactureerd": 0,
    "creditorder_later_verrekend": 0, "ander_tarief": 1, "dagvergoeding_niet_op_order": 1, "dubbel_gefactureerd": 1,
    "vervangen_order": 2, "dubbele_order": 2, "restant_op_kaart": 2, "meer_betaald": 2,
}  # fmt: skip
WEKEN_UIT_ELKAAR = 5
# Het tariefverschil van de order ten opzichte van de factuur, per uur.
TARIEFVERSCHIL = {"ander_tarief": -1.50, "meer_betaald": 2.00, "restant_op_kaart": 3.00}
CREDIT_TARIEF = -2.00


@dataclass(frozen=True)
class Afwijking:
    """Wat één afwijking veranderde. `bedrag` is ex btw en heeft het teken van factuur min order (zie de docstring van
    de module voor de uitzonderingen); `weken` zijn de weken waar het om gaat (weeksleutels als '2029 wk 03'),
    `periode` het blok, `order` het ordernummer waar het om gaat (leeg als er geen is), `aantal` een telling waar die
    iets zegt, `extra` een tweede ordernummer (bij `vervangen_order` is `order` de vervangen order en `extra` de order
    die hem vervangt) en `facturen` de facturen van het bureau waar de afwijking op zit."""

    naam: str
    medewerker: str
    weken: tuple[str, ...]
    periode: str
    order: str
    bedrag: float
    aantal: int = 0
    extra: str = ""
    facturen: tuple[str, ...] = ()

    def incl(self, btw: float) -> float:
        return round(self.bedrag * (1 + btw), 2)


@dataclass
class Scenario:
    """De uitkomst van `bouw`: de map en wat er in zit.

    `afwijkingen` heeft elke gevraagde afwijking. `orders` geeft per (medewerker, blok) het ordernummer van de gewone
    order. `perioden` zijn de blokken van de eerste tot de laatste werkdag. `verwacht` heeft per medewerker de vier
    bedragen van de Stand (ex btw) zoals de bouwer ze kent: `urenstaat`, `gefactureerd`, `op_orders` en `betaald`.
    `incl` heeft per ordernummer (ook de creditorder) het ordertotaal incl. btw."""

    map: Path
    medewerkers: tuple[str, ...]
    btw: float
    blok: str
    perioden: list[str]
    afwijkingen: dict[str, Afwijking]
    orders: dict[tuple[str, str], str]
    verwacht: dict[str, dict[str, float]]
    bank: bool = True
    kaart: bool = True
    zonder_orders: tuple[str, ...] = ()
    incl: dict[str, float] = field(default_factory=dict)  # per ordernummer het bedrag incl. btw
    dubbele: list[str] = field(default_factory=list)  # ordernummers die `lees` er een tweede keer bij zet

    def lees(self) -> dossier.Dossier:
        """Leest het dossier. Een dubbele order komt er hier als kopie bij (zoals een pdf onder een tweede naam)."""
        d = dossier.lees(self.map)
        for nr in self.dubbele:
            kop = d.kop[d.kop.order == nr]
            regels = d.regels[d.regels.order == nr]
            naam = f"kopie van {kop.bestand.iloc[0]}"
            d = replace(
                d,
                kop=pd.concat([d.kop, kop.assign(bestand=naam)], ignore_index=True),
                regels=pd.concat([d.regels, regels.assign(bestand=naam)], ignore_index=True),
            )
        return d


# ---------- kalender ----------


def _dag(week: int, dag: int) -> pd.Timestamp:
    """De datum van een dag (0 = maandag) in week `week` van het jaar (1-based)."""
    return EERSTE_MAANDAG + pd.Timedelta(days=7 * (week - 1) + dag)


def _week(datum: pd.Timestamp) -> str:
    iso = datum.isocalendar()
    return f"{iso.year} wk {iso.week:02d}"


def _blok(datum: pd.Timestamp, blok: str) -> str:
    """Het blok waarin de datum valt, met de naam die de tool eraan geeft."""
    if blok == "week":
        return _week(datum)
    if blok == "maand":
        return f"{datum.year}-{datum.month:02d}"
    iso = datum.isocalendar()
    eerste = (iso.week - 1) // 4 * 4 + 1
    return f"{iso.year} wk {eerste:02d}-{eerste + 3:02d}"


def _plan(afwijkingen: tuple[str, ...], aantal: int, weken: int) -> dict[str, tuple[int, int]]:
    """Per afwijking de medewerker (plaats in de lijst) en de week waarin hij speelt."""
    rang = [0] * aantal
    plan = {}
    for naam in (x for x in AFWIJKINGEN if x in afwijkingen and x != "overlappende_bankexport"):
        m = min(_VAN[naam], aantal - 1)
        if naam == "creditorder_later_verrekend":
            if aantal < 2:
                raise ValueError("creditorder_later_verrekend vraagt minstens twee medewerkers")
            j = max(rang[m], rang[m + 1])  # de betaler houdt in dat blok geen afwijking
            rang[m] = rang[m + 1] = j + 1
        else:
            j, rang[m] = rang[m], rang[m] + 1
        week = 1 + WEKEN_UIT_ELKAAR * j
        nodig = week + (1 if naam == "dagvergoeding_niet_op_order" else 0)
        if nodig > weken:
            raise ValueError(f"{naam} komt in week {nodig} en het dossier heeft {weken} weken: vraag meer weken")
        plan[naam] = (m, week)
    return plan


# ---------- het dossier ----------


@dataclass
class _Order:
    nr: str
    m: int
    blok: str
    datum: pd.Timestamp
    entiteit: str
    regels: list[dict]
    betaald: bool = True

    @property
    def excl(self) -> float:
        return round(sum(r["bedrag"] for r in self.regels), 2)


def _regel_uren(datum: pd.Timestamp, tarief: float) -> dict:
    return dict(datum=datum, aantal=UREN_PER_DAG, tarief=tarief, bedrag=round(UREN_PER_DAG * tarief, 2),
                omschrijving="Tekenaar", eenheid="Uren")  # fmt: skip


def _factuurregel(datum: pd.Timestamp, tarief: float, groep) -> dict:
    return dict(datum=datum, uren=UREN_PER_DAG, tarief=tarief, bedrag=round(UREN_PER_DAG * tarief, 2),
                soort="Normale Uren", groep=groep)  # fmt: skip


def _incl(excl: float, btw: float) -> float:
    return round(excl * (1 + btw), 2)


def _deel(incl: float, aandeel: float) -> float:
    """Het deel van een ordertotaal dat een rekening ontvangt, op de cent afgerond zoals de tool het uitrekent
    (`bank.koppel` rondt met pandas af, en dat rondt een half bedrag niet altijd naar boven)."""
    return float(pd.Series([incl * aandeel]).round(2).iloc[0])


def _rekeningen(aantal: int) -> list[tuple[str, str, float]]:
    """De rekeningen (naam, iban, aandeel) van de instellingen: met één rekening is dat de gewone met aandeel 1."""
    if aantal == 2:
        return [("G-rekening", IBAN_G, 0.30), ("Gewone rekening", IBAN_GEWOON, 0.70)]
    return [("Gewone rekening", IBAN_GEWOON, 1.0)]


def _instellingen(entiteiten: int, rekeningen: int, btw: float, blok: str, betaaltermijn: int) -> str:
    tekst = f'bureau = "Studio Noord"\nopdrachtgever = "Oeverland"\nbtw = {btw}\nblok = "{blok}"\n'
    tekst += f"betaaltermijn = {betaaltermijn}\n\n"
    for code, naam in ENTITEITEN[:entiteiten]:
        tekst += f'[[entiteit]]\ncode = "{code}"\nnaam = "{naam}"\n\n'
    for naam, iban, aandeel in _rekeningen(rekeningen):
        tekst += f'[[rekening]]\nnaam = "{naam}"\niban = "{iban}"\naandeel = {aandeel}\n\n'
    return tekst


def bouw(
    map_: Path,
    *,
    medewerkers: tuple[str, ...] = ("Sanne Bakker", "Jeroen Hendriks"),
    weken: int = 8,
    entiteiten: int = 2,
    rekeningen: int = 2,
    bank: bool = True,
    kaart: bool = True,
    btw: float = 0.21,
    blok: str = "4 weken",
    orders_als: str = "sjabloon",
    sjabloon: str = "csv",
    betaaltermijn: int = 30,
    tarieven: tuple[float, ...] = TARIEVEN,
    gelijk_betaald: bool = False,
    kaart_ongelet: bool = False,
    zonder_orders: tuple[str, ...] = (),
    afwijkingen: tuple[str, ...] = (),
) -> Scenario:
    """Schrijft een verzonnen dossier in sjabloonvorm met `dossier.schrijf` en geeft een `Scenario` terug met de map en
    wat er veranderd is. Zonder afwijkingen sluit alles aan (zie de docstring van de module).

    `weken` is het aantal weken vanaf week 1 van 2029. `entiteiten` (1 of 2) en `rekeningen` (1 of 2) zijn het aantal
    in de instellingen; medewerker i staat op entiteit i modulo het aantal. `bank=False` laat `bank.csv` weg,
    `kaart=False` laat `debiteurenkaart.csv` weg. `orders_als` kan alleen "sjabloon" zijn: de pdf-lezer heeft eigen
    tests. `sjabloon="xlsx"` schrijft alle sjablonen als Excel. `betaaltermijn` is die van de instellingen en
    `tarieven` zijn de uurtarieven per medewerker (zelfde tarief geeft even grote orders). `zonder_orders` zijn
    medewerkers die op geen enkele order staan.

    Met `gelijk_betaald` betaalt de opdrachtgever de orders van alle medewerkers in een blok op dezelfde dag. Met
    `kaart_ongelet` heeft de boekhouding de betalingen nog niet afgeletterd: de facturen van betaalde orders staan
    helemaal open en de ontvangsten staan los op de kaart.

    Een afwijking die niet in `AFWIJKINGEN` staat, of die meer weken of medewerkers vraagt dan er zijn, is een
    `ValueError`."""
    if orders_als != "sjabloon":
        raise ValueError(f"orders_als is 'sjabloon'; de pdf-lezer heeft eigen tests, niet '{orders_als}'")
    if sjabloon not in ("csv", "xlsx"):
        raise ValueError(f"sjabloon is 'csv' of 'xlsx', niet '{sjabloon}'")
    if blok not in BLOKKEN or entiteiten not in (1, 2) or rekeningen not in (1, 2):
        raise ValueError("blok is week, 4 weken of maand; entiteiten en rekeningen zijn 1 of 2")
    onbekend = set(afwijkingen) - set(AFWIJKINGEN)
    if onbekend:
        raise ValueError(f"onbekende afwijking: {', '.join(sorted(onbekend))}")
    if set(zonder_orders) - set(medewerkers):
        raise ValueError("zonder_orders noemt een medewerker die niet in medewerkers staat")
    map_ = Path(map_)
    n = len(medewerkers)
    plan = _plan(afwijkingen, n, weken)
    a = set(afwijkingen)

    # ---------- de dagen: wat het bureau factureerde en wat op de orders staat ----------
    fact = [[] for _ in range(n)]
    order_regels = [[] for _ in range(n)]
    for m in range(n):
        tarief = tarieven[m % len(tarieven)]
        for w in range(1, weken + 1):
            for d in range(DAGEN_PER_WEEK):
                datum = _dag(w, d)
                fact[m].append(_factuurregel(datum, tarief, _blok(datum, blok)))
                order_regels[m].append(_regel_uren(datum, tarief))
    laatste: dict[str, pd.Timestamp] = {}  # laatste werkdag van elk blok
    for r in fact[0]:
        laatste[r["groep"]] = max(laatste.get(r["groep"], r["datum"]), r["datum"])
    perioden = list(laatste)

    def van_week(lijst: list[dict], week: int, extra=lambda r: True) -> list[dict]:
        return [r for r in lijst if _week(r["datum"]) == _week(_dag(week, 0)) and extra(r)]

    geen_order: set[tuple[int, str]] = set()
    onbetaald: set[tuple[int, str]] = set()
    gegevens: dict[str, dict] = {}  # per afwijking wat de bouwer nog moet invullen na het bouwen van de orders
    urenstaat_afwijkend = []
    krediet: _Order | None = None
    krediet_betaler = None

    for naam, (m, week) in plan.items():
        donderdag = _dag(week, 3)
        wie, b = medewerkers[m], _blok(_dag(week, 0), blok)
        weeknaam = _week(_dag(week, 0))
        gegevens[naam] = dict(m=m, week=week, blok=b, groepen={b})
        eigen_order, eigen_fact = order_regels[m], fact[m]
        if naam == "blok_zonder_order":
            geen_order.add((m, b))
            uit = [r for r in eigen_fact if r["groep"] == b]
            gegevens[naam] |= dict(bedrag=round(sum(r["bedrag"] for r in uit), 2),
                                   weken=tuple(dict.fromkeys(_week(r["datum"]) for r in uit)))  # fmt: skip
        elif naam == "dag_zonder_order":
            dag = next(r for r in eigen_order if r["datum"] == donderdag)
            eigen_order.remove(dag)
            gegevens[naam] |= dict(bedrag=dag["bedrag"], weken=(weeknaam,))
        elif naam == "niet_gefactureerd":
            dag = next(r for r in eigen_fact if r["datum"] == donderdag)
            eigen_fact.remove(dag)
            gegevens[naam] |= dict(bedrag=-dag["bedrag"], weken=(weeknaam,))
        elif naam in TARIEFVERSCHIL:
            for r in van_week(eigen_order, week):
                r["tarief"] = round(r["tarief"] + TARIEFVERSCHIL[naam], 2)
                r["bedrag"] = round(r["aantal"] * r["tarief"], 2)
            fact_som = sum(r["bedrag"] for r in van_week(eigen_fact, week))
            order_som = sum(r["bedrag"] for r in van_week(eigen_order, week))
            gegevens[naam] |= dict(bedrag=round(fact_som - order_som, 2), weken=(weeknaam,))
        elif naam == "dagvergoeding_niet_op_order":
            weeknamen = []
            for w in (week, week + 1):
                weeknamen.append(_week(_dag(w, 0)))
                for d in range(DAGEN_PER_WEEK):
                    datum = _dag(w, d)
                    eigen_fact.append(dict(datum=datum, uren=0.0, tarief=DAGVERGOEDING, bedrag=DAGVERGOEDING,
                                           soort="Dagvergoeding", groep=_blok(datum, blok)))  # fmt: skip
                    gegevens[naam]["groepen"].add(_blok(datum, blok))
            gegevens[naam] |= dict(bedrag=round(2 * DAGEN_PER_WEEK * DAGVERGOEDING, 2), weken=tuple(weeknamen))
        elif naam == "dubbel_gefactureerd":
            dag = next(r for r in eigen_fact if r["datum"] == donderdag)
            eigen_fact.append({**dag, "groep": ("dubbel", b)})
            goed = round(sum(r["bedrag"] for r in van_week(eigen_fact, week) if r["groep"] == b), 2)
            urenstaat_afwijkend.append((wie, weeknaam, goed))
            gegevens[naam]["groepen"].add(("dubbel", b))
            gegevens[naam] |= dict(bedrag=dag["bedrag"], weken=(weeknaam,))
        elif naam == "order_niet_betaald":
            onbetaald.add((m, b))
            gegevens[naam] |= dict(weken=(weeknaam,))
        elif naam == "creditorder_later_verrekend":
            maandag = _dag(week, 0)
            regel = _regel_uren(maandag, CREDIT_TARIEF) | dict(omschrijving="Correctie tarief")
            krediet = _Order("CR-0001", m, b, pd.NaT, "", [regel])
            krediet_betaler = m + 1
            gegevens[naam] |= dict(bedrag=-regel["bedrag"], weken=(weeknaam,))
        else:  # vervangen_order en dubbele_order: ze hangen aan de order van het blok
            gegevens[naam] |= dict(weken=(weeknaam,))

    # ---------- de orders ----------
    entiteit_van = [ENTITEITEN[m % entiteiten][0] for m in range(n)]
    orders: list[_Order] = []
    volgnummer = 0
    for b in perioden:
        for m in range(n):
            if medewerkers[m] in zonder_orders or (m, b) in geen_order:
                continue
            regels = [r for r in order_regels[m] if _blok(r["datum"], blok) == b]
            if not regels:
                continue
            volgnummer += 1
            orders.append(_Order(f"OR-{volgnummer:04d}", m, b, laatste[b] + pd.Timedelta(days=4), entiteit_van[m],
                                 regels, (m, b) not in onbetaald))  # fmt: skip
    gewoon = {(o.m, o.blok): o for o in orders}
    if krediet:
        betaler = gewoon[(krediet_betaler, krediet.blok)]
        # De verrekening hoort bij één entiteit: de creditorder staat op die van de order waarmee hij wordt verrekend.
        krediet.datum, krediet.entiteit = gewoon[(krediet.m, krediet.blok)].datum, betaler.entiteit
    vervallen = []
    extra_orders: list[_Order] = []
    for naam in ("vervangen_order", "dubbele_order"):
        if naam not in gegevens:
            continue
        g = gegevens[naam]
        nieuw = gewoon[(g["m"], g["blok"])]
        g |= dict(order=nieuw.nr, bedrag=0.0)
        if naam == "vervangen_order":
            volgnummer += 1
            oud = _Order(f"OR-{volgnummer:04d}", nieuw.m, nieuw.blok, nieuw.datum - pd.Timedelta(days=7),
                         nieuw.entiteit, [dict(r) for r in nieuw.regels], False)  # fmt: skip
            extra_orders.append(oud)
            vervallen.append((oud.nr, nieuw.nr, "Vervangen door een latere order met dezelfde regels."))
            g |= dict(order=oud.nr, extra=nieuw.nr, bedrag=nieuw.excl)
    if "order_niet_betaald" in gegevens:
        g = gegevens["order_niet_betaald"]
        g |= dict(order=gewoon[(g["m"], g["blok"])].nr, bedrag=gewoon[(g["m"], g["blok"])].excl)
    for naam, g in gegevens.items():
        if naam not in ("vervangen_order", "dubbele_order", "order_niet_betaald", "overlappende_bankexport"):
            o = gewoon.get((g["m"], g["blok"]))
            g["order"] = o.nr if o else ""
    if krediet:
        gegevens["creditorder_later_verrekend"]["order"] = krediet.nr
    alle_orders = orders + extra_orders + ([krediet] if krediet else [])

    # ---------- de facturen ----------
    uren_rijen, factuurnummers = [], {}
    nummer = 0
    for groep in [*perioden, *[("dubbel", b) for b in perioden]]:
        for m in range(n):
            regels = [r for r in fact[m] if r["groep"] == groep]
            if not regels:
                continue
            nummer += 1
            b = groep if isinstance(groep, str) else groep[1]
            datum = laatste[b] + pd.Timedelta(days=1 if isinstance(groep, str) else 8)
            fnr = f"{JAAR}-{nummer:04d}"
            factuurnummers[(m, groep)] = (fnr, datum)
            for r in regels:
                uren_rijen.append(dict(
                    medewerker=medewerkers[m], urenstaat=f"TS-{m + 1}{_week(r['datum'])[-2:]}", datum=r["datum"],
                    uren=r["uren"], soort=r["soort"], tarief=r["tarief"], status="Goedgekeurd", factuurnummer=fnr,
                    factuurdatum=datum, bedrag=r["bedrag"], creditnummer="", creditbedrag=0.0,
                    factuurbedrag=r["bedrag"],
                ))  # fmt: skip
    uren = pd.DataFrame(uren_rijen)

    # ---------- de bank ----------
    rekening_lijst = _rekeningen(rekeningen)
    entiteit_naam = dict(ENTITEITEN)

    def betaaldag(o: _Order) -> pd.Timestamp:
        return o.datum + pd.Timedelta(days=3 + (0 if gelijk_betaald else o.m))

    betalingen = []  # datum, iban, bedrag, tegenpartij, omschrijving
    krediet_incl = _incl(krediet.excl, btw) if krediet else 0.0
    for o in orders:
        if not o.betaald:
            continue
        betaler = krediet and (o.m, o.blok) == (krediet_betaler, krediet.blok)
        datum = betaaldag(o)
        for _, iban, aandeel in rekening_lijst:
            bedrag = _deel(_incl(o.excl, btw), aandeel)
            if betaler:
                bedrag = round(bedrag + _deel(krediet_incl, aandeel), 2)
            betalingen.append((datum, iban, bedrag, f"{entiteit_naam[o.entiteit]} B.V.", f"Betaling {o.nr}"))
    betalingen.sort(key=lambda x: (x[0], x[1]))
    if bank and "order_niet_betaald" in gegevens and betalingen:
        order = gewoon[(gegevens["order_niet_betaald"]["m"], gegevens["order_niet_betaald"]["blok"])]
        if order.datum + pd.Timedelta(days=betaaltermijn) > betalingen[-1][0]:
            raise ValueError(
                "order_niet_betaald: de order is bij het einde van de bankexport nog niet vervallen: meer weken"
            )
    bankrij = pd.DataFrame(betalingen, columns=["datum", "rekening", "bedrag", "tegenpartij", "omschrijving"])
    bankrij["referentie"] = [f"REF{i:05d}" for i in range(1, len(bankrij) + 1)]
    if "overlappende_bankexport" in a:
        tweede = bankrij.iloc[len(bankrij) // 2 :]
        bankrij = pd.concat([bankrij, tweede], ignore_index=True)
        gegevens["overlappende_bankexport"] = dict(m=0, week=0, blok="", weken=(), bedrag=0.0, order="",
                                                   aantal=len(tweede))  # fmt: skip

    # ---------- de kaart: wat niet is afgeletterd ----------
    kaartrijen = []
    gewone_iban = rekening_lijst[-1][1]
    for m in range(n):
        for b in perioden:
            fnr, fdatum = factuurnummers.get((m, b), (None, None))
            if fnr is None:
                continue
            incl_f = _incl(sum(r["bedrag"] for r in fact[m] if r["groep"] == b), btw)
            o = gewoon.get((m, b))
            betaler = krediet and (m, b) == (krediet_betaler, krediet.blok)
            if betaler or (kaart_ongelet and o is not None and o.betaald):
                kaartrijen.append((fdatum, "factuur", fnr, "", "", incl_f))
                for _, iban, aandeel in rekening_lijst:
                    deel = _deel(_incl(o.excl, btw), aandeel) + (_deel(krediet_incl, aandeel) if betaler else 0)
                    kaartrijen.append((betaaldag(o), "ontvangst", "", "", iban, -round(deel, 2)))
                continue
            betaald = _incl(o.excl, btw) if o is not None and o.betaald else 0.0
            rest = round(incl_f - betaald, 2)
            if rest > 0.005:
                kaartrijen.append((fdatum, "factuur", fnr, "", "", rest))
            elif rest < -0.005 and gegevens.get("restant_op_kaart", {}).get("m") == m \
                    and gegevens["restant_op_kaart"]["blok"] == b:  # fmt: skip
                # wat de bank binnenkreeg is de som van de afgeronde delen per rekening, en dat kan een cent afwijken
                werkelijk = sum(_deel(_incl(o.excl, btw), aandeel) for _, _, aandeel in rekening_lijst)
                kaartrijen.append((betaaldag(o), "ontvangst", "", "", gewone_iban,
                                   round(incl_f - werkelijk, 2)))  # fmt: skip
    for (m, groep), (fnr, fdatum) in factuurnummers.items():
        if isinstance(groep, tuple):  # een tweede factuur voor dezelfde dag: niemand betaalde hem
            incl_f = _incl(sum(r["bedrag"] for r in fact[m] if r["groep"] == groep), btw)
            kaartrijen.append((fdatum, "factuur", fnr, "", "", incl_f))
    kaartrij = pd.DataFrame(kaartrijen, columns=["datum", "soort", "nummer", "omschrijving", "rekening", "bedrag"])
    kaartrij = kaartrij.sort_values(["datum", "soort", "nummer", "rekening"], kind="stable").reset_index(drop=True)

    # ---------- wat de bouwer verwacht ----------
    def betaald_ex(o: _Order) -> float:
        """Wat de bank voor deze order binnenkreeg, ex btw: het ordertotaal naar de verhouding van wat per rekening
        (afgerond) is betaald en het bedrag incl. btw. Dat verschilt hooguit een cent van het bedrag van de order."""
        incl = _incl(o.excl, btw)
        ontvangen = sum(_deel(incl, aandeel) for _, _, aandeel in rekening_lijst)
        return round(o.excl * ontvangen / incl, 2) if incl else 0.0

    verwacht = {}
    for m, wie in enumerate(medewerkers):
        eigen = [o for o in orders + ([krediet] if krediet else []) if o.m == m]
        verwacht[wie] = dict(
            gefactureerd=round(sum(r["bedrag"] for r in fact[m]), 2),
            op_orders=round(sum(o.excl for o in eigen), 2),
            betaald=round(sum(betaald_ex(o) for o in eigen if o.betaald), 2) if bank else 0.0,
        )
        verwacht[wie]["urenstaat"] = round(verwacht[wie]["gefactureerd"] - sum(
            g["bedrag"] for nm, g in gegevens.items() if nm == "dubbel_gefactureerd" and g["m"] == m), 2)  # fmt: skip

    # ---------- schrijven ----------
    kop = pd.DataFrame(
        [dict(bestand=f"{o.nr}.pdf", order=o.nr, entiteit=o.entiteit, factuurdatum=o.datum, referentie=o.blok,
              excl=o.excl, incl=_incl(o.excl, btw)) for o in alle_orders if o.regels]
    )  # fmt: skip
    regels = pd.DataFrame(
        [dict(bestand=f"{o.nr}.pdf", order=o.nr, entiteit=o.entiteit, medewerker=medewerkers[o.m],
              project="Project Noord", **r) for o in alle_orders for r in o.regels]
    )  # fmt: skip
    bronnen = dict(
        instellingen=_instellingen(entiteiten, rekeningen, btw, blok, betaaltermijn), uren=uren, kop=kop, regels=regels
    )
    if bank:
        bronnen["bank"] = bankrij
    if kaart:
        bronnen["kaart"] = kaartrij
    if vervallen:
        bronnen["vervallen"] = pd.DataFrame(vervallen, columns=["order", "vervangen_door", "toelichting"])
    if urenstaat_afwijkend:
        bronnen["urenstaat_afwijkend"] = pd.DataFrame(
            [(wie, week, bedrag, "De goedgekeurde urenstaat telt de dag één keer; een tweede factuur is te veel.")
             for wie, week, bedrag in urenstaat_afwijkend],
            columns=inlezen.URENSTAAT_AFWIJKEND,
        )  # fmt: skip
    dossier.schrijf(map_, **bronnen)
    if sjabloon == "xlsx":
        _naar_excel(map_)

    uit = {}
    for naam in AFWIJKINGEN:
        if naam in gegevens:
            g = gegevens[naam]
            facturen = tuple(sorted(factuurnummers[(g["m"], x)][0] for x in g.get("groepen", ())
                                    if (g["m"], x) in factuurnummers))  # fmt: skip
            uit[naam] = Afwijking(naam, medewerkers[g["m"]], g.get("weken", ()), g["blok"], g.get("order", ""),
                                  g.get("bedrag", 0.0), g.get("aantal", 0), g.get("extra", ""), facturen)  # fmt: skip
    return Scenario(
        map_, tuple(medewerkers), btw, blok, perioden, uit, {(medewerkers[o.m], o.blok): o.nr for o in orders},
        verwacht, bank, kaart, tuple(zonder_orders),
        incl={o.nr: _incl(o.excl, btw) for o in alle_orders},
        dubbele=[gegevens["dubbele_order"]["order"]] if "dubbele_order" in gegevens else [],
    )  # fmt: skip


# ---------- Excel ----------

_TEKSTKOLOMMEN = {"nr", "nummer", "referentie", "order", "urenstaat", "factuurnummer", "creditnummer", "medewerker"}


def _naar_excel(map_: Path) -> None:
    """Schrijft elk csv-sjabloon in de map als werkboek: datums en bedragen als echte waarden, de rest als tekst. De
    csv's verdwijnen."""
    for pad in sorted(map_.glob("*.csv")):
        df = pd.read_csv(pad, sep=";", dtype=str, keep_default_na=False)
        for k in df.columns:
            waarden = df[k][df[k] != ""]
            if waarden.empty or k in _TEKSTKOLOMMEN:
                continue
            if waarden.str.fullmatch(r"\d{2}-\d{2}-\d{4}").all():
                df[k] = [datetime.datetime.strptime(v, "%d-%m-%Y") if v else None for v in df[k]]
            elif waarden.str.fullmatch(r"-?\d+(,\d+)?").all():
                df[k] = [float(v.replace(",", ".")) if v else None for v in df[k]]
        df.to_excel(pad.with_suffix(".xlsx"), index=False)
        pad.unlink()


# ---------- de uitvoer van de opdrachtregel ----------

VERDIEPING = ("Opmerkingen", "Per week", "Per dag", "Orders", "Facturen", "Bank", "Urenregels", "Orderregels")
VOORKANT_TOTAAL = ("Samenvatting", "Actielijst", "Zonder order alle")
KOLOMMEN_BEVINDINGEN = bevindingen.KOLOMMEN
HOOFDBLADEN = ("Stand", "Zonder order", "Debiteurenkaart", "Per periode", "Verschillen")  # letterlijk: niet uit de code
KOP_BLOK_1 = "1. Van urenstaat tot betaling"
GEEN_ACTIES = "Er zijn geen acties: alles sluit aan."


def controleer(map_: Path, uit: Path | None = None) -> dict:
    """Draait `cli.main(["aansluiten", map_, "--uit", uit])` en doet de vaste controles van elk scenario.

    De opdracht eindigt met 0 en schrijft per medewerker een werkboek, het totaalbestand en `bevindingen.csv`. Per
    medewerker tellen de verschillen op tot gefactureerd min op orders en is betaald nooit meer dan op orders (uit het
    model). De bladen van elk werkboek zijn zichtbaar, verborgen of afwezig zoals het hoort: Stand, Zonder order, Per
    periode en Verschillen zichtbaar; Debiteurenkaart zichtbaar als er een kaart en een bank is en anders afwezig; Bank
    verborgen als er een bank is en anders afwezig; Opmerkingen verborgen als er een bevinding voor de medewerker is; de
    rest verborgen. Het totaalbestand heeft vooraan Samenvatting, Actielijst en Zonder order alle, een kolom per
    medewerker plus Totaal in blok 1 en per medewerker de bladen Stand en Zonder order. `bevindingen.csv` is leesbaar.

    Geeft een dict met `uit` (de map), `overzichten` (naam naar `Overzicht`), `bevindingen` (de tabel uit het csv),
    `werkboeken` (naam naar werkboek) en `totaal` (het totaalbestand). Doorrekenen met `formulas` doet `reken_na`."""
    map_ = Path(map_)
    uit = Path(uit) if uit else map_.parent / f"{map_.name}-uit"
    assert main(["aansluiten", str(map_), "--uit", str(uit)]) == 0
    d = dossier.lees(map_)
    namen = sorted(d.uren.medewerker.unique())
    overzichten = {}
    for naam in namen:
        a = aansluiten.sluit_aan(
            d.uren, d.kop, d.regels, naam, d.inst, bankregels=d.bank, vervallen=d.vervallen, saldo=d.saldo(naam)
        )
        overzichten[naam] = overzicht.maak(a, d)

    gevonden = pd.read_csv(uit / "bevindingen.csv", sep=";", decimal=",", keep_default_na=False)
    assert list(gevonden.columns) == KOLOMMEN_BEVINDINGEN
    assert gevonden.sleutel.is_unique

    werkboeken = {}
    for naam, o in overzichten.items():
        stand, v = o.stand, o.verschillen
        assert round((v.fact - v.order).sum(), 2) == round(stand["gefactureerd"] - stand["op_orders"], 2), naam
        assert stand["betaald"] <= stand["op_orders"] + 0.01 * len(o.orders), naam
        wb = werkboeken[naam] = load_workbook(uit / f"Overzicht {dossier.in_naam(naam)}.xlsx")
        met_bank, met_kaart = d.bank is not None, d.kaart is not None and d.bank is not None
        verwacht_zichtbaar = [b for b in HOOFDBLADEN if b != "Debiteurenkaart" or met_kaart]  # de namen staan hierboven
        assert [ws.title for ws in wb if ws.sheet_state == "visible"] == verwacht_zichtbaar, naam
        assert ("Debiteurenkaart" in wb.sheetnames) == met_kaart, naam
        assert ("Bank" in wb.sheetnames) == met_bank, naam
        assert ("Opmerkingen" in wb.sheetnames) == bool((gevonden.medewerker == naam).any()), naam
        verborgen = [ws.title for ws in wb if ws.sheet_state != "visible"]
        assert set(verborgen) <= set(VERDIEPING) and {"Per week", "Per dag", "Orders", "Urenregels"} <= set(verborgen)
        assert [ws.title for ws in wb][: len(verwacht_zichtbaar)] == verwacht_zichtbaar, naam  # de hoofdbladen vooraan

    wt = load_workbook(uit / "Overzicht totaal.xlsx")
    assert wt.sheetnames[:3] == list(VOORKANT_TOTAAL)
    assert all(ws.sheet_state == "visible" for ws in wt)
    korte = totaal.korte_namen(namen)
    for naam in namen:
        assert {f"Stand {korte[naam]}", f"Zonder order {korte[naam]}"} <= set(wt.sheetnames), naam
    inst = d.inst
    blad = [list(r) for r in wt[totaal.TOT].iter_rows(values_only=True)]
    kop = next(r for r in blad if isinstance(r[0], str) and r[0].startswith(KOP_BLOK_1))
    assert kop[1 : len(namen) + 2] == [*(korte[n] for n in namen), "Totaal"]  # een kolom per medewerker plus Totaal
    regels = _regels_blok_1(inst, d.bank is not None)
    for begin, sleutel in regels.items():
        rij = next(r for r in blad if isinstance(r[0], str) and r[0].startswith(begin))
        for i, naam in enumerate(namen, start=1):
            assert _waarde(wt, rij[i]) == pytest.approx(overzichten[naam].stand[sleutel], abs=0.005), (begin, naam)
        assert str(rij[len(namen) + 1]).startswith("=SUM(")  # het totaal is de som, als formule
    if d.bank is None:
        assert not [r for r in blad if isinstance(r[0], str) and r[0].startswith(f"Door {inst.opdrachtgever} betaald")]
    for naam in namen:  # en de bladen Stand per medewerker (als waarden geschreven) hebben dezelfde getallen
        stand = [list(r) for r in wt[f"Stand {korte[naam]}"].iter_rows(values_only=True)]
        for begin, sleutel in regels.items():
            rij = next(r for r in stand if isinstance(r[0], str) and r[0].startswith(begin))
            assert rij[1] == pytest.approx(overzichten[naam].stand[sleutel], abs=0.005), (begin, naam)
    actielijst = [r[0] for r in wt["Actielijst"].iter_rows(values_only=True)]
    assert (GEEN_ACTIES in actielijst) == gevonden.empty
    return dict(uit=uit, overzichten=overzichten, bevindingen=gevonden, werkboeken=werkboeken, totaal=wt)


def _regels_blok_1(inst, met_bank: bool) -> dict[str, str]:
    """Het begin van elke rij van blok 1 (in Totaal uitgebreid en op blad Stand) en de sleutel in `Overzicht.stand`."""
    return {
        "Gewerkt volgens": "urenstaat",
        f"Door {inst.bureau} gefactureerd": "gefactureerd",
        f"Door {inst.opdrachtgever} op orders gezet": "op_orders",
        **({f"Door {inst.opdrachtgever} betaald": "betaald"} if met_bank else {}),
    }


def _waarde(wb, cel):
    """De waarde van een cel; staat er een verwijzing naar één cel van een ander blad (`='Stand Sanne'!B7`), dan de
    waarde daar. Een getal blijft een getal."""
    if isinstance(cel, str):
        m = re.fullmatch(r"='?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)", cel)
        if m:
            return wb[m.group(1)][f"{m.group(2)}{m.group(3)}"].value
    return cel


def _doorgerekend(pad: Path) -> dict[str, list[list]]:
    """Het werkboek doorgerekend met `formulas`: per blad (bladnaam in hoofdletters) de rijen met waarden."""
    formulas = pytest.importorskip("formulas")
    with tempfile.TemporaryDirectory() as tijdelijk:
        model = formulas.ExcelModel().loads(str(pad)).finish()
        model.calculate()
        model.write(dirpath=tijdelijk)
        wb = load_workbook(next(Path(tijdelijk).glob("*")), data_only=True)
        return {ws.title.upper(): [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb}


def _rijen_van_blok(rijen: list[list], nummer: int) -> list[list]:
    """De rijen van blok `nummer` van Totaal uitgebreid, zonder de kop, tot de eerste lege rij."""
    kop = next(i for i, r in enumerate(rijen) if isinstance(r[0], str) and r[0].startswith(f"{nummer}. "))
    blok = []
    for r in rijen[kop + 1 :]:
        if not any(x is not None for x in r):
            break
        blok.append(r)
    return blok


def reken_na(uit: Path, *bestanden: str, overzichten: dict | None = None) -> None:
    """Rekent de werkboeken in `uit` door met `formulas` (alleen de genoemde bestanden, anders alle): geen cel die met
    `#` begint en nergens "sluit niet aan". Traag: roep dit voor een paar scenario's aan, niet voor alle.

    Met `overzichten` (naam naar `Overzicht`, zoals `controleer` ze geeft) worden ook de getallen nagelopen: in het
    totaalbestand is de kolom Totaal van blok 1, 2 en 4 de som van de kolommen van de medewerkers, en in het
    doorgerekende blad Stand van elke medewerker komen blok 1 en het totaal van blok 2 overeen met het model."""
    vergeleken = 0
    for pad in sorted(Path(uit).glob("*.xlsx")):
        if bestanden and pad.name not in bestanden:
            continue
        bladen = _doorgerekend(pad)
        for titel, rijen in bladen.items():
            for rij in rijen:
                for cel in rij:
                    assert not (isinstance(cel, str) and cel.startswith("#")), (pad.name, titel, cel)
                    assert not (isinstance(cel, str) and "sluit niet aan" in cel), (pad.name, titel, cel)
        if overzichten is None:
            continue
        n = len(overzichten)
        if pad.name == "Overzicht totaal.xlsx":
            for nummer in (1, 2, 4):
                for rij in _rijen_van_blok(bladen["TOTAAL UITGEBREID"], nummer):
                    if all(isinstance(x, int | float) for x in rij[1 : n + 2]):
                        assert rij[n + 1] == pytest.approx(sum(rij[1 : n + 1]), abs=0.005), (nummer, rij[0])
                        vergeleken += 1
        else:
            o = next((o for n, o in overzichten.items() if pad.name == f"Overzicht {dossier.in_naam(n)}.xlsx"), None)
            if o is None:
                continue
            stand = bladen["STAND"]
            for begin, sleutel in _regels_blok_1(o.inst, o.bank is not None).items():
                rij = next(r for r in _rijen_van_blok(stand, 1) if r[0].startswith(begin))
                assert rij[1] == pytest.approx(o.stand[sleutel], abs=0.005), (pad.name, sleutel)
            totaal_2 = next(r for r in _rijen_van_blok(stand, 2) if r[0] == "Totaal")[1]
            assert totaal_2 == pytest.approx(round(o.stand["gefactureerd"] - o.stand["op_orders"], 2), abs=0.005)
            vergeleken += 1
    if overzichten is not None:  # een controle die niets vergelijkt, bewijst niets
        assert vergeleken > 0, "reken_na heeft geen enkele rij vergeleken: bestanden of overzichten kloppen niet"
